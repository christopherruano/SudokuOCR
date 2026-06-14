#!/usr/bin/env python3
"""MASTER TRACKING DOCUMENT: one row per crop image we have, mapping it to its volume
and printed page, with everything relevant up front — exact title, what the page
contains (sections/religions, age range, columns), where it sits in its multi-page
table, and its pipeline status (OCR'd? validated? assembled? complete?).
Outputs master_tracking.csv + master_tracking.xlsx (Images + Volumes sheets)."""
import os, re, json, csv, glob
from collections import defaultdict

csv.field_size_limit(10 ** 7)
ROOT = os.path.dirname(os.path.abspath(__file__))
V1 = os.path.join(ROOT, "export_final/sendout_final/sendout_v1")
P = os.path.join(V1, "pageno_parts")
HEX = re.compile(r"([0-9a-f]{40})")
DOC = re.compile(r"digitized_tables/([0-9a-f]{16})_t(\d+)_([^/]*)/")
AGE = re.compile(r"(\d+)\s*[-–]\s*(\d+)|(\d+)\s*and over|under\s*(\d+)", re.I)


def load():
    # 1) every crop image -> identity (dedup by imghash; prefer referenced/results dir)
    crops = {}
    order = {"ocr_results_referenced": 0, "digitized_tables_referenced": 1, "suspicious_review": 2, "never_ocrd_images": 3}
    for root in order:
        for pth in glob.glob(f"export/{root}/digitized_tables/*/*.png"):
            m = HEX.search(os.path.basename(pth)); dm = DOC.search(pth)
            if not m or not dm:
                continue
            ih = m.group(1)
            if ih in crops and order[root] >= order[crops[ih]["root"]]:
                continue
            crops[ih] = {"path": pth, "root": root, "dochash": dm.group(1),
                         "tseq": dm.group(2), "dname": dm.group(3)}
    # 2) merged authoritative page numbers
    pageno = {}
    for ln in open(os.path.join(P, "gemini_reads.jsonl")):
        d = json.loads(ln); pn = d.get("page_number")
        if isinstance(pn, int) and d.get("page_visible"):
            pageno[d["imghash"]] = (pn, "gemini")
    for ln in open(os.path.join(P, "roman_reads.jsonl")):
        d = json.loads(ln); pn = d.get("page_number_int")
        if isinstance(pn, int):
            pageno[d["imghash"]] = (pn, "opus")
    for ih, v in json.load(open(os.path.join(P, "human_pageno_decisions.json"))).items():
        if isinstance(v.get("page_number"), int):
            pageno[ih] = (v["page_number"], "human")
    # 3) table_id + family/number + AUTHORITATIVE page number per imghash (from
    #    page_reads_fixed.json — the user's consolidated page-ordering work; this
    #    OVERRIDES the per-engine reads above as the primary page source).
    tid_by_ih, fam_by_tid = {}, {}
    for r in json.load(open(os.path.join(V1, "page_reads_fixed.json"))):
        m = HEX.search(r.get("table_id", ""))
        if m:
            tid_by_ih[m.group(1)] = r["table_id"]
            fam_by_tid[r["table_id"]] = (r.get("family", ""), r.get("table_number", ""))
            pn = r.get("page_number")
            if isinstance(pn, int) or (str(pn).lstrip("-").isdigit()):
                pageno[m.group(1)] = (int(pn), "page_reads_fixed")
    # 4) validation + verdict per table_id
    sc = {r["table_id"]: r for r in csv.DictReader(open(os.path.join(V1, "tableVII_scorecard.csv")))} \
        if os.path.exists(os.path.join(V1, "tableVII_scorecard.csv")) else {}
    vd = {r["table_id"]: r for r in csv.DictReader(open(os.path.join(V1, "tableVII_verdicts.csv")))} \
        if os.path.exists(os.path.join(V1, "tableVII_verdicts.csv")) else {}
    # 5) assembly membership + position
    members = json.load(open(os.path.join(V1, "assembled_members.json"))) \
        if os.path.exists(os.path.join(V1, "assembled_members.json")) else {}
    asm_of = {}     # member_tid -> (asm_id, ordered_member_tids)
    for asm, mem in members.items():
        # order members by page number
        mo = sorted(mem, key=lambda t: (pageno.get(HEX.search(t).group(1), (9999, ""))[0] if HEX.search(t) else 9999))
        for t in mem:
            asm_of[t] = (asm, mo)
    # 6) document meta
    docs = {r["dochash"]: r for r in csv.DictReader(open(os.path.join(V1, "documents.csv")))} \
        if os.path.exists(os.path.join(V1, "documents.csv")) else {}
    # 7) ORIGINAL PDF page offset per table (from the digitization manifest's
    #    resolved_start_page) -> per-volume printed->PDF offset (primary working key)
    DT = re.compile(r"([0-9a-f]{16})_t(\d+)")
    pdf_start = {}     # dochash_tNNNN -> resolved PDF start page
    mani = os.path.join(ROOT, "export/cluster_sync/organized_manifest.jsonl")
    if os.path.exists(mani):
        for line in open(mani):
            try:
                j = json.loads(line)
            except Exception:
                continue
            mm = DT.search(j.get("source_image", "") or j.get("result_json", "") or "")
            s = (j.get("queue_row", {}) or {}).get("resolved_start_page")
            if mm and isinstance(s, int):
                pdf_start[f"{mm.group(1)}_t{mm.group(2)}"] = s
    import statistics
    vol_offset = {}    # dochash -> median(pdf_page - printed_page)
    diffs = defaultdict(list)
    for tid, (fam, tnum) in fam_by_tid.items():
        mm = DT.search(tid); ihm = HEX.search(tid)
        if not mm or not ihm:
            continue
        key = f"{mm.group(1)}_t{mm.group(2)}"
        if key in pdf_start and ihm.group(1) in pageno:
            diffs[mm.group(1)].append(pdf_start[key] - pageno[ihm.group(1)][0])
    for dh, ds in diffs.items():
        vol_offset[dh] = int(statistics.median(ds))
    return crops, pageno, tid_by_ih, fam_by_tid, sc, vd, asm_of, docs, pdf_start, vol_offset


_ONESHOT_INDEX = None


def _build_oneshot_index():
    """imghash -> oneshot.json path, built ONCE (single tree walk)."""
    idx = {}
    for pth in glob.glob("export/*/digitized_tables/*/*_oneshot.json"):
        m = HEX.search(os.path.basename(pth))
        if m and m.group(1) not in idx:
            idx[m.group(1)] = pth
    return idx


def oneshot_for(ih):
    global _ONESHOT_INDEX
    if _ONESHOT_INDEX is None:
        _ONESHOT_INDEX = _build_oneshot_index()
    pth = _ONESHOT_INDEX.get(ih)
    if not pth:
        return None
    try:
        return json.load(open(pth))
    except Exception:
        return None


def content_of(js):
    """What the page contains: sections, age range, columns, rows, acceptance."""
    if not js:
        return {"title": "", "sections": "", "age_range": "", "columns": "", "n_rows": 0, "accept": ""}
    md = js.get("metadata", {})
    secs = [s.get("name", "") for s in js.get("data", []) if s.get("name")]
    lo, hi, nr = 999, -1, 0
    for s in js.get("data", []):
        for r in s.get("rows", []):
            nr += 1
            lab = " ".join(str(v) for v in r.values() if isinstance(v, str))
            mm = AGE.search(lab)
            if mm:
                a = int(mm.group(1) or mm.group(3) or mm.group(4) or 0)
                b = int(mm.group(2) or mm.group(3) or 99)
                lo = min(lo, a); hi = max(hi, b)
    age = f"{lo}-{hi}" if hi >= 0 else ""
    acc = (js.get("acceptance", {}) or {}).get("decision", "")
    return {"title": md.get("title", ""), "sections": "; ".join(secs[:6]),
            "age_range": age, "columns": "; ".join(md.get("column_groups", [])),
            "n_rows": nr, "accept": acc}


def main():
    crops, pageno, tid_by_ih, fam_by_tid, sc, vd, asm_of, docs, pdf_start, vol_offset = load()
    DT = re.compile(r"([0-9a-f]{16})_t(\d+)")
    VII = re.compile(r"VI[IL]|civil|age.{0,3}sex|C-?III", re.I)
    rows = []
    for ih, c in crops.items():
        tid = tid_by_ih.get(ih, "")
        js = oneshot_for(ih)
        cont = content_of(js)
        fam, tnum = fam_by_tid.get(tid, ("", ""))
        is_vii = bool(VII.search(c["dname"]) or VII.search(fam) or VII.search(cont["title"]))
        pg, psrc = pageno.get(ih, ("", ""))
        # PDF page (the primary working key) = book page + the volume's printed->PDF
        # offset (per-crop, consistent). Falls back to the manifest's resolved_start_page
        # only when this crop has no read book page.
        dh = c["dochash"]; off = vol_offset.get(dh, "")
        mm = DT.search(tid)
        if isinstance(pg, int) and off != "":
            pdf_page = pg + off
        elif mm and pdf_start.get(f"{mm.group(1)}_t{mm.group(2)}") is not None:
            pdf_page = pdf_start[f"{mm.group(1)}_t{mm.group(2)}"]
        else:
            pdf_page = ""
        # OCR status
        if js is None:
            ocr = "NOT_OCRd"
        elif cont["accept"] in ("accepted_high_confidence",):
            ocr = "ocrd_accepted"
        elif cont["accept"]:
            ocr = "ocrd_" + cont["accept"].replace("accepted_", "")
        else:
            ocr = "ocrd"
        # assembly + position
        asm, pos, total = "", "", ""
        if tid in asm_of:
            asm, mo = asm_of[tid]
            total = len(mo)
            try:
                pos = mo.index(tid) + 1
            except ValueError:
                pos = ""
        # validation — for an assembled member, inherit the ASSEMBLED table's status
        vtid = asm if asm else tid
        s = sc.get(vtid, {}); v = vd.get(vtid, {})
        perfect = s.get("perfect", ""); issues = s.get("issues", "")
        verdict = v.get("status", "")
        # headline status
        if ocr == "NOT_OCRd":
            status = "⏳ NOT OCR'd"
        elif ocr.startswith("ocrd_") and ocr not in ("ocrd_accepted", "ocrd_with_risk"):
            status = f"❌ OCR low/failed ({cont['accept']})"
        elif asm:
            if verdict == "perfect" or perfect:
                status = f"✅ Complete — assembled (page {pos}/{total})"
            else:
                status = f"⚠ Assembled, imperfect: {issues or verdict}"
        elif is_vii:
            if perfect:
                status = "✅ Complete — single page"
            elif issues:
                status = f"⚠ Imperfect: {issues}"
            elif verdict:
                status = f"… {verdict}"
            else:
                status = "ℹ OCR'd (not in Table VII scorecard)"
        else:
            status = "ℹ Non-Table-VII OCR'd"
        dm = docs.get(c["dochash"], {})
        rows.append({
            "status": status, "volume": c["dochash"], "year": dm.get("year", ""),
            "province": dm.get("region", ""), "pdf_page": pdf_page, "offset": off,
            "printed_page": pg, "page_source": psrc,
            "table_number": tnum, "exact_title": cont["title"] or c["dname"].replace("-", " ")[:80],
            "contains_sections": cont["sections"], "age_range": cont["age_range"],
            "columns": cont["columns"], "n_rows": cont["n_rows"],
            "is_tableVII": "Y" if is_vii else "", "family": fam,
            "assembled_table": asm, "position_in_table": f"{pos}/{total}" if asm else "",
            "ocr_status": ocr, "validation": ("perfect" if perfect else issues or verdict or ""),
            "verdict": verdict, "table_id": tid, "imghash": ih,
            "image_path": c["path"], "source_pdf": os.path.basename(dm.get("source_pdf", "") or ""),
        })
    # sort: volume, then PDF page (primary working key), then printed page
    rows.sort(key=lambda r: (str(r["volume"]),
                             (r["pdf_page"] if isinstance(r["pdf_page"], int) else 9999),
                             (r["printed_page"] if isinstance(r["printed_page"], int) else 9999)))
    cols = ["status", "volume", "year", "province", "pdf_page", "offset", "printed_page", "page_source",
            "table_number", "exact_title", "contains_sections", "age_range", "columns", "n_rows",
            "is_tableVII", "family", "assembled_table", "position_in_table", "ocr_status", "validation",
            "verdict", "table_id", "imghash", "image_path", "source_pdf"]
    with open(os.path.join(V1, "master_tracking.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader(); [w.writerow(r) for r in rows]

    # volume summary
    vol = defaultdict(lambda: {"imgs": 0, "ocrd": 0, "notocrd": 0, "vii": 0, "perfect": 0,
                               "imperfect": 0, "pages": set()})
    for r in rows:
        v = vol[r["volume"]]
        v["imgs"] += 1
        v["ocrd"] += r["ocr_status"] != "NOT_OCRd"
        v["notocrd"] += r["ocr_status"] == "NOT_OCRd"
        v["vii"] += r["is_tableVII"] == "Y"
        v["perfect"] += r["validation"] == "perfect"
        v["imperfect"] += bool(r["validation"]) and r["validation"] != "perfect"
        if isinstance(r["printed_page"], int):
            v["pages"].add(r["printed_page"])
        v["year"] = r["year"]; v["province"] = r["province"]; v["pdf"] = r["source_pdf"]
    vrows = []
    for dh, v in vol.items():
        pg = v["pages"]; span = (max(pg) - min(pg) + 1) if pg else 0
        missing = span - len(pg) if pg else 0
        vrows.append({"volume": dh, "year": v.get("year", ""), "province": v.get("province", ""),
                      "source_pdf": v.get("pdf", ""), "images": v["imgs"], "ocrd": v["ocrd"],
                      "never_ocrd": v["notocrd"], "tableVII_imgs": v["vii"], "perfect": v["perfect"],
                      "imperfect": v["imperfect"], "pages_have": len(pg),
                      "page_span": f"{min(pg)}-{max(pg)}" if pg else "", "pages_missing_in_span": missing,
                      "coverage_pct": f"{100*len(pg)//span}%" if span else ""})
    vrows.sort(key=lambda r: -r["pages_missing_in_span"])
    vcols = ["volume", "year", "province", "source_pdf", "images", "ocrd", "never_ocrd",
             "tableVII_imgs", "perfect", "imperfect", "pages_have", "page_span",
             "pages_missing_in_span", "coverage_pct"]
    with open(os.path.join(V1, "master_tracking_volumes.csv"), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=vcols, extrasaction="ignore")
        w.writeheader(); [w.writerow(r) for r in vrows]

    # XLSX (2 sheets)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active; ws.title = "Images"
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(k, "") for k in cols])
        ws.freeze_panes = "A2"
        ws2 = wb.create_sheet("Volumes")
        ws2.append(vcols)
        for c in ws2[1]:
            c.font = Font(bold=True)
        for r in vrows:
            ws2.append([r.get(k, "") for k in vcols])
        ws2.freeze_panes = "A2"
        wb.save(os.path.join(V1, "master_tracking.xlsx"))
        xlsx = " + master_tracking.xlsx"
    except Exception as e:
        xlsx = f" (xlsx skipped: {e})"
    print(f"master tracker: {len(rows)} images across {len(vrows)} volumes")
    from collections import Counter
    st = Counter(r["status"].split(" (")[0].split(":")[0] for r in rows)
    for k, n in st.most_common():
        print(f"  {n:5d}  {k}")
    print(f"wrote master_tracking.csv + master_tracking_volumes.csv{xlsx}")


if __name__ == "__main__":
    main()
