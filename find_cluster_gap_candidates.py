"""Find cluster image candidates for province-year gaps in primary tables.

Uses export/data_guide.xlsx as the authoritative inventory of validated tables.
Primary coverage follows the workbook definition:
  primary_category == "AgeSexCivilCondition" and sub_category == "General"

Then scans digitized_tables/*/metadata.json (typically on FASRC) to find
candidate rendered-table directories that could fill province-year gaps where
the guide currently has no primary tables.

Typical usage on cluster:
    python3 find_cluster_gap_candidates.py \
      --guide ~/SudokuOCR/export/data_guide.xlsx \
      --digitized-root ~/census_toc/digitized_tables \
      --output-json ~/census_toc/gap_candidates.json \
      --output-csv ~/census_toc/gap_candidates.csv
"""

import argparse
import csv
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


MAJOR_CENSUS_YEARS = [1872, 1881, 1891, 1901, 1911, 1921, 1931, 1941, 1951, 1961]

RELIGION_TERMS = {
    "religion", "religions", "hindu", "hindus", "muslim", "musalmans",
    "musalman", "mohammedan", "christian", "christians", "sikh", "jain",
    "jains", "parsi", "buddhist", "arya", "arya samaj",
}

NARRATIVE_TERMS = {
    "chapter", "contents", "appendix", "appendices", "introductory",
    "introduction", "general remarks", "remarks", "reference",
    "references", "diagram", "diagrams", "illustrates", "illustration",
    "methods of recording", "proportion of adults", "proportion of the sexes",
    "commentary", "statistics of", "statistical reference",
}

OFF_TARGET_TERMS = {
    "education", "infirmities", "infirmity", "blind", "deaf",
    "deaf mutes", "deaf-mutes", "deafmute", "leper", "lepers",
    "insane", "occupation", "occupations", "language", "caste",
    "castes", "literacy", "literate", "bilingual", "mortality",
    "fertility",
}

PRIMARY_HINT_PATTERNS = [
    re.compile(r"\bage\s+sex\s+and\s+civil\s+condition\b", re.I),
    re.compile(r"\btable\s+(?:no\.\s*)?(?:vii|vil|vi1|v1i|7)\b", re.I),
    re.compile(r"\bimperial\s+table\s+(?:vii|vil|7)\b", re.I),
]

SECONDARY_HINT_PATTERNS = [
    re.compile(r"\bage\s+distribution\b", re.I),
    re.compile(r"\bage\s+by\s+religion\b", re.I),
    re.compile(r"\bage\s+sex\b", re.I),
    re.compile(r"\bmarital\b", re.I),
]

PROVINCE_ALIASES = {
    "ajmer merwara": "Ajmer Merwara",
    "andaman nicobar": "Andaman and Nicobar",
    "andaman and nicobar": "Andaman and Nicobar",
    "andaman nicobar islands": "Andaman and Nicobar",
    "assam bengal england": "Assam Bengal and England",
    "baroda state": "Baroda",
    "bengal presidency": "Bengal Presidency",
    "central provinces and berar": "Central Provinces and Berar",
    "central provinces": "Central Provinces",
    "central india": "Central India",
    "city of calcutta": "City of Calcutta",
    "hill districts and manipur": "Hill districts and Manipur",
    "jammu and kashmir": "Jammu and Kashmir",
    "kashmir": "Kashmir",
    "madras presidency": "Madras",
    "north west frontier province": "North-West Frontier Province",
    "northwest frontier province": "North-West Frontier Province",
    "north western provinces and oudh": "North Western Provinces Oudh",
    "north western provinces oudh": "North Western Provinces Oudh",
    "north western provinces of agra and oudh": "North Western Provinces Oudh",
    "north western provinces": "North Western Provinces Oudh",
    "north western provinces o": "North Western Provinces Oudh",
    "north west provinces and oudh": "North Western Provinces Oudh",
    "north west provinces oudh": "North Western Provinces Oudh",
    "northwestern provinces and oudh": "North Western Provinces Oudh",
    "northwestern provinces oudh": "North Western Provinces Oudh",
    "rajputana agency": "Rajputana",
    "united provinces of agra and oudh": "United Provinces",
    "united provinces": "United Provinces",
}


def normalize_text(value):
    value = (value or "").strip().lower()
    value = value.replace("&", " and ")
    value = value.replace("_", " ")
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return " ".join(value.split())


def canonicalize_province(value):
    text = normalize_text(value)
    if not text:
        return None
    if text in PROVINCE_ALIASES:
        return PROVINCE_ALIASES[text]
    return " ".join(part.capitalize() for part in text.split())


def is_religionish(text):
    normalized = normalize_text(text)
    return any(term in normalized for term in RELIGION_TERMS)


def has_any_term(text, terms):
    normalized = normalize_text(text)
    return any(term in normalized for term in terms)


def has_table_marker(text):
    return bool(re.search(r"\b(?:imperial\s+)?table\s+(?:no\s+)?[a-z0-9ivxlcdm\.\-]+\b", text or "", re.I))


def looks_like_narrative_noise(text):
    normalized = normalize_text(text)
    if not normalized:
        return True
    if has_any_term(normalized, NARRATIVE_TERMS):
        return True
    if normalized.startswith("chapter "):
        return True
    if normalized.startswith("contents "):
        return True
    if len(normalized.split()) > 35 and not has_table_marker(text):
        return True
    return False


def candidate_bucket(row):
    topic = row.get("canonical_topic") or ""
    title = row.get("title") or ""
    title_norm = normalize_text(title)
    primary_hint = any(p.search(title) for p in PRIMARY_HINT_PATTERNS)
    secondary_hint = any(p.search(title) for p in SECONDARY_HINT_PATTERNS)
    religionish = is_religionish(title)
    table_marker = has_table_marker(title)
    narrative = looks_like_narrative_noise(title)
    off_target = has_any_term(title, OFF_TARGET_TERMS)

    if topic == "age_sex_civil_condition":
        if narrative or off_target:
            if primary_hint or secondary_hint:
                return "secondary_age_related"
            return None
        if religionish:
            return "religion_primary_like"
        if table_marker or primary_hint:
            return "strong_primary"
        return "secondary_structured"
    if topic in ("age_sex_population", "age_sex_marital_status"):
        if narrative or off_target:
            if secondary_hint:
                return "secondary_age_related"
            return None
        if religionish:
            return "religion_secondary_like"
        if table_marker or secondary_hint:
            return "secondary_structured"
        return None
    if topic in ("age_distribution", "age_by_religion", "age_civil_condition"):
        if off_target:
            return None
        return "secondary_age_related"
    if primary_hint and not religionish and not narrative and not off_target:
        return "strong_primary"
    if secondary_hint and not narrative and not off_target:
        return "secondary_age_related"
    if "civil condition" in title_norm and religionish:
        return "religion_primary_like"
    return None


def candidate_rank(row):
    bucket = row["candidate_bucket"]
    title = row.get("title") or ""
    score = {
        "strong_primary": 400,
        "secondary_structured": 300,
        "secondary_age_related": 200,
        "religion_primary_like": 150,
        "religion_secondary_like": 100,
    }.get(bucket, 0)

    if has_table_marker(title):
        score += 40
    if any(p.search(title) for p in PRIMARY_HINT_PATTERNS):
        score += 25
    if any(p.search(title) for p in SECONDARY_HINT_PATTERNS):
        score += 10
    score -= min(len(normalize_text(title).split()), 80)
    return score


def load_guide_records(guide_path):
    wb = load_workbook(guide_path, read_only=True, data_only=True)
    ws = wb["Full Inventory"]
    rows = ws.iter_rows(values_only=True)
    headers = next(rows)
    header_map = {name: idx for idx, name in enumerate(headers)}

    records = []
    for row in rows:
        if not row or not row[0]:
            continue
        province = canonicalize_province(row[header_map["Province"]])
        year = row[header_map["Year"]]
        try:
            year = int(year)
        except Exception:
            continue
        record = {
            "province": province,
            "year": year,
            "region": row[header_map["Region"]],
            "primary_category": row[header_map["Primary Category"]],
            "sub_category": row[header_map["Sub-Category"]],
            "title": row[header_map["Title"]],
            "source_json": row[header_map["Source JSON"]],
        }
        records.append(record)
    return records


def compute_gap_rows(records, years=None, include_zero_total=False):
    provinces = sorted({r["province"] for r in records if r["province"]})
    years = list(years or MAJOR_CENSUS_YEARS)

    total_counts = defaultdict(int)
    primary_counts = defaultdict(int)
    for row in records:
        key = (row["province"], row["year"])
        total_counts[key] += 1
        if (row["primary_category"] == "AgeSexCivilCondition"
                and row["sub_category"] == "General"):
            primary_counts[key] += 1

    gaps = []
    for province in provinces:
        for year in years:
            key = (province, year)
            total = total_counts.get(key, 0)
            primary = primary_counts.get(key, 0)
            if primary > 0:
                continue
            if total == 0 and not include_zero_total:
                continue
            gap_type = "no_primary_but_other_tables" if total > 0 else "no_tables_in_guide"
            gaps.append({
                "province": province,
                "year": year,
                "guide_total_tables": total,
                "guide_primary_tables": primary,
                "gap_type": gap_type,
            })

    gaps.sort(key=lambda r: (r["year"], r["province"]))
    return gaps


def iter_digitized_candidates(root):
    root = Path(root)
    for table_dir in sorted(root.iterdir()):
        if not table_dir.is_dir():
            continue
        metadata_path = table_dir / "metadata.json"
        if not metadata_path.exists():
            continue
        try:
            payload = json.loads(metadata_path.read_text(encoding="utf-8"))
        except Exception:
            continue
        if payload.get("status") != "ok":
            continue

        digitized = payload.get("digitized_table") or {}
        queue_row = digitized.get("queue_row") or {}
        province = canonicalize_province(queue_row.get("province"))
        year = queue_row.get("year")
        try:
            year = int(year)
        except Exception:
            year = None

        rendered_pages = digitized.get("rendered_pages") or []
        image_paths = [page.get("image_path") for page in rendered_pages if page.get("image_path")]
        if not image_paths:
            image_paths = [str(p) for p in sorted(table_dir.glob("*.png"))]

        row = {
            "province": province,
            "year": year,
            "canonical_topic": queue_row.get("canonical_topic"),
            "relevance_tier": queue_row.get("relevance_tier"),
            "title": queue_row.get("display_identifier") or queue_row.get("title") or table_dir.name,
            "table_dir": str(table_dir),
            "dir_name": table_dir.name,
            "image_count": len(image_paths),
            "sample_image": image_paths[0] if image_paths else "",
        }
        bucket = candidate_bucket(row)
        if bucket:
            row["candidate_bucket"] = bucket
            yield row


def match_candidates_to_gaps(gaps, candidate_rows):
    by_key = defaultdict(list)
    for row in candidate_rows:
        key = (row["province"], row["year"])
        by_key[key].append(row)

    out_rows = []
    for gap in gaps:
        key = (gap["province"], gap["year"])
        matches = sorted(by_key.get(key, []), key=candidate_rank, reverse=True)
        bucket_counts = defaultdict(int)
        page_counts = defaultdict(int)
        titles = []
        sample_paths = []
        for row in matches:
            bucket = row["candidate_bucket"]
            bucket_counts[bucket] += 1
            page_counts[bucket] += row["image_count"]
            if len(titles) < 6:
                titles.append(row["title"])
            if len(sample_paths) < 6:
                sample_paths.append(row["sample_image"] or row["table_dir"])

        out = dict(gap)
        out.update({
            "cluster_candidate_tables": len(matches),
            "cluster_strong_primary_tables": bucket_counts.get("strong_primary", 0),
            "cluster_strong_primary_pages": page_counts.get("strong_primary", 0),
            "cluster_secondary_structured_tables": bucket_counts.get("secondary_structured", 0),
            "cluster_secondary_structured_pages": page_counts.get("secondary_structured", 0),
            "cluster_secondary_age_related_tables": bucket_counts.get("secondary_age_related", 0),
            "cluster_secondary_age_related_pages": page_counts.get("secondary_age_related", 0),
            "cluster_religion_primary_like_tables": bucket_counts.get("religion_primary_like", 0),
            "cluster_religion_secondary_like_tables": bucket_counts.get("religion_secondary_like", 0),
            "sample_titles": titles,
            "sample_paths": sample_paths,
        })
        out_rows.append(out)

    out_rows.sort(
        key=lambda r: (
            -r["cluster_strong_primary_tables"],
            -r["cluster_secondary_structured_tables"],
            -r["cluster_candidate_tables"],
            r["year"],
            r["province"],
        )
    )
    return out_rows


def write_csv(path, rows):
    if not rows:
        Path(path).write_text("", encoding="utf-8")
        return
    fieldnames = [
        "province", "year", "gap_type", "guide_total_tables", "guide_primary_tables",
        "cluster_candidate_tables",
        "cluster_strong_primary_tables", "cluster_strong_primary_pages",
        "cluster_secondary_structured_tables", "cluster_secondary_structured_pages",
        "cluster_secondary_age_related_tables", "cluster_secondary_age_related_pages",
        "cluster_religion_primary_like_tables", "cluster_religion_secondary_like_tables",
        "sample_titles", "sample_paths",
    ]
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            out = dict(row)
            out["sample_titles"] = " | ".join(out["sample_titles"])
            out["sample_paths"] = " | ".join(out["sample_paths"])
            writer.writerow(out)


def main():
    parser = argparse.ArgumentParser(description="Find cluster image candidates for primary-table gaps")
    parser.add_argument("--guide", required=True, help="Path to export/data_guide.xlsx")
    parser.add_argument("--digitized-root", required=True, help="Path to digitized_tables root on cluster")
    parser.add_argument("--years", nargs="+", type=int, default=MAJOR_CENSUS_YEARS,
                        help="Years to evaluate (default: major census years)")
    parser.add_argument("--include-zero-total", action="store_true",
                        help="Also include guide gaps where there are currently no validated tables at all")
    parser.add_argument("--output-json", help="Write detailed JSON report")
    parser.add_argument("--output-csv", help="Write flat CSV report")
    parser.add_argument("--top", type=int, default=40, help="How many rows to print")
    args = parser.parse_args()

    records = load_guide_records(Path(args.guide).expanduser())
    gaps = compute_gap_rows(records, years=args.years, include_zero_total=args.include_zero_total)
    candidates = list(iter_digitized_candidates(Path(args.digitized_root).expanduser()))
    report_rows = match_candidates_to_gaps(gaps, candidates)

    print("Gap analysis")
    print("------------")
    print(f"Guide records: {len(records)}")
    print(f"Gap province-years: {len(gaps)}")
    print(f"Cluster candidate table dirs: {len(candidates)}")

    rows_with_primary = sum(1 for row in report_rows if row["cluster_strong_primary_tables"] > 0)
    rows_with_secondary = sum(
        1 for row in report_rows
        if row["cluster_strong_primary_tables"] == 0
        and row["cluster_secondary_structured_tables"] > 0
    )
    print(f"Gaps with strong primary candidates: {rows_with_primary}")
    print(f"Gaps with only secondary structured candidates: {rows_with_secondary}")

    print("\nTop candidate gaps")
    print("------------------")
    for row in report_rows[:args.top]:
        print(
            f"{row['year']:>4}  {row['province']:<32}  "
            f"{row['gap_type']:<24}  "
            f"guide={row['guide_primary_tables']}p/{row['guide_total_tables']}t  "
            f"cluster strong={row['cluster_strong_primary_tables']}  "
            f"secondary={row['cluster_secondary_structured_tables']}  "
            f"other_age={row['cluster_secondary_age_related_tables']}"
        )
        for title in row["sample_titles"][:3]:
            print(f"      title: {title}")

    if args.output_json:
        payload = {
            "guide_records": len(records),
            "gap_count": len(gaps),
            "cluster_candidate_dir_count": len(candidates),
            "rows": report_rows,
        }
        Path(args.output_json).expanduser().write_text(
            json.dumps(payload, indent=2), encoding="utf-8")
        print(f"\nWrote JSON: {args.output_json}")

    if args.output_csv:
        write_csv(Path(args.output_csv).expanduser(), report_rows)
        print(f"Wrote CSV: {args.output_csv}")


if __name__ == "__main__":
    main()
