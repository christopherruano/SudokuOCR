#!/usr/bin/env python3
"""Audit final_project_exports and export_final for structural quality."""

from __future__ import annotations

import csv
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


SRC = Path("export/final_project_exports")
OUT = Path("export_final")

REQUIRED_SRC = [
    SRC / "README.md",
    SRC / "PROJECT_CODEBOOK.md",
    SRC / "COLUMN_CODEBOOK.md",
    SRC / "core_curated_v2_final" / "source_catalog.csv",
    SRC / "core_curated_v2_final" / "cell_facts_rich.csv",
    SRC / "tail_exploratory_v2" / "source_catalog.csv",
    SRC / "tail_exploratory_v2" / "cell_facts_rich.csv",
    SRC / "assembled_strong_v2" / "source_catalog.csv",
    SRC / "assembled_strong_v2" / "cell_facts_rich.csv",
    SRC / "preview_pack" / "index.csv",
    SRC / "typed_views" / "summary.json",
]

REQUIRED_OUT = [
    OUT / "README.md",
    OUT / "docs" / "README.md",
    OUT / "datasets" / "core_curated_v2_final",
    OUT / "datasets" / "tail_exploratory_v2",
    OUT / "datasets" / "assembled_strong_v2",
    OUT / "browse" / "preview_pack",
    OUT / "browse" / "typed_views",
    OUT / "indexes" / "dataset_entrypoints.csv",
    OUT / "indexes" / "typed_category_summary.csv",
]


def parse_csv(path: Path) -> dict[str, object]:
    csv.field_size_limit(sys.maxsize)
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.reader(handle)
        header = next(reader)
        rows = 0
        for _ in reader:
            rows += 1
    return {"columns": len(header), "rows": rows, "header": header[:12]}


def parse_json(path: Path) -> dict[str, object]:
    obj = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(obj, dict):
        return {"type": "dict", "keys": sorted(obj.keys())[:12]}
    if isinstance(obj, list):
        return {"type": "list", "length": len(obj)}
    return {"type": type(obj).__name__}


def parse_jsonl(path: Path) -> dict[str, object]:
    count = 0
    first_keys: list[str] = []
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            obj = json.loads(line)
            count += 1
            if count == 1 and isinstance(obj, dict):
                first_keys = sorted(obj.keys())[:12]
    return {"rows": count, "first_keys": first_keys}


def parse_xlsx(path: Path) -> dict[str, object]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return {"sheets": wb.sheetnames, "n_sheets": len(wb.sheetnames)}
    finally:
        wb.close()


def audit_preview_pack() -> tuple[dict[str, object], list[str]]:
    issues: list[str] = []
    preview_root = SRC / "preview_pack"
    index_path = preview_root / "index.csv"
    with index_path.open(newline="", encoding="utf-8") as handle:
        index_rows = list(csv.DictReader(handle))

    summary = {
        "index_rows": len(index_rows),
        "page_previews": 0,
        "assembled_previews": 0,
    }

    for row in index_rows:
        folder = preview_root / row["folder"]
        if not folder.exists():
            issues.append(f"Missing preview folder: {folder}")
            continue
        dataset = row["dataset"]
        if dataset in {"core_example", "tail_example"}:
            summary["page_previews"] += 1
            required = ["facts.csv", "matrix.csv", "source_catalog_row.json", "summary.txt", "interpretation.txt"]
            for name in required:
                path = folder / name
                if not path.exists():
                    issues.append(f"Missing page preview file: {path}")
            if (folder / "facts.csv").exists():
                try:
                    facts_info = parse_csv(folder / "facts.csv")
                    if facts_info["rows"] == 0:
                        issues.append(f"Empty facts.csv: {folder}")
                except Exception as exc:  # noqa: BLE001
                    issues.append(f"Bad facts.csv {folder}: {exc}")
            if (folder / "matrix.csv").exists():
                try:
                    matrix_info = parse_csv(folder / "matrix.csv")
                    header = matrix_info["header"]
                    if len(header) < 2 or header[0] != "section" or header[1] != "row_label":
                        issues.append(f"Unexpected matrix header in {folder}: {header[:4]}")
                except Exception as exc:  # noqa: BLE001
                    issues.append(f"Bad matrix.csv {folder}: {exc}")
            if (folder / "source_catalog_row.json").exists():
                try:
                    parse_json(folder / "source_catalog_row.json")
                except Exception as exc:  # noqa: BLE001
                    issues.append(f"Bad source_catalog_row.json {folder}: {exc}")
        elif dataset == "assembled_reviewed":
            summary["assembled_previews"] += 1
            required = ["assembled_group.json", "member_pages.csv", "sections_index.csv", "source_catalog_row.json", "summary.txt"]
            for name in required:
                path = folder / name
                if not path.exists():
                    issues.append(f"Missing assembled preview file: {path}")
            try:
                sections_info = parse_csv(folder / "sections_index.csv")
                if sections_info["rows"] == 0:
                    issues.append(f"Empty sections_index.csv: {folder}")
                else:
                    with (folder / "sections_index.csv").open(newline="", encoding="utf-8") as handle:
                        for section_row in csv.DictReader(handle):
                            csv_name = section_row.get("csv_file", "")
                            if not csv_name:
                                issues.append(f"Blank csv_file in sections_index: {folder}")
                                continue
                            section_path = folder / csv_name
                            if not section_path.exists():
                                issues.append(f"Missing section CSV {section_path}")
                                continue
                            section_info = parse_csv(section_path)
                            if section_info["rows"] == 0:
                                issues.append(f"Empty section CSV {section_path}")
            except Exception as exc:  # noqa: BLE001
                issues.append(f"Bad sections_index.csv {folder}: {exc}")
            for json_name in ["assembled_group.json", "source_catalog_row.json"]:
                path = folder / json_name
                if path.exists():
                    try:
                        parse_json(path)
                    except Exception as exc:  # noqa: BLE001
                        issues.append(f"Bad {json_name} {folder}: {exc}")
        else:
            issues.append(f"Unknown preview dataset type: {dataset}")

    return summary, issues


def audit_typed_views() -> tuple[dict[str, object], list[str]]:
    issues: list[str] = []
    root = SRC / "typed_views"
    summary = json.loads((root / "summary.json").read_text(encoding="utf-8"))
    for category in summary.get("categories", {}):
        cat_dir = root / category
        for rel in ["tables.csv", "README.md", "family_counts.csv", "preview_links", "by_family"]:
            path = cat_dir / rel
            if not path.exists():
                issues.append(f"Missing typed view component: {path}")
        if (cat_dir / "tables.csv").exists():
            try:
                parse_csv(cat_dir / "tables.csv")
            except Exception as exc:  # noqa: BLE001
                issues.append(f"Bad typed tables.csv {cat_dir}: {exc}")
        by_family = cat_dir / "by_family"
        if by_family.exists():
            for path in by_family.iterdir():
                if path.suffix == ".csv":
                    try:
                        parse_csv(path)
                    except Exception as exc:  # noqa: BLE001
                        issues.append(f"Bad by_family CSV {path}: {exc}")
        preview_links = cat_dir / "preview_links"
        if preview_links.exists():
            for link in preview_links.iterdir():
                if not link.is_symlink():
                    issues.append(f"Preview link is not symlink: {link}")
                elif not link.resolve().exists():
                    issues.append(f"Broken preview symlink: {link}")
    return summary, issues


def main() -> None:
    qa_dir = OUT / "qa"
    qa_dir.mkdir(parents=True, exist_ok=True)
    issues: list[str] = []
    warnings: list[str] = []

    for path in REQUIRED_SRC:
        if not path.exists():
            issues.append(f"Missing required source path: {path}")
    for path in REQUIRED_OUT:
        if not path.exists():
            issues.append(f"Missing required export_final path: {path}")

    if list(SRC.rglob("cell_facts.csv")):
        issues.append("Duplicate cell_facts.csv still present inside final_project_exports")
    if list(OUT.rglob("cell_facts.csv")):
        issues.append("Duplicate cell_facts.csv still present inside export_final")

    file_stats: dict[str, dict[str, object]] = {}
    counts = Counter()
    csv.field_size_limit(sys.maxsize)

    for path in sorted(p for p in SRC.rglob("*") if p.is_file() and not p.name.startswith(".")):
        suffix = path.suffix.lower()
        counts[suffix or "<no_suffix>"] += 1
        rel = str(path.relative_to(SRC))
        try:
            if suffix == ".csv":
                file_stats[rel] = parse_csv(path)
            elif suffix == ".json":
                file_stats[rel] = parse_json(path)
            elif suffix == ".jsonl":
                file_stats[rel] = parse_jsonl(path)
            elif suffix == ".xlsx":
                file_stats[rel] = parse_xlsx(path)
        except Exception as exc:  # noqa: BLE001
            issues.append(f"Failed to parse {path}: {exc}")

    preview_summary, preview_issues = audit_preview_pack()
    typed_summary, typed_issues = audit_typed_views()
    issues.extend(preview_issues)
    issues.extend(typed_issues)

    warnings.append(
        "Full raw OCR page JSONs are not bundled in the final release; source_json fields are references, usually to cluster paths."
    )
    warnings.append(
        "Preview matrices are convenience views. For analysis, trust cell_facts_rich.csv and source_catalog.csv first."
    )
    warnings.append(
        "Some assembled section names are OCR-noisy; section CSV content and source_catalog_row.json remain the authoritative shipped artifacts."
    )

    report = {
        "status": "pass" if not issues else "fail",
        "n_issues": len(issues),
        "n_warnings": len(warnings),
        "required_source_paths_checked": len(REQUIRED_SRC),
        "required_export_final_paths_checked": len(REQUIRED_OUT),
        "file_type_counts": dict(counts),
        "preview_summary": preview_summary,
        "typed_view_summary": typed_summary,
        "issues": issues,
        "warnings": warnings,
        "parsed_file_stats_sample": {
            key: file_stats[key] for key in list(file_stats.keys())[:40]
        },
    }

    (qa_dir / "package_qa_report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")

    md_lines = [
        "# Package QA Report",
        "",
        f"Status: `{report['status']}`",
        f"Issues: `{len(issues)}`",
        f"Warnings: `{len(warnings)}`",
        "",
        "## Scope",
        f"- Source package audited: `{SRC}`",
        f"- Final view audited: `{OUT}`",
        "",
        "## Preview Audit",
        f"- index rows: `{preview_summary['index_rows']}`",
        f"- page previews: `{preview_summary['page_previews']}`",
        f"- assembled previews: `{preview_summary['assembled_previews']}`",
        "",
        "## File Counts",
    ]
    for suffix, count in sorted(counts.items()):
        md_lines.append(f"- `{suffix}`: `{count}`")
    md_lines.extend(["", "## Warnings"])
    for warning in warnings:
        md_lines.append(f"- {warning}")
    md_lines.extend(["", "## Issues"])
    if issues:
        for issue in issues:
            md_lines.append(f"- {issue}")
    else:
        md_lines.append("- none")

    (qa_dir / "package_qa_report.md").write_text("\n".join(md_lines) + "\n", encoding="utf-8")
    print(f"Wrote {qa_dir / 'package_qa_report.json'}")
    print(f"Wrote {qa_dir / 'package_qa_report.md'}")
    print(f"Status: {report['status']} ({len(issues)} issues)")


if __name__ == "__main__":
    main()
