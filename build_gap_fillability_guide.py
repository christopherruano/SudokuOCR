"""Build an Excel workbook showing how fillable current primary-table gaps are.

This reads the validated guide workbook plus the gap-candidate report produced by
find_cluster_gap_candidates.py and creates a companion workbook with:

  - Summary: counts by fillability status and difficulty
  - Fillability Matrix: province × year matrix mirroring the guide coverage view
  - Gap Detail: one row per gap with candidate counts, status, and samples

Typical usage on FASRC:
    python3 build_gap_fillability_guide.py \
      --guide ~/SudokuOCR/data_guide.xlsx \
      --gap-json ~/census_toc/gap_candidates.json \
      --output ~/census_toc/gap_fillability_guide.xlsx
"""

import argparse
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


MAJOR_CENSUS_YEARS = [1872, 1881, 1891, 1901, 1911, 1921, 1931, 1941, 1951, 1961]

THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)

HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
SECTION_FONT = Font(bold=True, size=11, color="2F5496")

COVERED_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
LIKELY_FILL = PatternFill(start_color="9FD5B3", end_color="9FD5B3", fill_type="solid")
REVIEW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
SECONDARY_FILL = PatternFill(start_color="F4B183", end_color="F4B183", fill_type="solid")
MANUAL_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
NO_DATA_FILL = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")

STATUS_INFO = {
    "covered": {
        "label": "Covered",
        "difficulty": "None",
        "fill": COVERED_FILL,
        "description": "Guide already has at least one validated primary table.",
    },
    "likely_now": {
        "label": "Likely",
        "difficulty": "Low",
        "fill": LIKELY_FILL,
        "description": "Strong cluster evidence for strict primary-table recovery now.",
    },
    "review_first": {
        "label": "Review",
        "difficulty": "Medium",
        "fill": REVIEW_FILL,
        "description": "Some primary-like candidates exist, but sample titles need manual inspection first.",
    },
    "secondary_only": {
        "label": "Secondary",
        "difficulty": "High",
        "fill": SECONDARY_FILL,
        "description": "Only secondary structured candidates were found on cluster.",
    },
    "manual": {
        "label": "Manual",
        "difficulty": "Very High",
        "fill": MANUAL_FILL,
        "description": "No structured cluster candidate found; manual search or new source work is likely required.",
    },
    "no_data": {
        "label": "No Data",
        "difficulty": "",
        "fill": NO_DATA_FILL,
        "description": "Province/year not represented in the current guide inventory.",
    },
}

STRICT_PRIMARY_PATTERNS = [
    re.compile(r"\bage\s+sex\s+and\s+civil\s+condition\b", re.I),
    re.compile(r"\b(?:imperial\s+)?table\s+(?:vii|vil|vh|7)\b", re.I),
]


def normalize_text(value):
    value = (value or "").strip().lower()
    value = value.replace("&", " and ")
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return " ".join(value.split())


def style_header_row(ws, row, ncols):
    for ci in range(1, ncols + 1):
        cell = ws.cell(row=row, column=ci)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER


def auto_width(ws, min_width=10, max_width=40):
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        max_len = 0
        for cell in col_cells:
            if cell.value is None:
                continue
            value = str(cell.value)
            if "\n" in value:
                value = max(value.split("\n"), key=len)
            max_len = max(max_len, len(value))
        ws.column_dimensions[letter].width = max(min_width, min(max_len + 2, max_width))


def load_guide_coverage(guide_path):
    wb = load_workbook(guide_path, read_only=True, data_only=True)
    ws = wb["Full Inventory"]
    rows = ws.iter_rows(values_only=True)
    headers = next(rows)
    header_map = {name: idx for idx, name in enumerate(headers)}

    total_counts = Counter()
    primary_counts = Counter()
    provinces = set()

    for row in rows:
        if not row or not row[0]:
            continue
        province = row[header_map["Province"]]
        year = row[header_map["Year"]]
        try:
            year = int(year)
        except Exception:
            continue
        if not province:
            continue

        provinces.add(province)
        total_counts[(province, year)] += 1
        if (row[header_map["Primary Category"]] == "AgeSexCivilCondition"
                and row[header_map["Sub-Category"]] == "General"):
            primary_counts[(province, year)] += 1

    return sorted(provinces), total_counts, primary_counts


def coverage_from_records(records):
    total_counts = Counter()
    primary_counts = Counter()
    provinces = set()

    for row in records:
        province = row.get("province")
        year = row.get("year")
        if not province:
            continue
        try:
            year = int(year)
        except Exception:
            continue

        provinces.add(province)
        total_counts[(province, year)] += 1
        if (row.get("primary_category") == "AgeSexCivilCondition"
                and row.get("sub_category") == "General"):
            primary_counts[(province, year)] += 1

    return sorted(provinces), total_counts, primary_counts


def load_gap_rows(gap_json_path):
    with Path(gap_json_path).expanduser().open(encoding="utf-8") as fh:
        payload = json.load(fh)
    return enrich_gap_rows(payload["rows"])


def is_strict_primary_title(title):
    title = title or ""
    normalized = normalize_text(title)
    if not normalized:
        return False
    if "subsidiary table" in normalized:
        return False
    return any(pattern.search(title) for pattern in STRICT_PRIMARY_PATTERNS)


def classify_fillability(row):
    strong = row.get("cluster_strong_primary_tables", 0)
    secondary = row.get("cluster_secondary_structured_tables", 0)
    titles = row.get("sample_titles") or []
    strict_title_count = sum(1 for title in titles if is_strict_primary_title(title))

    if strong == 0 and secondary == 0:
        return "manual"
    if strict_title_count >= 2:
        return "likely_now"
    if strong > 0:
        return "review_first"
    if secondary > 0:
        return "secondary_only"
    return "manual"


def action_for_status(status):
    if status == "likely_now":
        return "Targeted OCR now on strict Table VII-style pages."
    if status == "review_first":
        return "Inspect sample paths first; OCR only explicit primary-table pages."
    if status == "secondary_only":
        return "Use only if secondary structured tables are acceptable substitutes."
    if status == "manual":
        return "Manual search, re-rendering, or new source acquisition needed."
    return "No action needed."


def enrich_gap_rows(gap_rows):
    enriched = []
    for row in gap_rows:
        status = classify_fillability(row)
        info = STATUS_INFO[status]
        out = dict(row)
        out["fillability_status"] = status
        out["fillability_label"] = info["label"]
        out["difficulty"] = info["difficulty"]
        out["recommended_action"] = action_for_status(status)
        enriched.append(out)
    return enriched


def write_summary_sheet(wb, gap_rows):
    ws = wb.active
    ws.title = "Summary"

    counts = Counter(row["fillability_status"] for row in gap_rows)
    likely = [row for row in gap_rows if row["fillability_status"] == "likely_now"]
    review = [row for row in gap_rows if row["fillability_status"] == "review_first"]

    ws.merge_cells("A1:D1")
    ws["A1"].value = "Primary Gap Fillability Guide"
    ws["A1"].font = Font(bold=True, size=16, color="2F5496")

    ws.merge_cells("A2:D2")
    ws["A2"].value = (
        f"Generated from {len(gap_rows)} current primary gaps and the cluster candidate scan."
    )
    ws["A2"].font = Font(size=11, italic=True, color="666666")

    row = 4
    ws.cell(row=row, column=1, value="Overview").font = Font(bold=True, size=13, color="2F5496")
    row += 1

    overview = [
        ("Total primary gaps", len(gap_rows)),
        ("Likely fillable now", counts["likely_now"]),
        ("Needs review first", counts["review_first"]),
        ("Secondary-only fallback", counts["secondary_only"]),
        ("Likely manual/new source", counts["manual"]),
    ]
    for label, value in overview:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Legend").font = Font(bold=True, size=13, color="2F5496")
    row += 1
    for key in ("covered", "likely_now", "review_first", "secondary_only", "manual"):
        info = STATUS_INFO[key]
        cell = ws.cell(row=row, column=1, value=info["label"])
        cell.fill = info["fill"]
        cell.border = THIN_BORDER
        ws.cell(row=row, column=2, value=info["difficulty"]).font = Font(bold=True)
        ws.cell(row=row, column=3, value=info["description"])
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Immediate Targets").font = Font(bold=True, size=13, color="2F5496")
    row += 1
    ws.cell(row=row, column=1, value="Province").font = SECTION_FONT
    ws.cell(row=row, column=2, value="Year").font = SECTION_FONT
    ws.cell(row=row, column=3, value="Strong").font = SECTION_FONT
    ws.cell(row=row, column=4, value="Secondary").font = SECTION_FONT
    row += 1
    for gap in sorted(likely, key=lambda r: (r["year"], r["province"])):
        ws.cell(row=row, column=1, value=gap["province"])
        ws.cell(row=row, column=2, value=gap["year"])
        ws.cell(row=row, column=3, value=gap["cluster_strong_primary_tables"])
        ws.cell(row=row, column=4, value=gap["cluster_secondary_structured_tables"])
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Review-First Targets").font = Font(bold=True, size=13, color="2F5496")
    row += 1
    ws.cell(row=row, column=1, value="Province").font = SECTION_FONT
    ws.cell(row=row, column=2, value="Year").font = SECTION_FONT
    ws.cell(row=row, column=3, value="Strong").font = SECTION_FONT
    ws.cell(row=row, column=4, value="Secondary").font = SECTION_FONT
    row += 1
    for gap in sorted(review, key=lambda r: (r["year"], r["province"])):
        ws.cell(row=row, column=1, value=gap["province"])
        ws.cell(row=row, column=2, value=gap["year"])
        ws.cell(row=row, column=3, value=gap["cluster_strong_primary_tables"])
        ws.cell(row=row, column=4, value=gap["cluster_secondary_structured_tables"])
        row += 1

    auto_width(ws, max_width=50)


def append_summary_section(ws, gap_rows, start_row=None):
    counts = Counter(row["fillability_status"] for row in gap_rows)

    row = start_row or (ws.max_row + 2)
    ws.cell(row=row, column=1, value="Gap Fillability").font = Font(bold=True, size=13, color="2F5496")
    row += 1

    overview = [
        ("Current primary gaps", len(gap_rows)),
        ("Likely fillable now", counts["likely_now"]),
        ("Needs review first", counts["review_first"]),
        ("Secondary-only fallback", counts["secondary_only"]),
        ("Likely manual/new source", counts["manual"]),
    ]
    for label, value in overview:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    legend_rows = [
        ("Likely", LIKELY_FILL, "Good immediate recovery target from cluster."),
        ("Review", REVIEW_FILL, "Inspect sample paths before OCR."),
        ("Secondary", SECONDARY_FILL, "Only secondary structured candidates found."),
        ("Manual", MANUAL_FILL, "No structured cluster candidate found."),
    ]
    for label, fill, description in legend_rows:
        cell = ws.cell(row=row, column=1, value=label)
        cell.fill = fill
        cell.border = THIN_BORDER
        ws.cell(row=row, column=2, value=description)
        row += 1

    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 35)
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width or 0, 18)
    ws.column_dimensions["C"].width = max(ws.column_dimensions["C"].width or 0, 70)

    return row


def write_matrix_sheet(wb, provinces, total_counts, primary_counts, gap_rows, title="Fillability Matrix"):
    ws = wb.create_sheet(title=title)
    gap_map = {(row["province"], row["year"]): row for row in gap_rows}

    ws.cell(row=1, column=1, value="Province / Year")
    for ci, year in enumerate(MAJOR_CENSUS_YEARS, 2):
        ws.cell(row=1, column=ci, value=year)
    style_header_row(ws, 1, len(MAJOR_CENSUS_YEARS) + 1)

    for ri, province in enumerate(provinces, 2):
        ws.cell(row=ri, column=1, value=province).font = Font(bold=True)
        ws.cell(row=ri, column=1).border = THIN_BORDER

        for ci, year in enumerate(MAJOR_CENSUS_YEARS, 2):
            cell = ws.cell(row=ri, column=ci)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            key = (province, year)
            total = total_counts.get(key, 0)
            primary = primary_counts.get(key, 0)

            if primary > 0:
                cell.value = "Covered"
                cell.fill = STATUS_INFO["covered"]["fill"]
                cell.comment = Comment(
                    f"Guide coverage exists.\nPrimary tables: {primary}\nTotal tables: {total}",
                    "SudokuOCR",
                )
                continue

            gap_row = gap_map.get(key)
            if not gap_row:
                cell.value = "No Data"
                cell.fill = STATUS_INFO["no_data"]["fill"]
                continue

            status = gap_row["fillability_status"]
            strong = gap_row["cluster_strong_primary_tables"]
            secondary = gap_row["cluster_secondary_structured_tables"]
            label = STATUS_INFO[status]["label"]
            if status == "manual":
                cell.value = label
            elif status == "secondary_only":
                cell.value = f"{label}\n0/{secondary}"
            else:
                cell.value = f"{label}\n{strong}/{secondary}"
            cell.fill = STATUS_INFO[status]["fill"]

            comment_lines = [
                f"{province} {year}",
                f"Guide gap: {gap_row['guide_primary_tables']}p/{gap_row['guide_total_tables']}t",
                f"Status: {label}",
                f"Difficulty: {gap_row['difficulty']}",
                f"Strong primary candidates: {strong}",
                f"Secondary structured candidates: {secondary}",
                f"Other age-related candidates: {gap_row['cluster_secondary_age_related_tables']}",
                f"Action: {gap_row['recommended_action']}",
            ]
            if gap_row.get("sample_titles"):
                comment_lines.append("Sample titles:")
                for title in gap_row["sample_titles"][:4]:
                    comment_lines.append(f"- {title}")
            cell.comment = Comment("\n".join(comment_lines), "SudokuOCR")

    legend_row = len(provinces) + 3
    ws.cell(row=legend_row, column=1, value="Legend").font = Font(bold=True)
    legend_items = [
        ("Covered", COVERED_FILL, "Guide already has primary coverage."),
        ("Likely", LIKELY_FILL, "Good immediate recovery target from cluster."),
        ("Review", REVIEW_FILL, "Inspect sample paths before OCR."),
        ("Secondary", SECONDARY_FILL, "Only secondary structured candidates found."),
        ("Manual", MANUAL_FILL, "No structured cluster candidate found."),
    ]
    for idx, (label, fill, description) in enumerate(legend_items, 1):
        cell = ws.cell(row=legend_row + idx, column=1, value=label)
        cell.fill = fill
        cell.border = THIN_BORDER
        ws.cell(row=legend_row + idx, column=2, value=description)

    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 30
    for ci in range(2, len(MAJOR_CENSUS_YEARS) + 2):
        ws.column_dimensions[get_column_letter(ci)].width = 15


def write_detail_sheet(wb, gap_rows, title="Gap Detail"):
    ws = wb.create_sheet(title=title)

    headers = [
        "Province",
        "Year",
        "Gap Type",
        "Status",
        "Difficulty",
        "Guide",
        "Strong Primary",
        "Secondary Structured",
        "Other Age-Related",
        "Recommended Action",
        "Sample Titles",
        "Sample Paths",
    ]
    for ci, header in enumerate(headers, 1):
        ws.cell(row=1, column=ci, value=header)
    style_header_row(ws, 1, len(headers))

    for ri, row in enumerate(
            sorted(gap_rows, key=lambda r: (r["year"], r["province"])), 2):
        values = [
            row["province"],
            row["year"],
            row["gap_type"],
            row["fillability_label"],
            row["difficulty"],
            f"{row['guide_primary_tables']}p/{row['guide_total_tables']}t",
            row["cluster_strong_primary_tables"],
            row["cluster_secondary_structured_tables"],
            row["cluster_secondary_age_related_tables"],
            row["recommended_action"],
            "\n".join(row.get("sample_titles") or []),
            "\n".join(row.get("sample_paths") or []),
        ]
        for ci, value in enumerate(values, 1):
            cell = ws.cell(row=ri, column=ci, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=ci in (10, 11, 12))

        status_fill = STATUS_INFO[row["fillability_status"]]["fill"]
        for ci in range(1, len(headers) + 1):
            ws.cell(row=ri, column=ci).fill = status_fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(gap_rows) + 1}"
    auto_width(ws, max_width=70)
    ws.column_dimensions["K"].width = 50
    ws.column_dimensions["L"].width = 80


def main():
    parser = argparse.ArgumentParser(description="Build an Excel workbook for gap fillability.")
    parser.add_argument("--guide", required=True, help="Path to data_guide.xlsx")
    parser.add_argument("--gap-json", required=True, help="Path to gap_candidates.json")
    parser.add_argument("--output", required=True, help="Path to output workbook (.xlsx)")
    args = parser.parse_args()

    guide_path = Path(args.guide).expanduser()
    gap_json_path = Path(args.gap_json).expanduser()
    output_path = Path(args.output).expanduser()

    if not guide_path.exists():
        raise SystemExit(f"Guide workbook not found: {guide_path}")
    if not gap_json_path.exists():
        raise SystemExit(f"Gap JSON not found: {gap_json_path}")

    gap_rows = load_gap_rows(gap_json_path)

    provinces, total_counts, primary_counts = load_guide_coverage(guide_path)

    wb = Workbook()
    write_summary_sheet(wb, gap_rows)
    write_matrix_sheet(wb, provinces, total_counts, primary_counts, gap_rows)
    write_detail_sheet(wb, gap_rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    counts = Counter(row["fillability_status"] for row in gap_rows)
    print(f"Wrote workbook: {output_path}")
    print(f"Total gaps: {len(gap_rows)}")
    print(f"  Likely now: {counts['likely_now']}")
    print(f"  Review first: {counts['review_first']}")
    print(f"  Secondary only: {counts['secondary_only']}")
    print(f"  Manual: {counts['manual']}")


if __name__ == "__main__":
    main()
