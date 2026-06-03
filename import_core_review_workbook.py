"""Import Excel review decisions back into core_review_decisions.csv."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_WORKBOOK = "export/local_research_bundle/final_core_prep/core_review_workbook.xlsx"
DEFAULT_DECISIONS_CSV = "export/local_research_bundle/final_core_prep/core_review_decisions.csv"
VALID_DECISIONS = {"promote", "reject", "hold", ""}


def normalize_decision(value) -> str:
    text = str(value or "").strip().lower()
    if text in {"accept", "accepted", "promote_now"}:
        return "promote"
    if text in {"reject_now", "rejected"}:
        return "reject"
    if text in {"review", "manual_review", "pending", "skip"}:
        return "hold"
    if text in VALID_DECISIONS:
        return text
    return ""


def read_csv_rows(path: Path):
    with path.open(encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def write_csv(path: Path, rows, fieldnames):
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser(description="Import review decisions from the Excel workbook.")
    parser.add_argument("--workbook", default=DEFAULT_WORKBOOK)
    parser.add_argument("--decisions-csv", default=DEFAULT_DECISIONS_CSV)
    args = parser.parse_args()

    workbook_path = Path(args.workbook).expanduser()
    decisions_csv = Path(args.decisions_csv).expanduser()

    rows = read_csv_rows(decisions_csv)
    fieldnames = list(rows[0].keys()) if rows else []
    by_result = {row.get("result_json", ""): row for row in rows}

    wb = load_workbook(workbook_path)
    ws = wb["Review"]

    header_map = {cell.value: idx for idx, cell in enumerate(ws[1], start=1)}
    required = ["review_decision", "review_notes", "remote_result_json"]
    missing = [name for name in required if name not in header_map]
    if missing:
        raise SystemExit(f"Workbook is missing required columns: {missing}")

    updated = 0
    invalid = 0
    for row_idx in range(2, ws.max_row + 1):
        result_json = str(ws.cell(row=row_idx, column=header_map["remote_result_json"]).value or "").strip()
        if not result_json:
            continue
        target = by_result.get(result_json)
        if not target:
            continue

        review_decision = normalize_decision(ws.cell(row=row_idx, column=header_map["review_decision"]).value)
        review_notes = str(ws.cell(row=row_idx, column=header_map["review_notes"]).value or "").strip()
        if review_decision == "" and str(ws.cell(row=row_idx, column=header_map["review_decision"]).value or "").strip():
            invalid += 1

        target["review_decision"] = review_decision
        target["review_notes"] = review_notes
        target["final_decision"] = review_decision if review_decision in {"promote", "reject", "hold"} else ""
        updated += 1

    write_csv(decisions_csv, rows, fieldnames)
    print("Imported workbook decisions")
    print("---------------------------")
    print(f"Rows touched:       {updated}")
    print(f"Invalid decisions:  {invalid}")
    print(f"Updated CSV:        {decisions_csv}")


if __name__ == "__main__":
    main()
