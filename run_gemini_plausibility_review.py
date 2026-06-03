#!/usr/bin/env python3
"""Run Gemini plausibility review over validated table packets.

This script is designed for the final-validation pass:

- streams JSONL packets instead of loading the fact CSV;
- sends each table as metadata plus a compact matrix/table preview;
- supports multiple Gemini keys from `.env`;
- writes append-only JSONL results so interrupted runs can resume;
- emits a CSV for the review app.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from openpyxl import load_workbook


load_dotenv()

ROOT = Path("export_final/final_validated")
SENDOUT = Path("export_final/sendout_final")
DEFAULT_INPUT = ROOT / "llm_plausibility_review" / "validated_table_packets.jsonl"
DEFAULT_OUTPUT_DIR = ROOT / "llm_plausibility_review" / "gemini_run"
DEFAULT_MODEL = os.environ.get("GEMINI_MODEL", "gemini-3.1-pro-preview")

RESULT_FIELDS = [
    "table_id",
    "status",
    "llm_decision",
    "confidence",
    "flags",
    "reason",
    "suggested_next_action",
    "review_questions",
    "model",
    "provider",
    "key_slot",
    "priority_score",
    "priority_reasons",
    "source_image_local",
    "excel_path",
    "province",
    "year",
    "table_family",
    "measure_category",
    "format_class",
    "title",
    "elapsed_seconds",
    "error",
    "raw_response_preview",
]


@dataclass
class KeyState:
    api_key: str
    provider: str
    base_url: str = ""
    slot: int | str = 0
    lock: threading.Lock = field(default_factory=threading.Lock)
    last_call: float = 0.0


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    rows = []
    if not path.exists():
        return rows
    with path.open(encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            try:
                rows.append(json.loads(line))
            except json.JSONDecodeError:
                continue
    return rows


def iter_jsonl(path: Path):
    with path.open(encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if line:
                yield json.loads(line)


def compact_text(value: Any, limit: int = 180) -> str:
    text = str(value or "").replace("\n", " ").strip()
    return text[:limit]


def markdown_escape(value: Any) -> str:
    text = "" if value is None else str(value)
    return text.replace("\n", " ").replace("|", "\\|").strip()


def resolve_excel_path(excel_path: str) -> Path | None:
    if not excel_path:
        return None
    candidates = []
    p = Path(excel_path)
    if p.is_absolute():
        candidates.append(p)
    else:
        candidates.append(SENDOUT / p)
        candidates.append(Path(excel_path))
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def excel_matrix_markdown(excel_path: str, max_rows: int, max_cols: int, max_chars: int) -> tuple[str, str]:
    path = resolve_excel_path(excel_path)
    if not path:
        return "", "excel_not_found"
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        sheet_name = "Matrix" if "Matrix" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]
        total_rows, total_cols = ws.max_row, ws.max_column
        rows = []
        for row in ws.iter_rows(max_row=min(total_rows, max_rows), max_col=min(total_cols, max_cols), values_only=True):
            rows.append(["" if cell is None else cell for cell in row])
        wb.close()
    except Exception as exc:
        return "", f"excel_read_error:{type(exc).__name__}:{compact_text(exc, 120)}"
    if not rows:
        return "", "excel_empty"
    header = [markdown_escape(cell) for cell in rows[0]]
    body = [[markdown_escape(cell) for cell in row] for row in rows[1:]]
    lines = [
        f"Matrix sheet preview from `{path}`",
        f"Original sheet size: {total_rows} rows x {total_cols} columns. Preview cap: {max_rows} rows x {max_cols} columns.",
        "",
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * len(header)) + " |",
    ]
    for row in body:
        if len(row) < len(header):
            row = row + [""] * (len(header) - len(row))
        lines.append("| " + " | ".join(row[: len(header)]) + " |")
        if sum(len(line) for line in lines) > max_chars:
            lines.append("")
            lines.append("[table preview truncated for prompt size]")
            break
    return "\n".join(lines)[:max_chars], "excel_matrix"


def facts_matrix_markdown(packet: dict[str, Any], max_chars: int) -> str:
    facts = packet.get("summary", {}).get("sample_facts", [])
    if not facts:
        return ""
    row_labels: list[str] = []
    col_labels: list[str] = []
    cells: dict[tuple[str, str], str] = {}
    for fact in facts:
        row_label = compact_text(
            fact.get("row_label_std")
            or fact.get("row_label_raw")
            or fact.get("row_group_value")
            or fact.get("resolved_age_label_std")
            or "row",
            80,
        )
        parts = [
            fact.get("section", ""),
            fact.get("column_group_std") or fact.get("column_group_raw", ""),
            fact.get("sub_column_std") or fact.get("sub_column_raw", ""),
            fact.get("sex", ""),
            fact.get("civil_condition", ""),
            fact.get("religion", ""),
        ]
        col_label = compact_text(" | ".join(str(p) for p in parts if p), 120) or "value"
        if row_label not in row_labels:
            row_labels.append(row_label)
        if col_label not in col_labels:
            col_labels.append(col_label)
        cells[(row_label, col_label)] = str(fact.get("value", ""))
    col_labels = col_labels[:18]
    lines = [
        "Fallback matrix reconstructed from sampled facts.",
        "",
        "| row | " + " | ".join(markdown_escape(c) for c in col_labels) + " |",
        "| --- | " + " | ".join(["---"] * len(col_labels)) + " |",
    ]
    for row_label in row_labels[:35]:
        values = [markdown_escape(cells.get((row_label, c), "")) for c in col_labels]
        lines.append("| " + markdown_escape(row_label) + " | " + " | ".join(values) + " |")
        if sum(len(line) for line in lines) > max_chars:
            lines.append("")
            lines.append("[sampled-facts matrix truncated for prompt size]")
            break
    return "\n".join(lines)[:max_chars]


def build_prompt(packet: dict[str, Any], args: argparse.Namespace) -> str:
    meta = packet.get("metadata", {})
    audit = packet.get("deterministic_audit", {})
    summary = packet.get("summary", {})
    format_guidance = packet.get("format_review_guidance", {})
    matrix, matrix_source = excel_matrix_markdown(
        meta.get("excel_path", ""),
        max_rows=args.max_matrix_rows,
        max_cols=args.max_matrix_cols,
        max_chars=args.max_table_chars,
    )
    if not matrix:
        matrix = facts_matrix_markdown(packet, args.max_table_chars)
    prompt = {
        "task": "Review this extracted historical census table for semantic/numeric plausibility before final export.",
        "decision_schema": {
            "table_id": packet.get("table_id"),
            "decision": "pass | needs_image_compare | needs_manual_review | reject_or_rerun",
            "confidence": "high | medium | low",
            "flags": ["short machine-readable flags"],
            "reason": "short explanation. If decision is not pass, explicitly explain why.",
            "suggested_next_action": "specific action for a human reviewer or pipeline rerun",
            "review_questions": ["specific questions to answer in the review app"],
        },
        "rules": [
            "Do not invent corrected numbers.",
            "A deterministic arithmetic pass does not guarantee semantic correctness.",
            "Preserve units: counts, rates, percentages, per-1000, and distribution-of-1000 are different.",
            "Use format_review_guidance to choose the correct plausibility rules before judging the table.",
            "Distribution-of-1000 and distribution-of-10000 tables are normalized distributions; their meaningful check is usually within-base sums/scale, not raw population additivity.",
            "For bounded distribution-of-1000/per-mille cells, values above 1000 are scale errors unless the source clearly says this is a ratio/rate that can exceed 1000; a 10000-like value should be flagged for image comparison or rerun.",
            "For bounded distribution-of-10000 cells, values above 10000 are scale errors unless the source explicitly switches to raw counts.",
            "For percentages, rates, mortality rates, averages, and ratios, decimal points are part of the data. Flag likely dropped decimal points or integer-rounded rate tables when the title/headers imply decimal values.",
            "Rate, ratio, percentage, sex-ratio, index, mortality-rate, and birth-rate tables are not additive count tables unless the source explicitly provides counts.",
            "Flag impossible or implausible age/civil-condition patterns, such as large widowed/married child counts.",
            "The normalized sex field is opportunistic metadata extracted from labels like Male, Female, Persons, Total, or All. Do not assume every table should have Male/Female/All categories.",
            "Only apply Male + Female = Persons/All checks when the printed title, row labels, column labels, or section labels explicitly define a sex axis with comparable count cells.",
            "Do not flag a table merely because sex_values are blank, only All, or not relevant to the printed table.",
            "For mortality/birth/death tables, be stricter about unit interpretation and zero-heavy rows.",
            "If the table is unusual, infer appropriate plausibility constraints from title, labels, and units before deciding.",
            "Return valid JSON only.",
        ],
        "metadata": {
            "table_id": packet.get("table_id"),
            "province": meta.get("province", ""),
            "year": meta.get("year", ""),
            "title": meta.get("title", ""),
            "table_family": meta.get("table_family", ""),
            "format_class": meta.get("format_class", ""),
            "measure_category": meta.get("measure_category", ""),
            "constraint_regime": meta.get("constraint_regime", ""),
            "denominator_hint": meta.get("denominator_hint", ""),
            "source_image_local": meta.get("source_image_local", ""),
            "excel_path": meta.get("excel_path", ""),
        },
        "format_review_guidance": format_guidance,
        "deterministic_audit": audit,
        "numeric_summary": {
            key: summary.get(key)
            for key in [
                "fact_rows",
                "numeric_rows",
                "nonzero_rows",
                "zero_rows",
                "min_value",
                "max_value",
                "mean_value",
                "sex_values",
                "civil_condition_values",
                "religion_values",
                "age_labels",
                "measure_units",
                "heuristic_flags",
                "suspicious_examples",
            ]
        },
        "recommended_checks": packet.get("recommended_checks", []),
        "matrix_source": matrix_source,
    }
    return (
        "You are a rigorous data QA reviewer for extracted census tables.\n"
        "Review the metadata and table preview, then return one JSON object matching the schema.\n\n"
        f"{json.dumps(prompt, ensure_ascii=False, indent=2)}\n\n"
        "TABLE PREVIEW:\n"
        f"{matrix}\n"
    )


def parse_json_response(text: str) -> dict[str, Any]:
    raw = (text or "").strip()
    if not raw:
        raise ValueError("empty response")
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?\s*", "", raw)
        raw = re.sub(r"\s*```$", "", raw)
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", raw, re.S)
        if match:
            parsed = json.loads(match.group(0))
        else:
            raise
    return normalize_model_json(parsed)


def normalize_model_json(parsed: Any) -> dict[str, Any]:
    if isinstance(parsed, list):
        dict_items = [item for item in parsed if isinstance(item, dict)]
        if dict_items:
            parsed = dict_items[0]
        else:
            raise ValueError("model returned JSON list without object entries")
    if not isinstance(parsed, dict):
        raise ValueError(f"model returned {type(parsed).__name__}, expected object")
    if "decision" not in parsed:
        for key in ["review", "result", "assessment", "decision_schema"]:
            nested = parsed.get(key)
            if isinstance(nested, dict) and "decision" in nested:
                parsed = {**parsed, **nested}
                break
    if "decision" not in parsed and "plausibility_decision" in parsed:
        parsed["decision"] = parsed.get("plausibility_decision")
    if parsed.get("decision") not in {"pass", "needs_image_compare", "needs_manual_review", "reject_or_rerun"}:
        parsed = {
            **parsed,
            "decision": "needs_manual_review",
            "confidence": parsed.get("confidence") or "low",
            "flags": (parsed.get("flags") if isinstance(parsed.get("flags"), list) else [str(parsed.get("flags"))] if parsed.get("flags") else [])
            + ["malformed_or_missing_llm_decision"],
            "reason": parsed.get("reason") or "Model response did not contain a valid decision; review manually.",
            "suggested_next_action": parsed.get("suggested_next_action") or "Review this table in the LLM plausibility review app.",
        }
    return parsed


def load_key_states(exclude_slots: set[int] | None = None, include_google_direct: bool = False) -> list[KeyState]:
    load_dotenv()
    exclude_slots = exclude_slots or set()
    states: list[KeyState] = []
    base_url = os.environ.get("HARVARD_GEMINI_BASE_URL") or os.environ.get("GEMINI_BASE_URL") or ""
    multi = os.environ.get("HARVARD_GEMINI_API_KEYS", "")
    if multi.strip():
        for slot, key in enumerate([part.strip() for part in multi.split(",") if part.strip()]):
            if slot not in exclude_slots:
                states.append(KeyState(api_key=key, provider="harvard_gemini", base_url=base_url, slot=slot))
    elif os.environ.get("HARVARD_GEMINI_API_KEY"):
        if 0 not in exclude_slots:
            states.append(KeyState(api_key=os.environ["HARVARD_GEMINI_API_KEY"], provider="harvard_gemini", base_url=base_url, slot=0))
    elif os.environ.get("GEMINI_API_KEY"):
        if 0 not in exclude_slots:
            states.append(KeyState(api_key=os.environ["GEMINI_API_KEY"], provider="google_gemini", base_url=os.environ.get("GEMINI_BASE_URL", ""), slot=0))
    if include_google_direct and os.environ.get("GEMINI_API_KEY") and states:
        states.append(KeyState(api_key=os.environ["GEMINI_API_KEY"], provider="google_gemini", base_url=os.environ.get("GEMINI_BASE_URL", ""), slot="google"))
    if not states:
        raise RuntimeError("No Gemini key found. Expected HARVARD_GEMINI_API_KEYS, HARVARD_GEMINI_API_KEY, or GEMINI_API_KEY in .env/environment.")
    return states


def call_gemini(prompt: str, key_state: KeyState, model: str, min_interval: float, max_attempts: int) -> tuple[str, str]:
    from google import genai
    from google.genai import types

    timeout_ms = 300_000 if key_state.base_url else 120_000
    http_options: dict[str, Any] = {"timeout": timeout_ms}
    if key_state.base_url:
        http_options["base_url"] = key_state.base_url
    client = genai.Client(api_key=key_state.api_key, http_options=types.HttpOptions(**http_options))
    config = types.GenerateContentConfig(
        temperature=0.0,
        response_mime_type="application/json",
    )
    last_error = ""
    for attempt in range(max_attempts):
        with key_state.lock:
            if min_interval > 0 and key_state.last_call:
                wait = min_interval - (time.time() - key_state.last_call)
                if wait > 0:
                    time.sleep(wait)
            key_state.last_call = time.time()
        try:
            response = client.models.generate_content(model=model, contents=prompt, config=config)
            return response.text or "", key_state.provider
        except Exception as exc:
            last_error = f"{type(exc).__name__}: {exc}"
            transient = any(token in last_error for token in ["429", "500", "502", "503", "504", "UNAVAILABLE", "RESOURCE_EXHAUSTED", "INTERNAL"])
            if not transient or attempt == max_attempts - 1:
                raise
            time.sleep(min(90, 5 * (2**attempt)))
    raise RuntimeError(last_error or "Gemini call failed")


def result_row(
    packet: dict[str, Any],
    parsed: dict[str, Any],
    status: str,
    model: str,
    provider: str,
    elapsed: float,
    error: str = "",
    key_slot: int | str = "",
    raw_response: str = "",
) -> dict[str, Any]:
    meta = packet.get("metadata", {})
    priority = packet.get("review_priority", {})
    flags = parsed.get("flags", [])
    questions = parsed.get("review_questions", [])
    return {
        "table_id": packet.get("table_id", ""),
        "status": status,
        "llm_decision": parsed.get("decision", ""),
        "confidence": parsed.get("confidence", ""),
        "flags": ";".join(flags) if isinstance(flags, list) else str(flags or ""),
        "reason": parsed.get("reason", ""),
        "suggested_next_action": parsed.get("suggested_next_action", ""),
        "review_questions": " | ".join(questions) if isinstance(questions, list) else str(questions or ""),
        "model": model,
        "provider": provider,
        "key_slot": str(key_slot),
        "priority_score": str(priority.get("score", "")),
        "priority_reasons": ";".join(priority.get("reasons", [])),
        "source_image_local": meta.get("source_image_local", ""),
        "excel_path": meta.get("excel_path", ""),
        "province": meta.get("province", ""),
        "year": meta.get("year", ""),
        "table_family": meta.get("table_family", ""),
        "measure_category": meta.get("measure_category", ""),
        "format_class": meta.get("format_class", ""),
        "title": meta.get("title", ""),
        "elapsed_seconds": f"{elapsed:.2f}",
        "error": error,
        "raw_response_preview": raw_response[:2000],
    }


def load_done_ids(results_path: Path, rerun_errors: bool) -> set[str]:
    done = set()
    for row in load_jsonl(results_path):
        table_id = row.get("table_id", "")
        if not table_id:
            continue
        if row.get("status") == "ok" or not rerun_errors:
            done.add(table_id)
    return done


def append_jsonl(path: Path, row: dict[str, Any], lock: threading.Lock) -> None:
    with lock:
        with path.open("a", encoding="utf-8") as fh:
            fh.write(json.dumps(row, ensure_ascii=False) + "\n")


def write_results_csv(results_jsonl: Path, csv_path: Path) -> int:
    latest: dict[str, dict[str, Any]] = {}
    for row in load_jsonl(results_jsonl):
        table_id = row.get("table_id", "")
        if table_id:
            latest[table_id] = row
    rows = sorted(latest.values(), key=lambda r: (r.get("llm_decision", ""), r.get("province", ""), r.get("year", ""), r.get("table_id", "")))
    with csv_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=RESULT_FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


def review_one(task: tuple[int, dict[str, Any]], args: argparse.Namespace, key_states: list[KeyState]) -> dict[str, Any]:
    idx, packet = task
    key_state = key_states[idx % len(key_states)]
    started = time.time()
    try:
        prompt = build_prompt(packet, args)
        if args.dry_run:
            return result_row(
                packet,
                {
                    "decision": "dry_run",
                    "confidence": "high",
                    "flags": [],
                    "reason": prompt[: args.dry_run_chars],
                    "suggested_next_action": "Run without --dry-run to call Gemini.",
                    "review_questions": [],
                },
                status="dry_run",
                model=args.model,
                provider=key_state.provider,
                elapsed=time.time() - started,
                key_slot=key_state.slot,
            )
        text, provider = call_gemini(
            prompt=prompt,
            key_state=key_state,
            model=args.model,
            min_interval=args.min_interval,
            max_attempts=args.max_attempts,
        )
        parsed = parse_json_response(text)
        return result_row(packet, parsed, "ok", args.model, provider, time.time() - started, key_slot=key_state.slot, raw_response=text)
    except Exception as exc:
        return result_row(
            packet,
            {
                "decision": "api_error",
                "confidence": "low",
                "flags": ["api_error"],
                "reason": str(exc),
                "suggested_next_action": "Retry this table or review manually.",
                "review_questions": [],
            },
            status="error",
            model=args.model,
            provider=key_state.provider,
            elapsed=time.time() - started,
            error=f"{type(exc).__name__}: {exc}",
            key_slot=key_state.slot,
        )


def parse_slot_set(value: str) -> set[int]:
    slots = set()
    for part in str(value or "").split(","):
        part = part.strip()
        if part:
            slots.add(int(part))
    return slots


def check_keys(key_states: list[KeyState], model: str, min_interval: float, max_attempts: int) -> None:
    prompt = (
        "Return only this JSON object: "
        '{"decision":"pass","confidence":"high","flags":[],"reason":"key check","suggested_next_action":"none","review_questions":[]}'
    )
    print("Checking Gemini key slots:", flush=True)
    for state in key_states:
        try:
            text, provider = call_gemini(prompt, state, model, min_interval, max_attempts)
            parsed = parse_json_response(text)
            print(json.dumps({"slot": state.slot, "provider": provider, "ok": True, "decision": parsed.get("decision")}), flush=True)
        except Exception as exc:
            print(json.dumps({"slot": state.slot, "provider": state.provider, "ok": False, "error": f"{type(exc).__name__}: {exc}"}), flush=True)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-jsonl", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--model", default=DEFAULT_MODEL)
    parser.add_argument("--max-workers", type=int, default=8)
    parser.add_argument("--min-interval", type=float, default=float(os.environ.get("GEMINI_MIN_INTERVAL", "0")))
    parser.add_argument("--max-attempts", type=int, default=4)
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--priority-min", type=int, default=-1, help="Only review packets with priority score >= this value.")
    parser.add_argument("--rerun-errors", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--dry-run-chars", type=int, default=4000)
    parser.add_argument("--max-matrix-rows", type=int, default=45)
    parser.add_argument("--max-matrix-cols", type=int, default=28)
    parser.add_argument("--max-table-chars", type=int, default=14000)
    parser.add_argument("--exclude-key-slots", default="", help="Comma-separated original key slots to skip, e.g. 0,3.")
    parser.add_argument("--include-google-direct", action="store_true", help="Also use GEMINI_API_KEY as a direct Google slot even when Harvard keys are configured.")
    parser.add_argument("--check-keys", action="store_true", help="Call Gemini once per configured key slot, print health, then exit.")
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    results_jsonl = args.output_dir / "gemini_plausibility_results.jsonl"
    results_csv = args.output_dir / "gemini_plausibility_results.csv"
    done_ids = load_done_ids(results_jsonl, args.rerun_errors)
    key_states = load_key_states(parse_slot_set(args.exclude_key_slots), include_google_direct=args.include_google_direct)
    if args.check_keys:
        check_keys(key_states, args.model, args.min_interval, args.max_attempts)
        return
    print(f"Gemini model: {args.model}", flush=True)
    print(f"Key slots: {[k.slot for k in key_states]} ({', '.join(sorted({k.provider for k in key_states}))})", flush=True)
    print(f"Workers: {args.max_workers}; min_interval_per_key: {args.min_interval}", flush=True)
    print(f"Already complete/skipped: {len(done_ids)}", flush=True)

    tasks: list[tuple[int, dict[str, Any]]] = []
    for idx, packet in enumerate(iter_jsonl(args.input_jsonl)):
        table_id = packet.get("table_id", "")
        if table_id in done_ids:
            continue
        if args.priority_min >= 0 and int(packet.get("review_priority", {}).get("score", 0)) < args.priority_min:
            continue
        tasks.append((idx, packet))
        if args.limit and len(tasks) >= args.limit:
            break
    total = len(tasks)
    print(f"Queued for review: {total}", flush=True)
    if total == 0:
        written = write_results_csv(results_jsonl, results_csv)
        print(f"No new work. CSV rows: {written}", flush=True)
        return

    write_lock = threading.Lock()
    completed = 0
    started = time.time()
    counts = {"ok": 0, "error": 0, "dry_run": 0}
    with ThreadPoolExecutor(max_workers=max(1, args.max_workers)) as executor:
        futures = [executor.submit(review_one, task, args, key_states) for task in tasks]
        for future in as_completed(futures):
            row = future.result()
            append_jsonl(results_jsonl, row, write_lock)
            completed += 1
            counts[row.get("status", "error")] = counts.get(row.get("status", "error"), 0) + 1
            if completed == 1 or completed % 10 == 0 or completed == total:
                elapsed = time.time() - started
                rate = completed / elapsed if elapsed else 0
                print(
                    f"[{completed}/{total}] ok={counts.get('ok', 0)} error={counts.get('error', 0)} "
                    f"dry_run={counts.get('dry_run', 0)} rate={rate:.2f}/s",
                    flush=True,
                )
    written = write_results_csv(results_jsonl, results_csv)
    report = {
        "input_jsonl": str(args.input_jsonl),
        "results_jsonl": str(results_jsonl),
        "results_csv": str(results_csv),
        "model": args.model,
        "queued": total,
        "csv_rows": written,
        "status_counts": counts,
        "elapsed_seconds": round(time.time() - started, 2),
    }
    (args.output_dir / "run_summary.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2), flush=True)


if __name__ == "__main__":
    main()
