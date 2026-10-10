#!/usr/bin/env python3
"""District lineage 1872-1941 from the India State and District Evolution Database
(Jolad et al., India State Stories; Harvard Dataverse doi:10.7910/DVN/D1AGUR; CC BY 4.0;
file district_evolution_1872_1941.xlsx). Each row links a unit in one census to the unit
(s) it became in the next. Coverage: present-day India.

ancestors(name_1941) -> {year: [(admin_div, type, district), ...]} : the period unit(s)
whose territory a 1941 unit lay in, for every earlier census.
"""
import os, re, collections, unicodedata, functools
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "mapdata/district_evolution/district_evolution_1872_1941.xlsx")
YEARS = [1872, 1881, 1891, 1901, 1911, 1921, 1931, 1941]


def nm(s):
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    s = re.sub(r"\b(state|district|the)\b", " ", s)
    return re.sub(r"[^a-z]", "", s)


@functools.lru_cache(maxsize=1)
def _links():
    rows = list(openpyxl.load_workbook(SRC, read_only=True).worksheets[0].iter_rows(values_only=True))[1:]
    back = collections.defaultdict(list)          # (dest_year, dest_key) -> [source unit]
    names41 = collections.defaultdict(list)       # name -> [(admin, type, name)] at 1941
    for sa, st, sd, da, dt, dd, sy, dy, _ in rows:
        src, dst = (sa, st, sd), (da, dt, dd)
        back[(dy, dst)].append(src)
        if dy == 1941:
            names41[nm(dd)].append(dst)
    return back, names41


@functools.lru_cache(maxsize=None)
def units_in(year):
    """every unit the database lists in a census year"""
    rows = list(openpyxl.load_workbook(SRC, read_only=True).worksheets[0].iter_rows(values_only=True))[1:]
    return sorted({(sa, st, sd) for sa, st, sd, da, dt, dd, sy, dy, _ in rows if sy == year} |
                  {(da, dt, dd) for sa, st, sd, da, dt, dd, sy, dy, _ in rows if dy == year})


def ancestors(name_1941, admin_hint=None):
    back, names41 = _links()
    cands = names41.get(nm(name_1941), [])
    if admin_hint and len(cands) > 1:
        cands = [c for c in cands if nm(c[0])[:5] == nm(admin_hint)[:5]] or cands
    if not cands:
        return None
    out = {1941: [cands[0]]}
    cur = [cands[0]]
    for y0, y1 in zip(reversed(YEARS[:-1]), reversed(YEARS[1:])):
        prev = []
        for u in cur:
            for s in back.get((y1, u), []):
                if s not in prev:
                    prev.append(s)
        out[y0] = prev
        cur = prev or cur
    return out


if __name__ == "__main__":
    for n in ["Guntur", "Chittoor", "Ramnad", "Drug", "Palamau", "Ranchi", "Anantapur", "Koraput", "Yeotmal",
              "Bijapur", "Almora", "Nainital", "East Khandesh", "Godavari East", "Benares", "Coorg"]:
        a = ancestors(n)
        print(f"{n:14}", {y: [u[2] for u in a[y]] for y in YEARS[:-1]} if a else "not in database")
