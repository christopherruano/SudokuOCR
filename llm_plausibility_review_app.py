#!/usr/bin/env python3
"""Review app for Gemini plausibility decisions.

The app joins model output, table packets, source images, and per-table Excel
workbooks so the final human decision can be made quickly.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from flask import Flask, abort, redirect, render_template_string, request, send_file, url_for
from openpyxl import load_workbook


ROOT = Path("export_final/final_validated")
SENDOUT = Path("export_final/sendout_final")
DEFAULT_RESULTS = ROOT / "llm_plausibility_review" / "gemini_flash_full" / "gemini_plausibility_results.csv"
DEFAULT_PACKETS = ROOT / "llm_plausibility_review" / "validated_table_packets.jsonl"
DEFAULT_DECISIONS = ROOT / "llm_plausibility_review" / "human_review_decisions.csv"
DEFAULT_RERUN_FOCUS_QUEUE = ROOT / "llm_plausibility_review" / "rerun_focus_queue.csv"
DEFAULT_CONTINUATION_CANDIDATES = Path("export_final/assembled_all_read_v2/continuation_candidates.csv")
DEFAULT_ASSEMBLED_JSONL = Path("export_final/assembled_all_read_v2/assembled_tables.jsonl")
DEFAULT_ASSEMBLED_DECISIONS = Path("export_final/assembled_all_read_review/assembled_group_review_decisions.csv")
DEFAULT_ASSEMBLED_REVIEW_DIR = Path("export_final/assembled_all_read_review")
DEFAULT_GROUP_REVIEW_CSV = Path("export_final/assembled_all_read_v2/group_review_1901.csv")
DEFAULT_GROUP_REVIEW_DECISIONS = Path("export_final/assembled_all_read_review/group_review_decisions.csv")
DEFAULT_DECIMAL_ROLLOUT_DIR = Path("export_final/sendout_final/review/decimal_rollout")
DEFAULT_DECIMAL_DECISIONS = Path("export_final/sendout_final/review/decimal_rollout/decimal_review_decisions.csv")
DEFAULT_DECIMAL_DECISION_QUEUE = Path("export_final/sendout_final/review/decimal_rollout/decimal_decision_queue.csv")
DEFAULT_DECIMAL_DP_DECISIONS = Path("export_final/sendout_final/review/decimal_rollout/decimal_decision_review_decisions.csv")
DEFAULT_TABLES_CONFIRMED = SENDOUT / "data" / "tables_confirmed.csv"
DEFAULT_FACTS_CONFIRMED = SENDOUT / "data" / "facts_confirmed.csv"
DEFAULT_FACTS_CLEANED = SENDOUT / "data" / "facts_confirmed.cleaned.csv"
DEFAULT_REVALIDATED_TABLES = SENDOUT / "review" / "revalidate" / "revalidated_tables.csv"
DEFAULT_RECONCILIATION_ANNOTATIONS = SENDOUT / "review" / "validation_candidates" / "failure_annotations.json"
DEFAULT_RECONCILIATION_DECISIONS = SENDOUT / "review" / "validation_candidates" / "reconciliation_review_decisions.csv"
DIGITIZED_TABLES_REFERENCED = Path("export/digitized_tables_referenced/digitized_tables")
OCR_RESULTS_REFERENCED = Path("export/ocr_results_referenced/digitized_tables")
APP_PORT = 8771

DECISION_FIELDS = [
    "table_id",
    "human_decision",
    "rerun_focus_area",
    "human_notes",
    "reviewer",
    "reviewed_at",
    "llm_decision_at_review",
    "llm_reason_at_review",
]

RERUN_FOCUS_FIELDS = [
    "table_id",
    "human_decision",
    "rerun_focus_area",
    "llm_fix_instruction",
    "llm_decision",
    "llm_reason",
    "llm_flags",
    "source_image_local",
    "excel_path",
    "province",
    "year",
    "table_family",
    "measure_category",
    "format_class",
    "title",
]

ASSEMBLED_DECISION_FIELDS = [
    "group_id",
    "table_id",
    "review_decision",
    "review_notes",
    "reviewer",
    "reviewed_at",
    "province",
    "year",
    "table_family",
    "title",
    "n_pages",
    "overall_strategy",
    "assembled_constraint_checks",
    "assembled_constraint_failures",
    "assembled_all_passed",
]

ASSEMBLED_REVIEW_INDEX_FIELDS = [
    "group_id",
    "table_id",
    "review_decision",
    "review_notes",
    "province",
    "year",
    "table_family",
    "title",
    "n_pages",
    "overall_strategy",
    "assembled_constraint_checks",
    "assembled_constraint_failures",
    "assembled_all_passed",
    "use_as_compiled_table",
]

GROUP_REVIEW_DECISION_FIELDS = [
    "group_id",
    "human_decision",
    "human_notes",
    "reviewer",
    "reviewed_at",
    "claude_decision",
    "confidence",
    "approved_correct_first_source_json",
    "approved_member_order_source_jsons",
    "approved_merge_with_group_ids",
    "approved_missing_preceding_source_jsons",
    "approved_extra_member_source_jsons",
    "approved_split_subgroups",
    "province",
    "year",
    "table_family",
]

GROUP_REVIEW_ACTION_FIELDS = [
    "group_id",
    "human_decision",
    "claude_decision",
    "confidence",
    "correct_first_source_json",
    "member_order_source_jsons",
    "merge_with_group_ids",
    "missing_preceding_source_jsons",
    "extra_member_source_jsons",
    "proposed_split_subgroups",
    "human_notes",
    "province",
    "year",
    "table_family",
]

DECIMAL_DECISION_FIELDS = [
    "queue_id",
    "human_decision",
    "human_notes",
    "reviewer",
    "reviewed_at",
    "review_bucket",
    "table_id",
    "raw_table_id",
    "province",
    "year",
    "table_family",
    "title",
    "decimal_places",
    "reason",
    "validation_failures",
]

DECIMAL_REVIEW_INDEX_FIELDS = [
    "queue_id",
    "review_bucket",
    "human_decision",
    "table_id",
    "raw_table_id",
    "province",
    "year",
    "table_family",
    "format_class",
    "measure_category",
    "title",
    "decimal_places",
    "reason",
    "rejected_reason",
    "validation_failures",
    "facts_original",
    "facts_cleaned",
    "changed_values",
    "image_exists",
]

DECIMAL_DP_DECISION_FIELDS = [
    "queue_id",
    "table_id",
    "human_dp",
    "human_flag",
    "human_notes",
    "reviewer",
    "reviewed_at",
    "category",
    "opus_dp",
    "opus_uniform",
    "opus_sample",
    "rate_min",
    "rate_max",
    "province",
    "year",
    "table_family",
    "title",
]

DECIMAL_DP_REVIEW_INDEX_FIELDS = [
    "queue_id",
    "table_id",
    "category",
    "human_dp",
    "human_flag",
    "province",
    "year",
    "table_family",
    "title",
    "opus_dp",
    "opus_uniform",
    "opus_sample",
    "rate_min",
    "rate_max",
    "agreement",
    "action",
    "image_exists",
]

RECONCILIATION_DECISION_FIELDS = [
    "table_id",
    "human_decision",
    "human_notes",
    "reviewer",
    "reviewed_at",
    "province",
    "year",
    "table_family",
    "title",
    "n_failures",
]

RECONCILIATION_REVIEW_INDEX_FIELDS = [
    "table_id",
    "human_decision",
    "province",
    "year",
    "table_family",
    "title",
    "n_failures",
    "levels",
    "image_exists",
    "facts_original",
    "facts_cleaned",
]

LOCAL_DIGITIZED_PREFIX = "export/digitized_tables_referenced/digitized_tables"
CLUSTER_DIGITIZED_PREFIX = "/n/home11/cruano/census_toc/digitized_tables/"
CLUSTER_OCR_PREFIX = "/n/home11/cruano/census_toc/ocr_results/digitized_tables/"


APP_HTML = """
<!doctype html>
<html>
<head>
  <meta charset="utf-8">
  <title>LLM Plausibility Review</title>
  <style>
    :root {
      --ink:#191715;
      --muted:#706a61;
      --bg:#ece7dc;
      --paper:#fffaf0;
      --panel:#fffdf8;
      --line:#d6cbbb;
      --accent:#263d35;
      --accent2:#b0612d;
      --bad:#8f2828;
      --warn:#8a650e;
      --ok:#27603a;
    }
    * { box-sizing:border-box; }
    body { margin:0; padding-bottom:230px; color:var(--ink); background:radial-gradient(circle at top left, #fff7df 0, var(--bg) 42%, #e6ddd0 100%); font-family: "Iowan Old Style", Georgia, serif; }
    header { position:sticky; top:0; z-index:3; background:rgba(255,250,240,.96); backdrop-filter: blur(8px); border-bottom:1px solid var(--line); padding:12px 16px; display:flex; gap:14px; justify-content:space-between; align-items:flex-start; }
    h1 { font-size:20px; margin:0 0 4px 0; }
    h2 { font-size:18px; margin:0; line-height:1.25; }
    main { display:grid; gap:14px; padding:14px; }
    .review-grid { display:grid; grid-template-columns:minmax(520px,1fr) minmax(520px,1fr); gap:14px; align-items:start; }
    .context-grid { display:grid; grid-template-columns:minmax(420px,.75fr) minmax(520px,1.25fr); gap:14px; align-items:start; }
    .stack { display:grid; gap:14px; align-content:start; }
    .card { background:var(--paper); border:1px solid var(--line); border-radius:16px; box-shadow:0 10px 28px rgba(49,34,16,.07); overflow:hidden; }
    .pad { padding:14px; }
    .muted { color:var(--muted); font-size:12px; }
    .mono { font-family:Menlo, Consolas, monospace; }
    .pill { display:inline-block; border:1px solid var(--line); border-radius:999px; padding:4px 9px; background:#f5eddf; margin:2px; font-size:12px; }
    .pill.ok { color:var(--ok); border-color:#9cbea4; background:#eef8ef; }
    .pill.warn { color:var(--warn); border-color:#d8bf78; background:#fff6d8; }
    .pill.bad { color:var(--bad); border-color:#dca2a2; background:#fff0ef; }
    .btn, button { appearance:none; border:0; border-radius:999px; padding:9px 13px; background:var(--accent); color:white; text-decoration:none; font-weight:700; cursor:pointer; font-family:inherit; }
    .btn.secondary, button.secondary { background:#d8c7ac; color:var(--ink); }
    .btn.ghost { background:transparent; color:var(--ink); border:1px solid var(--line); }
    .btn.bad, button.bad { background:var(--bad); }
    .btn.warn, button.warn { background:var(--accent2); }
    .toolbar { display:flex; gap:8px; flex-wrap:wrap; align-items:center; }
    .quick-actions { display:grid; grid-template-columns:repeat(5,minmax(100px,1fr)); gap:8px; margin:0; }
    .quick-actions button { border-radius:14px; padding:14px 10px; min-height:58px; font-size:15px; }
    .quick-actions small { display:block; opacity:.8; font-weight:400; margin-top:3px; }
    .meta { display:grid; grid-template-columns:repeat(2,minmax(0,1fr)); gap:8px 14px; font-size:13px; }
    .callout { border-left:4px solid var(--accent2); background:#fff5df; padding:11px 13px; line-height:1.35; }
    .callout.pass { border-color:var(--ok); background:#eef8ef; }
    .callout.bad { border-color:var(--bad); background:#fff0ef; }
    select, textarea, input { width:100%; border:1px solid var(--line); border-radius:12px; padding:10px; background:#fffdf8; font:14px/1.35 Menlo, Consolas, monospace; }
    textarea { min-height:58px; resize:vertical; }
    label { font-weight:700; display:block; margin-bottom:6px; }
    pre { white-space:pre-wrap; word-break:break-word; overflow:auto; max-height:320px; background:#fffdf8; border:1px solid var(--line); border-radius:12px; padding:12px; font:12px/1.4 Menlo, Consolas, monospace; }
    .image-toolbar { border-top:1px solid var(--line); border-bottom:1px solid var(--line); background:#f6eedf; position:sticky; top:74px; z-index:2; }
    .image-toolbar .zoom-readout { min-width:54px; text-align:center; font-family:Menlo, Consolas, monospace; font-size:12px; color:var(--muted); }
    .image-toolbar button { padding:7px 11px; border-radius:12px; }
    .image-wrap { background:#f8f4eb; padding:10px; overflow:auto; height:72vh; min-height:520px; text-align:left; border-top:1px solid var(--line); cursor:grab; overscroll-behavior:contain; }
    .image-wrap.dragging { cursor:grabbing; user-select:none; }
    .image-wrap img { display:block; width:100%; max-width:none; max-height:none; border:1px solid var(--line); background:white; transform-origin:top left; }
    .table-wrap { overflow:auto; height:72vh; min-height:520px; background:#fffdf8; border-top:1px solid var(--line); }
    .facts-wrap { overflow:auto; max-height:360px; background:#fffdf8; border-top:1px solid var(--line); }
    .decision-bar { position:fixed; left:0; right:0; bottom:0; z-index:5; background:rgba(255,250,240,.98); backdrop-filter:blur(10px); border-top:1px solid var(--line); padding:10px 14px; box-shadow:0 -10px 30px rgba(49,34,16,.12); }
    .decision-form { display:grid; grid-template-columns:minmax(360px,1.2fr) minmax(280px,.8fr) minmax(520px,1.4fr); gap:10px; align-items:start; }
    .decision-form label { margin-bottom:4px; }
    .bottom-hint { font-size:12px; color:var(--muted); margin-top:4px; }
    .decision-form textarea { min-height:86px; max-height:180px; }
    .note-tools { display:flex; gap:8px; flex-wrap:wrap; margin-top:6px; }
    .note-tools button { padding:6px 10px; border-radius:10px; font-size:12px; }
    table { border-collapse:collapse; width:100%; font-size:12px; }
    th,td { border:1px solid var(--line); padding:5px 7px; vertical-align:top; text-align:left; }
    th { background:#efe4d1; position:sticky; top:0; z-index:1; }
    .topline { display:flex; align-items:center; gap:8px; flex-wrap:wrap; }
    @media (max-width:1100px) {
      body { padding-bottom:260px; }
      header { display:block; }
      .review-grid, .context-grid, .decision-form { grid-template-columns:1fr; }
      .image-wrap, .table-wrap { height:auto; max-height:70vh; min-height:360px; }
    }
  </style>
</head>
<body>
  <header>
    <div>
      <div class="muted">LLM plausibility review | {{ index + 1 }} / {{ total }} | filter={{ filter_mode }}</div>
      <h1>{{ row.province }} {{ row.year }} | {{ row.title }}</h1>
      <div class="mono muted">{{ row.table_id }}</div>
    </div>
    <div class="toolbar">
      {% if prev_idx is not none %}<a class="btn secondary" data-prev href="{{ url_for('review_page', idx=prev_idx, filter=filter_mode) }}">Prev</a>{% endif %}
      {% if next_idx is not none %}<a class="btn secondary" data-next href="{{ url_for('review_page', idx=next_idx, filter=filter_mode) }}">Next</a>{% endif %}
      {% if next_unresolved_idx is not none %}<a class="btn" href="{{ url_for('review_page', idx=next_unresolved_idx, filter='unresolved') }}">Next unresolved</a>{% endif %}
    </div>
  </header>
  <main>
    <section class="card">
      <div class="pad toolbar">
        {% for mode,label in filters %}
          <a class="btn ghost" href="{{ url_for('review_page', idx=first_index(mode), filter=mode) }}">{{ label }}</a>
        {% endfor %}
        {% if assembled_total %}
          <a class="btn" href="{{ url_for('assembled_review_page', idx=assembled_first_index('unreviewed'), filter='unreviewed') }}">Review assembled groups</a>
        {% endif %}
        {% if group_review_total %}
          <a class="btn warn" href="{{ url_for('group_review_page', idx=group_review_first_index('unreviewed'), filter='unreviewed') }}">Review 1901 validation</a>
        {% endif %}
        {% if decimal_total %}
          <a class="btn warn" href="{{ url_for('decimal_review_page', idx=decimal_first_index('residue'), filter='residue') }}">Decimal/residue review</a>
        {% endif %}
        {% if reconciliation_total %}
          <a class="btn warn" href="{{ url_for('reconciliation_review_page', idx=reconciliation_first_index('unreviewed'), filter='unreviewed') }}">Reconciliation review</a>
        {% endif %}
        {% if has_image %}<a class="btn" href="{{ url_for('source_image_file', idx=index) }}" target="_blank">Open image</a>{% endif %}
        {% if has_excel %}<a class="btn secondary" href="{{ url_for('excel_file', idx=index) }}">Open Excel</a>{% endif %}
      </div>
    </section>

    <section class="review-grid">
      <div class="card">
        <div class="pad"><h2>Source Image</h2><div class="muted">{{ row.source_image_local or "not local" }}</div></div>
        {% if has_image %}
          <div class="pad toolbar image-toolbar">
            <button type="button" class="secondary" onclick="zoomImage(-25)">Zoom -</button>
            <button type="button" class="secondary" onclick="zoomImage(25)">Zoom +</button>
            <button type="button" class="secondary" onclick="setImageZoom(100)">Fit</button>
            <button type="button" class="secondary" onclick="setImageZoom(175)">175%</button>
            <span class="zoom-readout" id="zoomReadout">100%</span>
            <span class="muted">Ctrl/Cmd + wheel zooms; drag image to pan.</span>
          </div>
          <div class="image-wrap" id="imageViewport"><img id="sourceImage" src="{{ url_for('source_image_file', idx=index) }}" draggable="false"></div>
        {% else %}
          <div class="pad muted">No local source image found at: {{ row.source_image_local }}</div>
        {% endif %}
      </div>
      <div class="card">
        <div class="pad"><h2>Recorded Facts Matrix</h2><div class="muted">{{ excel_status }}</div></div>
        <div class="table-wrap">{{ matrix_html|safe }}</div>
      </div>
    </section>

    <section class="context-grid">
      <div class="card">
        <div class="pad topline">
          <span class="pill {{ decision_class }}">{{ row.llm_decision or row.status }}</span>
          <span class="pill">{{ row.confidence or "confidence?" }}</span>
          <span class="pill">priority {{ row.priority_score }}</span>
          <span class="pill">{{ row.measure_category }}</span>
          <span class="pill">{{ row.table_family }}</span>
        </div>
        <div class="pad">
          <div class="callout {{ decision_class }}">
            <strong>Model reason:</strong> {{ row.reason }}
            {% if row.suggested_next_action %}<br><strong>Suggested next action:</strong> {{ row.suggested_next_action }}{% endif %}
          </div>
          {% if row.is_continuation_candidate == "1" %}
          <div class="callout warn" style="margin-top:10px">
            <strong>Continuation candidate:</strong>
            {{ row.continuation_auto_attach_reason or "not auto-attached" }}
            {% if row.continuation_candidate_title %}
              <br><strong>Candidate previous table:</strong> {{ row.continuation_candidate_title }}
            {% endif %}
            {% if row.continuation_needs_manual_multipage_review == "1" %}
              <br><strong>Action:</strong> manually check adjacent page/table and decide whether this should be assembled.
            {% endif %}
          </div>
          {% endif %}
        </div>
        <div class="pad meta">
          <div><strong>Flags</strong><br><span class="muted">{{ row.flags }}</span></div>
          <div><strong>Review questions</strong><br><span class="muted">{{ row.review_questions }}</span></div>
          <div><strong>Model</strong><br><span class="muted">{{ row.model }} / {{ row.provider }}</span></div>
          <div><strong>Status</strong><br><span class="muted">{{ row.status }} {{ row.error }}</span></div>
          <div><strong>Format</strong><br><span class="muted">{{ row.format_class }}</span></div>
          <div><strong>Priority reasons</strong><br><span class="muted">{{ row.priority_reasons }}</span></div>
        </div>
      </div>
      <div class="card">
        <div class="pad"><h2>Recorded Fact Rows</h2><div class="muted">These are the actual fact rows saved for this table. If a value is here but not in the matrix above, that is a display bug.</div></div>
        <div class="facts-wrap">{{ facts_html|safe }}</div>
      </div>
    </section>
  </main>

  <form class="decision-bar" method="post" action="{{ url_for('save_page', idx=index, filter=filter_mode) }}">
    <input type="hidden" name="quick_decision" value="">
    <div class="decision-form">
      <div>
        <label>Focus/fix note</label>
        <textarea name="rerun_focus_area" placeholder="Type: e.g. check child widowed counts; religion duplicated in columns; image looks OK">{{ row.rerun_focus_area }}</textarea>
        <div class="note-tools">
          <button type="button" class="secondary" onclick="insertMultiTableNote()">Insert multi-table split note</button>
        </div>
        <div class="bottom-hint">Saved with every decision. For rerun/reject/hold, it is exported to the rerun-focus queue for the next LLM/pipeline pass.</div>
      </div>
      <div>
        <label>Optional notes</label>
        <textarea name="human_notes" placeholder="Extra context, if needed">{{ row.human_notes }}</textarea>
      </div>
      <div>
        <label>Decision buttons</label>
        <div class="quick-actions">
          <button type="button" onclick="quickSubmit('accept')">Accept<small>1</small></button>
          <button type="button" onclick="quickSubmit('image_compared_accept')">Image OK<small>2</small></button>
          <button class="warn" type="button" onclick="quickSubmit('hold')">Hold<small>3</small></button>
          <button class="warn" type="button" onclick="quickSubmit('rerun')">Rerun<small>4</small></button>
          <button class="bad" type="button" onclick="quickSubmit('reject')">Reject<small>5</small></button>
        </div>
      </div>
    </div>
    <select name="human_decision" style="display:none">
      {% for opt in ["", "accept", "reject", "rerun", "hold", "image_compared_accept", "image_compared_reject"] %}
        <option value="{{ opt }}" {% if row.human_decision == opt %}selected{% endif %}>{{ opt or "(unset)" }}</option>
      {% endfor %}
    </select>
  </form>
  <script>
    let imageZoom = 100;
    function imageEls() {
      return {
        viewport: document.getElementById('imageViewport'),
        image: document.getElementById('sourceImage'),
        readout: document.getElementById('zoomReadout')
      };
    }
    function setImageZoom(value) {
      const els = imageEls();
      if (!els.viewport || !els.image) return;
      const oldScrollW = Math.max(els.viewport.scrollWidth, 1);
      const oldScrollH = Math.max(els.viewport.scrollHeight, 1);
      const centerX = (els.viewport.scrollLeft + els.viewport.clientWidth / 2) / oldScrollW;
      const centerY = (els.viewport.scrollTop + els.viewport.clientHeight / 2) / oldScrollH;
      imageZoom = Math.max(50, Math.min(500, value));
      els.image.style.width = imageZoom + '%';
      if (els.readout) els.readout.textContent = imageZoom + '%';
      requestAnimationFrame(() => {
        els.viewport.scrollLeft = centerX * els.viewport.scrollWidth - els.viewport.clientWidth / 2;
        els.viewport.scrollTop = centerY * els.viewport.scrollHeight - els.viewport.clientHeight / 2;
      });
    }
    function zoomImage(delta) {
      setImageZoom(imageZoom + delta);
    }
    function insertMultiTableNote() {
      const field = document.querySelector('textarea[name="rerun_focus_area"]');
      if (!field) return;
      const note = 'Multiple printed tables on this image. Extract ALL printed tables, but keep each as a separate table/section with its own exact printed title or statement number, unit/base, row axis, and column headers. Do not merge values across subtables. If subtables have different formats, create separate logical outputs/table sections rather than forcing them into one shared schema.';
      const current = field.value.trim();
      field.value = current ? current + "\\n\\n" + note : note;
      field.focus();
      field.selectionStart = field.selectionEnd = field.value.length;
    }
    (function initImageViewer() {
      const els = imageEls();
      if (!els.viewport || !els.image) return;
      setImageZoom(100);
      els.viewport.addEventListener('wheel', function(evt) {
        if (!(evt.ctrlKey || evt.metaKey)) return;
        evt.preventDefault();
        zoomImage(evt.deltaY < 0 ? 25 : -25);
      }, { passive:false });
      let dragging = false;
      let startX = 0;
      let startY = 0;
      let startLeft = 0;
      let startTop = 0;
      els.viewport.addEventListener('pointerdown', function(evt) {
        if (evt.button !== 0) return;
        dragging = true;
        startX = evt.clientX;
        startY = evt.clientY;
        startLeft = els.viewport.scrollLeft;
        startTop = els.viewport.scrollTop;
        els.viewport.classList.add('dragging');
        els.viewport.setPointerCapture(evt.pointerId);
      });
      els.viewport.addEventListener('pointermove', function(evt) {
        if (!dragging) return;
        els.viewport.scrollLeft = startLeft - (evt.clientX - startX);
        els.viewport.scrollTop = startTop - (evt.clientY - startY);
      });
      function stopDrag(evt) {
        dragging = false;
        els.viewport.classList.remove('dragging');
        try { els.viewport.releasePointerCapture(evt.pointerId); } catch (err) {}
      }
      els.viewport.addEventListener('pointerup', stopDrag);
      els.viewport.addEventListener('pointercancel', stopDrag);
    })();
    function quickSubmit(decision) {
      const form = document.querySelector('form');
      const quick = document.querySelector('input[name="quick_decision"]');
      const select = document.querySelector('select[name="human_decision"]');
      if (quick) quick.value = decision;
      if (select) select.value = decision;
      if (form) form.submit();
    }
    document.addEventListener('keydown', function(evt) {
      if (evt.target && ['TEXTAREA','INPUT','SELECT'].includes(evt.target.tagName)) return;
      if (evt.key === 'j') { const el=document.querySelector('a[data-next]'); if (el) window.location=el.href; }
      if (evt.key === 'k') { const el=document.querySelector('a[data-prev]'); if (el) window.location=el.href; }
      if (evt.key === '1') { quickSubmit('accept'); }
      if (evt.key === '2') { quickSubmit('image_compared_accept'); }
      if (evt.key === '3') { quickSubmit('hold'); }
      if (evt.key === '4') { quickSubmit('rerun'); }
      if (evt.key === '5') { quickSubmit('reject'); }
      if (evt.key === '+' || evt.key === '=') { zoomImage(25); }
      if (evt.key === '-') { zoomImage(-25); }
      if (evt.key === '0') { setImageZoom(100); }
      if (evt.key === 's') { const form=document.querySelector('form'); if (form) form.submit(); }
    });
  </script>
</body>
</html>
"""


ASSEMBLED_HTML = """
<!doctype html>
<html>
<head>
  <meta charset="utf-8">
  <title>Assembled Group Review</title>
  <style>
    :root {
      --ink:#191715;
      --muted:#706a61;
      --bg:#ece7dc;
      --paper:#fffaf0;
      --panel:#fffdf8;
      --line:#d6cbbb;
      --accent:#263d35;
      --accent2:#b0612d;
      --bad:#8f2828;
      --warn:#8a650e;
      --ok:#27603a;
    }
    * { box-sizing:border-box; }
    body { margin:0; padding-bottom:154px; color:var(--ink); background:radial-gradient(circle at top left, #fff7df 0, var(--bg) 42%, #e6ddd0 100%); font-family:"Iowan Old Style", Georgia, serif; }
    header { position:sticky; top:0; z-index:4; background:rgba(255,250,240,.97); backdrop-filter:blur(8px); border-bottom:1px solid var(--line); padding:12px 16px; display:flex; justify-content:space-between; gap:12px; align-items:flex-start; }
    h1 { font-size:20px; margin:0 0 4px 0; line-height:1.2; }
    h2 { font-size:17px; margin:0; }
    h3 { font-size:15px; margin:0 0 6px 0; }
    main { display:grid; grid-template-columns:minmax(420px,42vw) minmax(620px,1fr); gap:14px; padding:14px; align-items:start; }
    .card { background:var(--paper); border:1px solid var(--line); border-radius:16px; box-shadow:0 10px 28px rgba(49,34,16,.07); overflow:hidden; margin-bottom:14px; }
    .pad { padding:14px; }
    .muted { color:var(--muted); font-size:12px; }
    .mono { font-family:Menlo, Consolas, monospace; }
    .toolbar { display:flex; gap:8px; flex-wrap:wrap; align-items:center; }
    .btn, button { appearance:none; border:0; border-radius:999px; padding:9px 13px; background:var(--accent); color:white; text-decoration:none; font-weight:700; cursor:pointer; font-family:inherit; }
    .btn.secondary, button.secondary { background:#d8c7ac; color:var(--ink); }
    .btn.ghost { background:transparent; color:var(--ink); border:1px solid var(--line); }
    .btn.bad, button.bad { background:var(--bad); }
    .btn.warn, button.warn { background:var(--accent2); }
    .pill { display:inline-block; border:1px solid var(--line); border-radius:999px; padding:4px 9px; background:#f5eddf; margin:2px; font-size:12px; }
    .pill.ok { color:var(--ok); border-color:#9cbea4; background:#eef8ef; }
    .pill.warn { color:var(--warn); border-color:#d8bf78; background:#fff6d8; }
    .pill.bad { color:var(--bad); border-color:#dca2a2; background:#fff0ef; }
    .meta { display:grid; grid-template-columns:145px minmax(0,1fr); gap:6px 12px; font-size:13px; }
    .pages { position:sticky; top:86px; max-height:calc(100vh - 112px); overflow:auto; padding-right:2px; }
    .image-card { background:#26231f; color:#f7eedf; border-radius:14px; padding:10px; margin-bottom:12px; border:1px solid #514a42; }
    .image-card.context { background:#382d20; border-color:#7d6344; }
    .image-card img { width:100%; display:block; border:1px solid #514a42; background:white; }
    .caption { margin-top:8px; font-size:12px; color:#dfd1bd; overflow-wrap:anywhere; }
    .compiled { min-width:0; }
    .section-head { display:flex; justify-content:space-between; gap:12px; align-items:baseline; border-bottom:1px solid var(--line); }
    .table-wrap { overflow:auto; max-height:72vh; background:#fffdf8; }
    table { border-collapse:collapse; width:100%; font-size:12px; }
    th,td { border:1px solid var(--line); padding:5px 7px; vertical-align:top; text-align:left; white-space:nowrap; }
    th { background:#efe4d1; position:sticky; top:0; z-index:1; }
    td:first-child, th:first-child { position:sticky; left:0; background:#fbf4e9; z-index:2; }
    th:first-child { background:#e6d9c3; z-index:3; }
    textarea { width:100%; min-height:82px; resize:vertical; border:1px solid var(--line); border-radius:12px; padding:10px; background:#fffdf8; font:14px/1.35 Menlo, Consolas, monospace; }
    .decision-bar { position:fixed; left:0; right:0; bottom:0; z-index:5; background:rgba(255,250,240,.98); backdrop-filter:blur(10px); border-top:1px solid var(--line); padding:10px 14px; box-shadow:0 -10px 30px rgba(49,34,16,.12); }
    .decision-form { display:grid; grid-template-columns:minmax(420px,1fr) minmax(520px,1.3fr); gap:12px; align-items:start; }
    .quick-actions { display:grid; grid-template-columns:repeat(4,minmax(130px,1fr)); gap:8px; }
    .quick-actions button { border-radius:14px; min-height:58px; font-size:15px; }
    .quick-actions small { display:block; opacity:.82; font-weight:400; margin-top:3px; }
    @media (max-width:1120px) {
      body { padding-bottom:240px; }
      header { display:block; }
      main, .decision-form { grid-template-columns:1fr; }
      .pages { position:relative; top:0; max-height:none; }
    }
  </style>
</head>
<body>
  <header>
    <div>
      <div class="muted">Assembled group review | {{ index + 1 }} / {{ total }} | filter={{ filter_mode }}</div>
      <h1>{{ group.province }} {{ group.year }} | {{ display_title }}</h1>
      <div class="mono muted">{{ table_id }}</div>
    </div>
    <div class="toolbar">
      <a class="btn secondary" href="{{ url_for('review_page', idx=normal_first_index, filter='continuation') }}">Back to anomaly review</a>
      {% if group_review_total %}
        <a class="btn warn" href="{{ url_for('group_review_page', idx=group_review_first_index('unreviewed'), filter='unreviewed') }}">Review 1901 validation</a>
      {% endif %}
      {% if decimal_total %}
        <a class="btn warn" href="{{ url_for('decimal_review_page', idx=decimal_first_index('residue'), filter='residue') }}">Decimal/residue review</a>
      {% endif %}
      {% if reconciliation_total %}
        <a class="btn warn" href="{{ url_for('reconciliation_review_page', idx=reconciliation_first_index('unreviewed'), filter='unreviewed') }}">Reconciliation review</a>
      {% endif %}
      {% if prev_idx is not none %}<a class="btn secondary" data-prev href="{{ url_for('assembled_review_page', idx=prev_idx, filter=filter_mode) }}">Prev group</a>{% endif %}
      {% if next_idx is not none %}<a class="btn secondary" data-next href="{{ url_for('assembled_review_page', idx=next_idx, filter=filter_mode) }}">Next group</a>{% endif %}
      {% for mode,label in filters %}
        <a class="btn ghost" href="{{ url_for('assembled_review_page', idx=first_index(mode), filter=mode) }}">{{ label }}</a>
      {% endfor %}
    </div>
  </header>

  <main>
    <aside class="pages">
      <div class="card">
        <div class="pad">
          <h2>Member Source Pages</h2>
          <div class="muted">These are the original and continuation pages that form this compiled group.</div>
        </div>
        <div class="pad meta">
          <strong>Primary title</strong><span>{{ anchor.title or display_title }}</span>
          <strong>Group title</strong><span>{{ group.title }}</span>
          <strong>Group</strong><span class="mono">{{ group.group_id }}</span>
          <strong>Pages</strong><span>{{ group.n_pages }}</span>
          <strong>Strategy</strong><span>{{ group.overall_strategy }}</span>
          <strong>Checks</strong><span>{{ report.total_checks }} total, {{ report.failed }} failed</span>
          <strong>All passed</strong><span>{{ report.all_passed }}</span>
          <strong>Decision</strong><span>{{ decision.review_decision or "unreviewed" }}</span>
          <strong>Anchor source</strong><span class="mono">{{ anchor.source_json }}</span>
          <strong>Anchor reason</strong><span>{{ anchor.reason }}</span>
        </div>
      </div>

      {% if member_images %}
        {% for image in member_images %}
          <div class="image-card">
            {% if image.exists %}
              <a href="{{ url_for('assembled_member_image_file', idx=index, member_idx=image.view_index) }}" target="_blank">
                <img src="{{ url_for('assembled_member_image_file', idx=index, member_idx=image.view_index) }}">
              </a>
            {% else %}
              <div class="muted">Image not found locally.</div>
            {% endif %}
            <div class="caption">
              <strong>{% if image.is_anchor %}Primary/original page{% else %}Page {{ image.member_index }}{% endif %}</strong>
              {% if image.reference_only %}<span class="pill warn">reference only</span>{% endif %}<br>
              {{ image.source_json }}<br>
              {{ image.source_image }}
            </div>
          </div>
        {% endfor %}
      {% else %}
        <div class="card pad muted">No member images listed for this assembled group.</div>
      {% endif %}

      {% if previous_context.images %}
        <div class="card">
          <div class="pad">
            <h2>Previous Table-Directory Context</h2>
            <div class="muted">Not included in this group. Use this only to check whether the apparent first page is missing a preceding page.</div>
            <div class="mono muted">{{ previous_context.dir_name }}</div>
          </div>
        </div>
        {% for image in previous_context.images %}
          <div class="image-card context">
            {% if image.exists %}
              <a href="{{ url_for('assembled_context_image_file', idx=index, context_idx=image.view_index) }}" target="_blank">
                <img src="{{ url_for('assembled_context_image_file', idx=index, context_idx=image.view_index) }}">
              </a>
            {% else %}
              <div class="muted">Context image not found locally.</div>
            {% endif %}
            <div class="caption">
              <strong>Previous context {{ image.view_index + 1 }}</strong><br>
              {{ image.title }}<br>
              {{ image.source_json }}<br>
              {{ image.source_image }}
            </div>
          </div>
        {% endfor %}
      {% endif %}
    </aside>

    <section class="compiled">
      <div class="card">
        <div class="pad">
          <h2>Compiled Table</h2>
          <div class="muted">One decision below applies to this whole compiled table. Once accepted, downstream export should use this as one table instead of its member page fragments.</div>
          <div>
            <span class="pill {% if report.failed %}bad{% elif report.total_checks %}ok{% else %}warn{% endif %}">constraints: {{ report.failed }} failed</span>
            <span class="pill">{{ group.table_family }}</span>
            <span class="pill">{{ group.overall_strategy }}</span>
            <span class="pill">{{ section_count }} sections</span>
            <span class="pill">{{ group.n_pages }} pages</span>
          </div>
        </div>
      </div>

      {% if not sections %}
        <div class="card pad muted">No assembled section rows available.</div>
      {% endif %}
      {% for section in sections %}
        <div class="card">
          <div class="pad section-head">
            <div>
              <h3>Section {{ section.index }}: {{ section.name }}</h3>
              <div class="muted">merge={{ section.merge_strategy }}; input_pages={{ section.input_pages }}; output_rows={{ section.output_rows }}; shape={{ section.shape }}</div>
            </div>
          </div>
          <div class="table-wrap">{{ section.html|safe }}</div>
        </div>
      {% endfor %}
    </section>
  </main>

  <form class="decision-bar" method="post" action="{{ url_for('save_assembled_page', idx=index, filter=filter_mode) }}">
    <input type="hidden" name="quick_decision" value="">
    <div class="decision-form">
      <div>
        <label><strong>Group review notes</strong></label>
        <textarea name="review_notes" placeholder="What you checked, anomalies, or rerun instructions for this whole multipage table">{{ decision.review_notes }}</textarea>
        <div class="muted">Saved at group level and included in the accepted/rejected/rerun exports.</div>
      </div>
      <div>
        <label><strong>Decision</strong></label>
        <div class="quick-actions">
          {% if report.failed %}
            <button class="warn" type="button" onclick="quickSubmit('grouping_ok_rerun_facts')">Grouping OK, rerun facts<small>1</small></button>
          {% else %}
            <button type="button" onclick="quickSubmit('keep')">Accept final compiled table<small>1</small></button>
          {% endif %}
          <button class="bad" type="button" onclick="quickSubmit('drop')">Reject group<small>2</small></button>
          <button class="warn" type="button" onclick="quickSubmit('rerun')">Rerun / split<small>3</small></button>
          <button class="secondary" type="button" onclick="quickSubmit('hold')">Hold<small>4</small></button>
        </div>
      </div>
    </div>
  </form>
  <script>
    const primaryDecision = "{{ 'grouping_ok_rerun_facts' if report.failed else 'keep' }}";
    function quickSubmit(decision) {
      const form = document.querySelector('form');
      const quick = document.querySelector('input[name="quick_decision"]');
      if (quick) quick.value = decision;
      if (form) form.submit();
    }
    document.addEventListener('keydown', function(evt) {
      if (evt.target && ['TEXTAREA','INPUT','SELECT'].includes(evt.target.tagName)) return;
      if (evt.key === 'j') { const el=document.querySelector('a[data-next]'); if (el) window.location=el.href; }
      if (evt.key === 'k') { const el=document.querySelector('a[data-prev]'); if (el) window.location=el.href; }
      if (evt.key === '1') quickSubmit(primaryDecision);
      if (evt.key === '2') quickSubmit('drop');
      if (evt.key === '3') quickSubmit('rerun');
      if (evt.key === '4') quickSubmit('hold');
    });
  </script>
</body>
</html>
"""


GROUP_REVIEW_HTML = """
<!doctype html>
<html>
<head>
  <meta charset="utf-8">
  <title>1901 Group Validation Review</title>
  <style>
    :root {
      --ink:#191715;
      --muted:#706a61;
      --bg:#ece7dc;
      --paper:#fffaf0;
      --line:#d6cbbb;
      --accent:#263d35;
      --accent2:#b0612d;
      --bad:#8f2828;
      --warn:#8a650e;
      --ok:#27603a;
    }
    * { box-sizing:border-box; }
    body { margin:0; padding-bottom:172px; color:var(--ink); background:radial-gradient(circle at top left, #fff7df 0, var(--bg) 42%, #e6ddd0 100%); font-family:"Iowan Old Style", Georgia, serif; }
    header { position:sticky; top:0; z-index:4; background:rgba(255,250,240,.97); backdrop-filter:blur(8px); border-bottom:1px solid var(--line); padding:12px 16px; display:flex; justify-content:space-between; gap:12px; align-items:flex-start; }
    h1 { font-size:20px; margin:0 0 4px 0; line-height:1.2; }
    h2 { font-size:17px; margin:0; }
    h3 { font-size:15px; margin:0 0 8px 0; }
    main { display:grid; grid-template-columns:minmax(430px,42vw) minmax(620px,1fr); gap:14px; padding:14px; align-items:start; }
    .card { background:var(--paper); border:1px solid var(--line); border-radius:16px; box-shadow:0 10px 28px rgba(49,34,16,.07); overflow:hidden; margin-bottom:14px; }
    .pad { padding:14px; }
    .muted { color:var(--muted); font-size:12px; }
    .mono { font-family:Menlo, Consolas, monospace; }
    .toolbar { display:flex; gap:8px; flex-wrap:wrap; align-items:center; }
    .btn, button { appearance:none; border:0; border-radius:999px; padding:9px 13px; background:var(--accent); color:white; text-decoration:none; font-weight:700; cursor:pointer; font-family:inherit; }
    .btn.secondary, button.secondary { background:#d8c7ac; color:var(--ink); }
    .btn.ghost { background:transparent; color:var(--ink); border:1px solid var(--line); }
    .btn.bad, button.bad { background:var(--bad); }
    .btn.warn, button.warn { background:var(--accent2); }
    .pill { display:inline-block; border:1px solid var(--line); border-radius:999px; padding:4px 9px; background:#f5eddf; margin:2px; font-size:12px; }
    .pill.ok { color:var(--ok); border-color:#9cbea4; background:#eef8ef; }
    .pill.warn { color:var(--warn); border-color:#d8bf78; background:#fff6d8; }
    .pill.bad { color:var(--bad); border-color:#dca2a2; background:#fff0ef; }
    .pages { position:sticky; top:86px; max-height:calc(100vh - 112px); overflow:auto; padding-right:2px; }
    .image-card { background:#26231f; color:#f7eedf; border-radius:14px; padding:10px; margin-bottom:12px; border:1px solid #514a42; }
    .image-card.proposed { background:#202b27; border-color:#536b5f; }
    .image-card.missing { background:#382d20; border-color:#7d6344; }
    .image-card img { width:100%; display:block; border:1px solid #514a42; background:white; }
    .caption { margin-top:8px; font-size:12px; color:#dfd1bd; overflow-wrap:anywhere; }
    .meta { display:grid; grid-template-columns:165px minmax(0,1fr); gap:7px 12px; font-size:13px; }
    .callout { border-left:4px solid var(--accent2); background:#fff5df; padding:11px 13px; line-height:1.35; }
    .callout.ok { border-color:var(--ok); background:#eef8ef; }
    .callout.bad { border-color:var(--bad); background:#fff0ef; }
    pre { white-space:pre-wrap; word-break:break-word; overflow:auto; max-height:390px; background:#fffdf8; border:1px solid var(--line); border-radius:12px; padding:12px; font:12px/1.4 Menlo, Consolas, monospace; }
    textarea { width:100%; min-height:98px; resize:vertical; border:1px solid var(--line); border-radius:12px; padding:10px; background:#fffdf8; font:14px/1.35 Menlo, Consolas, monospace; }
    .decision-bar { position:fixed; left:0; right:0; bottom:0; z-index:5; background:rgba(255,250,240,.98); backdrop-filter:blur(10px); border-top:1px solid var(--line); padding:10px 14px; box-shadow:0 -10px 30px rgba(49,34,16,.12); }
    .decision-form { display:grid; grid-template-columns:minmax(420px,1fr) minmax(680px,1.4fr); gap:12px; align-items:start; }
    .quick-actions { display:grid; grid-template-columns:repeat(5,minmax(120px,1fr)); gap:8px; }
    .quick-actions button { border-radius:14px; min-height:58px; font-size:15px; }
    .quick-actions small { display:block; opacity:.82; font-weight:400; margin-top:3px; }
    @media (max-width:1120px) {
      body { padding-bottom:260px; }
      header { display:block; }
      main, .decision-form { grid-template-columns:1fr; }
      .pages { position:relative; top:0; max-height:none; }
    }
  </style>
</head>
<body>
  <header>
    <div>
      <div class="muted">1901 group validation | {{ index + 1 }} / {{ total }} | filter={{ filter_mode }}</div>
      <h1>{{ row.province }} {{ row.year }} | {{ row.table_family }} | {{ row.decision }}</h1>
      <div class="mono muted">{{ row.group_id }}</div>
    </div>
    <div class="toolbar">
      <a class="btn secondary" href="{{ url_for('assembled_review_page', idx=assembled_first, filter='all') }}">Back to assembled groups</a>
      {% if prev_idx is not none %}<a class="btn secondary" data-prev href="{{ url_for('group_review_page', idx=prev_idx, filter=filter_mode) }}">Prev</a>{% endif %}
      {% if next_idx is not none %}<a class="btn secondary" data-next href="{{ url_for('group_review_page', idx=next_idx, filter=filter_mode) }}">Next</a>{% endif %}
      {% for mode,label in filters %}
        <a class="btn ghost" href="{{ url_for('group_review_page', idx=first_index(mode), filter=mode) }}">{{ label }}</a>
      {% endfor %}
    </div>
  </header>

  <main>
    <aside class="pages">
      <div class="card">
        <div class="pad">
          <h2>Current vs Proposed Pages</h2>
          <div class="muted">Top block is the current assembled group. Green block is Claude's proposed correct order. Brown block is recovered/missing-head evidence.</div>
        </div>
      </div>

      <div class="card"><div class="pad"><h3>Current Assembled Members</h3></div></div>
      {% for image in current_images %}
        <div class="image-card">
          {% if image.exists %}
            <a href="{{ url_for('group_review_image_file', idx=index, image_idx=image.view_index) }}" target="_blank">
              <img src="{{ url_for('group_review_image_file', idx=index, image_idx=image.view_index) }}">
            </a>
          {% else %}
            <div class="muted">Image not found locally.</div>
          {% endif %}
          <div class="caption"><strong>{{ image.label }}</strong><br>{{ image.source_json }}<br>{{ image.source_image }}</div>
        </div>
      {% endfor %}

      <div class="card"><div class="pad"><h3>Proposed Correct Order</h3></div></div>
      {% for image in proposed_images %}
        <div class="image-card proposed">
          {% if image.exists %}
            <a href="{{ url_for('group_review_image_file', idx=index, image_idx=image.view_index) }}" target="_blank">
              <img src="{{ url_for('group_review_image_file', idx=index, image_idx=image.view_index) }}">
            </a>
          {% else %}
            <div class="muted">Image not found locally.</div>
          {% endif %}
          <div class="caption"><strong>{{ image.label }}</strong><br>{{ image.source_json }}<br>{{ image.source_image }}</div>
        </div>
      {% endfor %}

      {% if evidence_images %}
        <div class="card"><div class="pad"><h3>Recovered / Missing-Head Evidence</h3></div></div>
        {% for image in evidence_images %}
          <div class="image-card missing">
            {% if image.exists %}
              <a href="{{ url_for('group_review_image_file', idx=index, image_idx=image.view_index) }}" target="_blank">
                <img src="{{ url_for('group_review_image_file', idx=index, image_idx=image.view_index) }}">
              </a>
            {% else %}
              <div class="muted">Image not found locally.</div>
            {% endif %}
            <div class="caption"><strong>{{ image.label }}</strong><br>{{ image.source_json }}<br>{{ image.source_image }}</div>
          </div>
        {% endfor %}
      {% endif %}
    </aside>

    <section>
      <div class="card">
        <div class="pad">
          <h2>Claude Proposal</h2>
          <div>
            <span class="pill {{ decision_class }}">{{ row.decision }}</span>
            <span class="pill {% if row.confidence == 'high' %}ok{% else %}warn{% endif %}">{{ row.confidence }}</span>
            <span class="pill">{{ row.n_pages }} pages</span>
            <span class="pill">{{ decision.human_decision or "unreviewed" }}</span>
          </div>
        </div>
        <div class="pad meta">
          <strong>Page sequence</strong><span>{{ row.page_sequence }}</span>
          <strong>Correct first</strong><span class="mono">{{ row.correct_first_source_json }}</span>
          <strong>Correct order</strong><span class="mono">{{ row.correct_member_order_source_jsons }}</span>
          <strong>Missing preceding</strong><span class="mono">{{ row.missing_preceding_source_jsons }}</span>
          <strong>Recovered head</strong><span class="mono">{{ row.recovered_head_source_json }}</span>
          <strong>Extra members</strong><span class="mono">{{ row.extra_member_source_jsons }}</span>
          <strong>Missing interior pages</strong><span>{{ row.missing_interior_page_numbers }}</span>
          <strong>Head pages</strong><span>{{ row.head_pages }}</span>
          <strong>Source PDF</strong><span>{{ row.source_pdf }}</span>
        </div>
        <div class="pad">
          <div class="callout {{ decision_class }}">
            <strong>Notes:</strong> {{ row.notes }}
          </div>
        </div>
      </div>

      {% if merge_targets %}
        <div class="card">
          <div class="pad">
            <h2>Merge Targets</h2>
            <div class="muted">If you approve, these group IDs should be merged into the same logical table/action set.</div>
          </div>
          <div class="pad">
            {% for target in merge_targets %}
              <div class="callout" style="margin-bottom:10px">
                <strong>{{ target.status }}</strong>
                {% if target.review_idx is not none %}
                  <a class="btn secondary" href="{{ url_for('group_review_page', idx=target.review_idx, filter='all') }}">Open target</a>
                {% endif %}
                <br><span class="mono">{{ target.group_id }}</span>
                {% if target.title %}<br>{{ target.title }}{% endif %}
                {% if target.decision %}<br><span class="muted">{{ target.decision }} / {{ target.confidence }}</span>{% endif %}
              </div>
            {% endfor %}
          </div>
        </div>
      {% endif %}

      {% if row.proposed_split_subgroups %}
        <div class="card">
          <div class="pad">
            <h2>Proposed Split Subgroups</h2>
            <div class="muted">Approve only if these sub-table boundaries match the images.</div>
          </div>
          <div class="pad"><pre>{{ row.proposed_split_subgroups }}</pre></div>
        </div>
      {% endif %}

      <div class="card">
        <div class="pad">
          <h2>Current Assembled Group Metadata</h2>
          {% if current_group %}
            <div class="meta">
              <strong>Title</strong><span>{{ current_group.title }}</span>
              <strong>Strategy</strong><span>{{ current_group.overall_strategy }}</span>
              <strong>Constraint checks</strong><span>{{ current_report.total_checks }} total, {{ current_report.failed }} failed</span>
              <strong>All passed</strong><span>{{ current_report.all_passed }}</span>
            </div>
          {% else %}
            <div class="muted">No current assembled JSONL group matched this group_id. Review the proposed image/order fields only.</div>
          {% endif %}
        </div>
      </div>
    </section>
  </main>

  <form class="decision-bar" method="post" action="{{ url_for('save_group_review_page', idx=index, filter=filter_mode) }}">
    <input type="hidden" name="quick_decision" value="">
    <div class="decision-form">
      <div>
        <label><strong>Human notes / corrections</strong></label>
        <textarea name="human_notes" placeholder="Confirm what you saw, or write corrected order/merge/split instructions here.">{{ decision.human_notes }}</textarea>
      </div>
      <div>
        <label><strong>Decision</strong></label>
        <div class="quick-actions">
          <button type="button" onclick="quickSubmit('approve_proposal')">Approve proposal<small>1</small></button>
          <button type="button" onclick="quickSubmit('approve_ok')">Clean OK<small>2</small></button>
          <button class="warn" type="button" onclick="quickSubmit('needs_manual')">Needs manual<small>3</small></button>
          <button class="warn" type="button" onclick="quickSubmit('rerun_or_reextract')">Rerun/reextract<small>4</small></button>
          <button class="bad" type="button" onclick="quickSubmit('reject_proposal')">Reject proposal<small>5</small></button>
        </div>
      </div>
    </div>
  </form>
  <script>
    function quickSubmit(decision) {
      const form = document.querySelector('form');
      const quick = document.querySelector('input[name="quick_decision"]');
      if (quick) quick.value = decision;
      if (form) form.submit();
    }
    document.addEventListener('keydown', function(evt) {
      if (evt.target && ['TEXTAREA','INPUT','SELECT'].includes(evt.target.tagName)) return;
      if (evt.key === 'j') { const el=document.querySelector('a[data-next]'); if (el) window.location=el.href; }
      if (evt.key === 'k') { const el=document.querySelector('a[data-prev]'); if (el) window.location=el.href; }
      if (evt.key === '1') quickSubmit('approve_proposal');
      if (evt.key === '2') quickSubmit('approve_ok');
      if (evt.key === '3') quickSubmit('needs_manual');
      if (evt.key === '4') quickSubmit('rerun_or_reextract');
      if (evt.key === '5') quickSubmit('reject_proposal');
    });
  </script>
</body>
</html>
"""


DECIMAL_REVIEW_HTML = """
<!doctype html>
<html>
<head>
  <meta charset="utf-8">
  <title>Decimal / Residue Review</title>
  <style>
    :root {
      --ink:#191715;
      --muted:#706a61;
      --bg:#ece7dc;
      --paper:#fffaf0;
      --line:#d6cbbb;
      --accent:#263d35;
      --accent2:#b0612d;
      --bad:#8f2828;
      --warn:#8a650e;
      --ok:#27603a;
    }
    * { box-sizing:border-box; }
    body { margin:0; padding-bottom:170px; color:var(--ink); background:radial-gradient(circle at top left, #fff7df 0, var(--bg) 42%, #e6ddd0 100%); font-family:"Iowan Old Style", Georgia, serif; }
    header { position:sticky; top:0; z-index:4; background:rgba(255,250,240,.97); backdrop-filter:blur(8px); border-bottom:1px solid var(--line); padding:12px 16px; display:flex; justify-content:space-between; gap:12px; align-items:flex-start; }
    h1 { font-size:20px; margin:0 0 4px 0; line-height:1.2; }
    h2 { font-size:17px; margin:0; }
    main { display:grid; grid-template-columns:minmax(500px,42vw) minmax(640px,1fr); gap:14px; padding:14px; align-items:start; }
    .card { background:var(--paper); border:1px solid var(--line); border-radius:16px; box-shadow:0 10px 28px rgba(49,34,16,.07); overflow:hidden; margin-bottom:14px; }
    .pad { padding:14px; }
    .muted { color:var(--muted); font-size:12px; }
    .mono { font-family:Menlo, Consolas, monospace; }
    .toolbar { display:flex; gap:8px; flex-wrap:wrap; align-items:center; }
    .btn, button { appearance:none; border:0; border-radius:999px; padding:9px 13px; background:var(--accent); color:white; text-decoration:none; font-weight:700; cursor:pointer; font-family:inherit; }
    .btn.secondary, button.secondary { background:#d8c7ac; color:var(--ink); }
    .btn.ghost { background:transparent; color:var(--ink); border:1px solid var(--line); }
    .btn.bad, button.bad { background:var(--bad); }
    .btn.warn, button.warn { background:var(--accent2); }
    .pill { display:inline-block; border:1px solid var(--line); border-radius:999px; padding:4px 9px; background:#f5eddf; margin:2px; font-size:12px; }
    .pill.ok { color:var(--ok); border-color:#9cbea4; background:#eef8ef; }
    .pill.warn { color:var(--warn); border-color:#d8bf78; background:#fff6d8; }
    .pill.bad { color:var(--bad); border-color:#dca2a2; background:#fff0ef; }
    .image-pane { position:sticky; top:86px; max-height:calc(100vh - 112px); overflow:auto; }
    .image-toolbar { background:#f6eedf; border-top:1px solid var(--line); border-bottom:1px solid var(--line); }
    .image-wrap { background:#26231f; padding:10px; overflow:auto; height:72vh; min-height:540px; cursor:grab; }
    .image-wrap img { display:block; width:100%; max-width:none; border:1px solid #514a42; background:white; transform-origin:top left; }
    .image-wrap.dragging { cursor:grabbing; user-select:none; }
    .meta { display:grid; grid-template-columns:150px minmax(0,1fr); gap:7px 12px; font-size:13px; }
    .callout { border-left:4px solid var(--accent2); background:#fff5df; padding:11px 13px; line-height:1.35; }
    .callout.ok { border-color:var(--ok); background:#eef8ef; }
    .callout.bad { border-color:var(--bad); background:#fff0ef; }
    .table-wrap { overflow:auto; max-height:46vh; background:#fffdf8; border-top:1px solid var(--line); }
    table { border-collapse:collapse; width:100%; font-size:12px; }
    th,td { border:1px solid var(--line); padding:5px 7px; vertical-align:top; text-align:left; white-space:nowrap; }
    th { background:#efe4d1; position:sticky; top:0; z-index:1; }
    td:first-child, th:first-child { position:sticky; left:0; background:#fbf4e9; z-index:2; }
    th:first-child { background:#e6d9c3; z-index:3; }
    textarea { width:100%; min-height:98px; resize:vertical; border:1px solid var(--line); border-radius:12px; padding:10px; background:#fffdf8; font:14px/1.35 Menlo, Consolas, monospace; }
    pre { white-space:pre-wrap; word-break:break-word; overflow:auto; max-height:230px; background:#fffdf8; border:1px solid var(--line); border-radius:12px; padding:12px; font:12px/1.4 Menlo, Consolas, monospace; }
    .decision-bar { position:fixed; left:0; right:0; bottom:0; z-index:5; background:rgba(255,250,240,.98); backdrop-filter:blur(10px); border-top:1px solid var(--line); padding:10px 14px; box-shadow:0 -10px 30px rgba(49,34,16,.12); }
    .decision-form { display:grid; grid-template-columns:minmax(420px,1fr) minmax(680px,1.4fr); gap:12px; align-items:start; }
    .quick-actions { display:grid; grid-template-columns:repeat(5,minmax(120px,1fr)); gap:8px; }
    .quick-actions button { border-radius:14px; min-height:58px; font-size:15px; }
    .quick-actions small { display:block; opacity:.82; font-weight:400; margin-top:3px; }
    @media (max-width:1120px) {
      body { padding-bottom:260px; }
      header { display:block; }
      main, .decision-form { grid-template-columns:1fr; }
      .image-pane { position:relative; top:0; max-height:none; }
    }
  </style>
</head>
<body>
  <header>
    <div>
      <div class="muted">Decimal / residue review | {{ index + 1 }} / {{ total }} | filter={{ filter_mode }}</div>
      <h1>{{ row.province }} {{ row.year }} | {{ row.review_bucket }} | {{ row.title }}</h1>
      <div class="mono muted">{{ row.table_id }}</div>
    </div>
    <div class="toolbar">
      <a class="btn secondary" href="{{ url_for('review_page', idx=normal_first_index, filter='flagged') }}">Back to anomaly review</a>
      {% if prev_idx is not none %}<a class="btn secondary" data-prev href="{{ url_for('decimal_review_page', idx=prev_idx, filter=filter_mode) }}">Prev</a>{% endif %}
      {% if next_idx is not none %}<a class="btn secondary" data-next href="{{ url_for('decimal_review_page', idx=next_idx, filter=filter_mode) }}">Next</a>{% endif %}
      {% for mode,label in filters %}
        <a class="btn ghost" href="{{ url_for('decimal_review_page', idx=first_index(mode), filter=mode) }}">{{ label }}</a>
      {% endfor %}
    </div>
  </header>

  <main>
    <aside class="image-pane">
      <div class="card">
        <div class="pad">
          <h2>Source Image</h2>
          <div class="muted">{{ row.source_image_local or "not local" }}</div>
        </div>
        {% if has_image %}
          <div class="pad toolbar image-toolbar">
            <button type="button" class="secondary" onclick="zoomImage(-25)">Zoom -</button>
            <button type="button" class="secondary" onclick="zoomImage(25)">Zoom +</button>
            <button type="button" class="secondary" onclick="setImageZoom(100)">Fit</button>
            <button type="button" class="secondary" onclick="setImageZoom(175)">175%</button>
            <span class="mono muted" id="zoomReadout">100%</span>
            <a class="btn secondary" href="{{ url_for('decimal_image_file', idx=index) }}" target="_blank">Open image</a>
          </div>
          <div class="image-wrap" id="imageViewport"><img id="sourceImage" src="{{ url_for('decimal_image_file', idx=index) }}" draggable="false"></div>
        {% else %}
          <div class="pad muted">No local source image found.</div>
        {% endif %}
      </div>
    </aside>

    <section>
      <div class="card">
        <div class="pad">
          <h2>Review Target</h2>
          <div>
            <span class="pill {{ bucket_class }}">{{ row.review_bucket }}</span>
            <span class="pill">{{ decision.human_decision or "unreviewed" }}</span>
            {% if row.decimal_places %}<span class="pill ok">dp={{ row.decimal_places }}</span>{% endif %}
            {% if row.validation_failures %}<span class="pill bad">failures={{ row.validation_failures }}</span>{% endif %}
            <span class="pill">{{ row.measure_category }}</span>
            <span class="pill">{{ row.format_class }}</span>
            <span class="pill">changed {{ changed_count }}</span>
          </div>
        </div>
        <div class="pad meta">
          <strong>Reason</strong><span>{{ row.reason }}</span>
          <strong>Rejected reason</strong><span>{{ row.rejected_reason }}</span>
          <strong>Opus verdict</strong><span>{{ opus_summary }}</span>
          <strong>Sample cells</strong><span>{{ opus_sample }}</span>
          <strong>Facts</strong><span>{{ original_count }} original, {{ cleaned_count }} cleaned</span>
          <strong>Raw ID</strong><span class="mono">{{ row.raw_table_id }}</span>
          <strong>Excel</strong><span class="mono">{{ row.excel_path }}</span>
        </div>
        {% if row.review_bucket == "auto_fixed" %}
          <div class="pad"><div class="callout ok"><strong>Auto-fix rule:</strong> values in rate/distribution cells were divided by 10^dp only because Opus reported uniform decimal places and its sample cells matched that dp. Counts were not touched.</div></div>
        {% elif row.review_bucket == "needs_review" %}
          <div class="pad"><div class="callout"><strong>Blocked auto-fix:</strong> this table had non-uniform columns or Opus sample mismatch. Check the image and decide whether to rerun with the decimal-preserving prompt or manually annotate the scale.</div></div>
        {% elif row.review_bucket == "validation_fail" %}
          <div class="pad"><div class="callout bad"><strong>Validation failure:</strong> this is one of the real post-cleaning constraint failures. It may need re-OCR/re-extraction, not a decimal-only correction.</div></div>
        {% endif %}
      </div>

      {% if cleaned_count == 0 and original_count > 0 %}
        <div class="card">
          <div class="pad">
            <div class="callout bad">
              <strong>Cleaned table has zero fact rows.</strong>
              The original table has {{ original_count }} facts, but the cleaned export has none for this table_id. This usually means the cleanup pass removed all recorded `persons` facts from a rate/distribution table. If the original table is valid, use <strong>Revert to original</strong>.
            </div>
          </div>
        </div>
      {% endif %}

      <div class="card">
        <div class="pad"><h2>Whole Table Before / After</h2><div class="muted">Each logical cell is shown as original and cleaned side-by-side. If cleaned is blank, that value was removed by cleanup. {{ comparison_status }}</div></div>
        <div class="table-wrap">{{ comparison_matrix_html|safe }}</div>
      </div>

      <div class="card">
        <div class="pad"><h2>Changed Values Only</h2><div class="muted">Rows where original and cleaned values differ by coordinate. Useful for auditing exact cleanup effects.</div></div>
        <div class="table-wrap">{{ changed_html|safe }}</div>
      </div>

      <div class="card">
        <div class="pad"><h2>Cleaned Facts Matrix</h2><div class="muted">{{ cleaned_status }}</div></div>
        <div class="table-wrap">{{ cleaned_matrix_html|safe }}</div>
      </div>

      <div class="card">
        <div class="pad"><h2>Original Facts Matrix</h2><div class="muted">{{ original_status }}</div></div>
        <div class="table-wrap">{{ original_matrix_html|safe }}</div>
      </div>
    </section>
  </main>

  <form class="decision-bar" method="post" action="{{ url_for('save_decimal_review_page', idx=index, filter=filter_mode) }}">
    <input type="hidden" name="quick_decision" value="">
    <div class="decision-form">
      <div>
        <label><strong>Human notes / fix focus</strong></label>
        <textarea name="human_notes" placeholder="Write what you see: decimal point correct, scale wrong, rerun prompt focus, or table should be excluded.">{{ decision.human_notes }}</textarea>
      </div>
      <div>
        <label><strong>Decision</strong></label>
        <div class="quick-actions">
          <button type="button" onclick="quickSubmit('revert_to_original')">Revert to original<small>1</small></button>
          <button type="button" onclick="quickSubmit('decimal_ok')">Cleaned OK<small>2</small></button>
          <button type="button" onclick="quickSubmit('manual_scale_fix')">Manual scale fix<small>3</small></button>
          <button class="warn" type="button" onclick="quickSubmit('rerun_decimal_prompt')">Rerun prompt<small>4</small></button>
          <button class="bad" type="button" onclick="quickSubmit('exclude_for_now')">Exclude<small>5</small></button>
        </div>
      </div>
    </div>
  </form>
  <script>
    let imageZoom = 100;
    function imageEls() {
      return {
        viewport: document.getElementById('imageViewport'),
        image: document.getElementById('sourceImage'),
        readout: document.getElementById('zoomReadout')
      };
    }
    function setImageZoom(value) {
      const els = imageEls();
      if (!els.viewport || !els.image) return;
      imageZoom = Math.max(50, Math.min(500, value));
      els.image.style.width = imageZoom + '%';
      if (els.readout) els.readout.textContent = imageZoom + '%';
    }
    function zoomImage(delta) { setImageZoom(imageZoom + delta); }
    (function initImageViewer() {
      const els = imageEls();
      if (!els.viewport || !els.image) return;
      setImageZoom(100);
      els.viewport.addEventListener('wheel', function(evt) {
        if (!(evt.ctrlKey || evt.metaKey)) return;
        evt.preventDefault();
        zoomImage(evt.deltaY < 0 ? 25 : -25);
      }, { passive:false });
      let dragging = false, startX = 0, startY = 0, startLeft = 0, startTop = 0;
      els.viewport.addEventListener('pointerdown', function(evt) {
        if (evt.button !== 0) return;
        dragging = true; startX = evt.clientX; startY = evt.clientY;
        startLeft = els.viewport.scrollLeft; startTop = els.viewport.scrollTop;
        els.viewport.classList.add('dragging'); els.viewport.setPointerCapture(evt.pointerId);
      });
      els.viewport.addEventListener('pointermove', function(evt) {
        if (!dragging) return;
        els.viewport.scrollLeft = startLeft - (evt.clientX - startX);
        els.viewport.scrollTop = startTop - (evt.clientY - startY);
      });
      function stopDrag(evt) {
        dragging = false; els.viewport.classList.remove('dragging');
        try { els.viewport.releasePointerCapture(evt.pointerId); } catch (err) {}
      }
      els.viewport.addEventListener('pointerup', stopDrag);
      els.viewport.addEventListener('pointercancel', stopDrag);
    })();
    function quickSubmit(decision) {
      const form = document.querySelector('form');
      const quick = document.querySelector('input[name="quick_decision"]');
      if (quick) quick.value = decision;
      if (form) form.submit();
    }
    document.addEventListener('keydown', function(evt) {
      if (evt.target && ['TEXTAREA','INPUT','SELECT'].includes(evt.target.tagName)) return;
      if (evt.key === 'j') { const el=document.querySelector('a[data-next]'); if (el) window.location=el.href; }
      if (evt.key === 'k') { const el=document.querySelector('a[data-prev]'); if (el) window.location=el.href; }
      if (evt.key === '1') quickSubmit('revert_to_original');
      if (evt.key === '2') quickSubmit('decimal_ok');
      if (evt.key === '3') quickSubmit('manual_scale_fix');
      if (evt.key === '4') quickSubmit('rerun_decimal_prompt');
      if (evt.key === '5') quickSubmit('exclude_for_now');
      if (evt.key === '+' || evt.key === '=') zoomImage(25);
      if (evt.key === '-') zoomImage(-25);
      if (evt.key === '0') setImageZoom(100);
    });
  </script>
</body>
</html>
"""


DECIMAL_DECISION_HTML = """
<!doctype html>
<html>
<head>
  <meta charset="utf-8">
  <title>Decimal Decision Review</title>
  <style>
    :root {
      --ink:#191715;
      --muted:#706a61;
      --bg:#ece7dc;
      --paper:#fffaf0;
      --line:#d6cbbb;
      --accent:#263d35;
      --accent2:#b0612d;
      --bad:#8f2828;
      --warn:#8a650e;
      --ok:#27603a;
    }
    * { box-sizing:border-box; }
    body { margin:0; padding-bottom:176px; color:var(--ink); background:radial-gradient(circle at top left, #fff7df 0, var(--bg) 42%, #e6ddd0 100%); font-family:"Iowan Old Style", Georgia, serif; }
    header { position:sticky; top:0; z-index:4; background:rgba(255,250,240,.97); backdrop-filter:blur(8px); border-bottom:1px solid var(--line); padding:12px 16px; display:flex; justify-content:space-between; gap:12px; align-items:flex-start; }
    h1 { font-size:20px; margin:0 0 4px 0; line-height:1.2; }
    h2 { font-size:17px; margin:0; }
    main { display:grid; grid-template-columns:minmax(520px,44vw) minmax(620px,1fr); gap:14px; padding:14px; align-items:start; }
    .card { background:var(--paper); border:1px solid var(--line); border-radius:16px; box-shadow:0 10px 28px rgba(49,34,16,.07); overflow:hidden; margin-bottom:14px; }
    .pad { padding:14px; }
    .muted { color:var(--muted); font-size:12px; }
    .mono { font-family:Menlo, Consolas, monospace; }
    .toolbar { display:flex; gap:8px; flex-wrap:wrap; align-items:center; }
    .btn, button { appearance:none; border:0; border-radius:999px; padding:9px 13px; background:var(--accent); color:white; text-decoration:none; font-weight:700; cursor:pointer; font-family:inherit; }
    .btn.secondary, button.secondary { background:#d8c7ac; color:var(--ink); }
    .btn.ghost { background:transparent; color:var(--ink); border:1px solid var(--line); }
    .btn.bad, button.bad { background:var(--bad); }
    .btn.warn, button.warn { background:var(--accent2); }
    .pill { display:inline-block; border:1px solid var(--line); border-radius:999px; padding:4px 9px; background:#f5eddf; margin:2px; font-size:12px; }
    .pill.ok { color:var(--ok); border-color:#9cbea4; background:#eef8ef; }
    .pill.warn { color:var(--warn); border-color:#d8bf78; background:#fff6d8; }
    .pill.bad { color:var(--bad); border-color:#dca2a2; background:#fff0ef; }
    .image-pane { position:sticky; top:86px; max-height:calc(100vh - 112px); overflow:auto; }
    .image-toolbar { background:#f6eedf; border-top:1px solid var(--line); border-bottom:1px solid var(--line); }
    .image-wrap { background:#26231f; padding:10px; overflow:auto; height:72vh; min-height:540px; cursor:grab; }
    .image-wrap img { display:block; width:100%; max-width:none; border:1px solid #514a42; background:white; transform-origin:top left; }
    .image-wrap.dragging { cursor:grabbing; user-select:none; }
    .meta { display:grid; grid-template-columns:150px minmax(0,1fr); gap:7px 12px; font-size:13px; }
    .callout { border-left:4px solid var(--accent2); background:#fff5df; padding:11px 13px; line-height:1.35; margin-bottom:8px; }
    .table-wrap { overflow:auto; max-height:68vh; background:#fffdf8; border-top:1px solid var(--line); }
    table { border-collapse:collapse; width:100%; font-size:12px; }
    th,td { border:1px solid var(--line); padding:5px 7px; vertical-align:top; text-align:left; white-space:nowrap; }
    th { background:#efe4d1; position:sticky; top:0; z-index:1; }
    td:first-child, th:first-child { position:sticky; left:0; background:#fbf4e9; z-index:2; }
    th:first-child { background:#e6d9c3; z-index:3; }
    textarea { width:100%; min-height:92px; resize:vertical; border:1px solid var(--line); border-radius:12px; padding:10px; background:#fffdf8; font:14px/1.35 Menlo, Consolas, monospace; }
    .decision-bar { position:fixed; left:0; right:0; bottom:0; z-index:5; background:rgba(255,250,240,.98); backdrop-filter:blur(10px); border-top:1px solid var(--line); padding:10px 14px; box-shadow:0 -10px 30px rgba(49,34,16,.12); }
    .decision-form { display:grid; grid-template-columns:minmax(380px,.8fr) minmax(760px,1.5fr); gap:12px; align-items:start; }
    .quick-actions { display:grid; grid-template-columns:repeat(6,minmax(104px,1fr)); gap:8px; }
    .quick-actions button { border-radius:14px; min-height:58px; font-size:15px; }
    .quick-actions small { display:block; opacity:.82; font-weight:400; margin-top:3px; }
    @media (max-width:1120px) {
      body { padding-bottom:280px; }
      header { display:block; }
      main, .decision-form { grid-template-columns:1fr; }
      .image-pane { position:relative; top:0; max-height:none; }
      .quick-actions { grid-template-columns:repeat(2,minmax(130px,1fr)); }
    }
  </style>
</head>
<body>
  <header>
    <div>
      <div class="muted">Decimal decision review | {{ index + 1 }} / {{ total }} | filter={{ filter_mode }}</div>
      <h1>{{ row.province }} {{ row.year }} | {{ row.category }} | {{ row.title or row.table_id }}</h1>
      <div class="mono muted">{{ row.table_id }}</div>
    </div>
    <div class="toolbar">
      <a class="btn secondary" href="{{ url_for('decimal_review_page', idx=decimal_first, filter='residue') }}">Decimal residue</a>
      {% if prev_idx is not none %}<a class="btn secondary" data-prev href="{{ url_for('decimal_decision_page', idx=prev_idx, filter=filter_mode) }}">Prev</a>{% endif %}
      {% if next_idx is not none %}<a class="btn secondary" data-next href="{{ url_for('decimal_decision_page', idx=next_idx, filter=filter_mode) }}">Next</a>{% endif %}
      {% for mode,label in filters %}
        <a class="btn ghost" href="{{ url_for('decimal_decision_page', idx=first_index(mode), filter=mode) }}">{{ label }}</a>
      {% endfor %}
    </div>
  </header>

  <main>
    <aside class="image-pane">
      <div class="card">
        <div class="pad">
          <h2>Source Image</h2>
          <div class="muted">{{ row.image }}</div>
        </div>
        {% if has_image %}
          <div class="pad toolbar image-toolbar">
            <button type="button" class="secondary" onclick="zoomImage(-25)">Zoom -</button>
            <button type="button" class="secondary" onclick="zoomImage(25)">Zoom +</button>
            <button type="button" class="secondary" onclick="setImageZoom(100)">Fit</button>
            <button type="button" class="secondary" onclick="setImageZoom(175)">175%</button>
            <span class="mono muted" id="zoomReadout">100%</span>
            <a class="btn secondary" href="{{ url_for('decimal_decision_image_file', idx=index) }}" target="_blank">Open image</a>
          </div>
          <div class="image-wrap" id="imageViewport"><img id="sourceImage" src="{{ url_for('decimal_decision_image_file', idx=index) }}" draggable="false"></div>
        {% else %}
          <div class="pad muted">No local source image found.</div>
        {% endif %}
      </div>
    </aside>

    <section>
      <div class="card">
        <div class="pad">
          <h2>Opus Prediction vs Your Call</h2>
          <div>
            <span class="pill {{ category_class }}">{{ row.category }}</span>
            <span class="pill {% if row.opus_dp in ['1','2','3','4'] %}ok{% else %}warn{% endif %}">Opus dp={{ row.opus_dp or "?" }}</span>
            <span class="pill">uniform={{ row.opus_uniform }}</span>
            <span class="pill">range {{ row.rate_min }}–{{ row.rate_max }}</span>
            <span class="pill">you: {{ decision_label }}</span>
          </div>
        </div>
        <div class="pad meta">
          <strong>Opus sample</strong><span class="mono">{{ row.opus_sample }}</span>
          <strong>Looks like</strong><span>{{ row.opus_looks_like }}</span>
          <strong>Current facts</strong><span>{{ facts_count }} rows</span>
          <strong>Image exists</strong><span>{{ has_image }}</span>
          <strong>Excel</strong><span class="mono">{{ row.excel_path }}</span>
        </div>
        <div class="pad">
          <div class="callout">
            Use `0/1/2/3+` only when the table has a clean uniform decimal scale. Use `mixed/weird` when columns or sections have different scales. Use `already correct` when the recorded facts should not be divided.
          </div>
        </div>
      </div>

      <div class="card">
        <div class="pad"><h2>Current Recorded Facts Matrix</h2><div class="muted">{{ matrix_status }}</div></div>
        <div class="table-wrap">{{ matrix_html|safe }}</div>
      </div>
    </section>
  </main>

  <form class="decision-bar" method="post" action="{{ url_for('save_decimal_decision_page', idx=index, filter=filter_mode) }}">
    <input type="hidden" name="quick_dp" value="">
    <input type="hidden" name="quick_flag" value="">
    <div class="decision-form">
      <div>
        <label><strong>Human notes</strong></label>
        <textarea name="human_notes" placeholder="Optional: exact issue, column scale, or rerun instruction.">{{ decision.human_notes }}</textarea>
      </div>
      <div>
        <label><strong>Human decimal decision</strong></label>
        <div class="quick-actions">
          <button type="button" onclick="quickSubmit('0','')">0 dp<small>1</small></button>
          <button type="button" onclick="quickSubmit('1','')">1 dp<small>2</small></button>
          <button type="button" onclick="quickSubmit('2','')">2 dp<small>3</small></button>
          <button type="button" onclick="quickSubmit('3+','')">3+ dp<small>4</small></button>
          <button class="warn" type="button" onclick="quickSubmit('', 'mixed')">Mixed/weird<small>5</small></button>
          <button class="secondary" type="button" onclick="quickSubmit('', 'already_correct')">Already correct<small>6</small></button>
        </div>
      </div>
    </div>
  </form>
  <script>
    let imageZoom = 100;
    function imageEls() {
      return {
        viewport: document.getElementById('imageViewport'),
        image: document.getElementById('sourceImage'),
        readout: document.getElementById('zoomReadout')
      };
    }
    function setImageZoom(value) {
      const els = imageEls();
      if (!els.viewport || !els.image) return;
      imageZoom = Math.max(50, Math.min(500, value));
      els.image.style.width = imageZoom + '%';
      if (els.readout) els.readout.textContent = imageZoom + '%';
    }
    function zoomImage(delta) { setImageZoom(imageZoom + delta); }
    (function initImageViewer() {
      const els = imageEls();
      if (!els.viewport || !els.image) return;
      setImageZoom(100);
      els.viewport.addEventListener('wheel', function(evt) {
        if (!(evt.ctrlKey || evt.metaKey)) return;
        evt.preventDefault();
        zoomImage(evt.deltaY < 0 ? 25 : -25);
      }, { passive:false });
      let dragging = false, startX = 0, startY = 0, startLeft = 0, startTop = 0;
      els.viewport.addEventListener('pointerdown', function(evt) {
        if (evt.button !== 0) return;
        dragging = true; startX = evt.clientX; startY = evt.clientY;
        startLeft = els.viewport.scrollLeft; startTop = els.viewport.scrollTop;
        els.viewport.classList.add('dragging'); els.viewport.setPointerCapture(evt.pointerId);
      });
      els.viewport.addEventListener('pointermove', function(evt) {
        if (!dragging) return;
        els.viewport.scrollLeft = startLeft - (evt.clientX - startX);
        els.viewport.scrollTop = startTop - (evt.clientY - startY);
      });
      function stopDrag(evt) {
        dragging = false; els.viewport.classList.remove('dragging');
        try { els.viewport.releasePointerCapture(evt.pointerId); } catch (err) {}
      }
      els.viewport.addEventListener('pointerup', stopDrag);
      els.viewport.addEventListener('pointercancel', stopDrag);
    })();
    function quickSubmit(dp, flag) {
      const form = document.querySelector('form');
      document.querySelector('input[name="quick_dp"]').value = dp;
      document.querySelector('input[name="quick_flag"]').value = flag;
      if (form) form.submit();
    }
    document.addEventListener('keydown', function(evt) {
      if (evt.target && ['TEXTAREA','INPUT','SELECT'].includes(evt.target.tagName)) return;
      if (evt.key === 'j') { const el=document.querySelector('a[data-next]'); if (el) window.location=el.href; }
      if (evt.key === 'k') { const el=document.querySelector('a[data-prev]'); if (el) window.location=el.href; }
      if (evt.key === '1') quickSubmit('0','');
      if (evt.key === '2') quickSubmit('1','');
      if (evt.key === '3') quickSubmit('2','');
      if (evt.key === '4') quickSubmit('3+','');
      if (evt.key === '5') quickSubmit('', 'mixed');
      if (evt.key === '6') quickSubmit('', 'already_correct');
      if (evt.key === '+' || evt.key === '=') zoomImage(25);
      if (evt.key === '-') zoomImage(-25);
      if (evt.key === '0') setImageZoom(100);
    });
  </script>
</body>
</html>
"""


RECONCILIATION_REVIEW_HTML = """
<!doctype html>
<html>
<head>
  <meta charset="utf-8">
  <title>Reconciliation Highlight Review</title>
  <style>
    :root {
      --ink:#191715;
      --muted:#706a61;
      --bg:#ece7dc;
      --paper:#fffaf0;
      --line:#d6cbbb;
      --accent:#263d35;
      --accent2:#b0612d;
      --bad:#8f2828;
      --warn:#8a650e;
      --ok:#27603a;
    }
    * { box-sizing:border-box; }
    body { margin:0; padding-bottom:168px; color:var(--ink); background:radial-gradient(circle at top left, #fff7df 0, var(--bg) 42%, #e6ddd0 100%); font-family:"Iowan Old Style", Georgia, serif; }
    header { position:sticky; top:0; z-index:4; background:rgba(255,250,240,.97); backdrop-filter:blur(8px); border-bottom:1px solid var(--line); padding:12px 16px; display:flex; justify-content:space-between; gap:12px; align-items:flex-start; }
    h1 { font-size:20px; margin:0 0 4px 0; line-height:1.2; }
    h2 { font-size:17px; margin:0; }
    main { display:grid; grid-template-columns:minmax(500px,42vw) minmax(640px,1fr); gap:14px; padding:14px; align-items:start; }
    .card { background:var(--paper); border:1px solid var(--line); border-radius:16px; box-shadow:0 10px 28px rgba(49,34,16,.07); overflow:hidden; margin-bottom:14px; }
    .pad { padding:14px; }
    .muted { color:var(--muted); font-size:12px; }
    .mono { font-family:Menlo, Consolas, monospace; }
    .toolbar { display:flex; gap:8px; flex-wrap:wrap; align-items:center; }
    .btn, button { appearance:none; border:0; border-radius:999px; padding:9px 13px; background:var(--accent); color:white; text-decoration:none; font-weight:700; cursor:pointer; font-family:inherit; }
    .btn.secondary, button.secondary { background:#d8c7ac; color:var(--ink); }
    .btn.ghost { background:transparent; color:var(--ink); border:1px solid var(--line); }
    .btn.bad, button.bad { background:var(--bad); }
    .btn.warn, button.warn { background:var(--accent2); }
    .pill { display:inline-block; border:1px solid var(--line); border-radius:999px; padding:4px 9px; background:#f5eddf; margin:2px; font-size:12px; }
    .pill.ok { color:var(--ok); border-color:#9cbea4; background:#eef8ef; }
    .pill.warn { color:var(--warn); border-color:#d8bf78; background:#fff6d8; }
    .pill.bad { color:var(--bad); border-color:#dca2a2; background:#fff0ef; }
    .image-pane { position:sticky; top:86px; max-height:calc(100vh - 112px); overflow:auto; }
    .image-toolbar { background:#f6eedf; border-top:1px solid var(--line); border-bottom:1px solid var(--line); }
    .image-wrap { background:#26231f; padding:10px; overflow:auto; height:72vh; min-height:540px; cursor:grab; }
    .image-wrap img { display:block; width:100%; max-width:none; border:1px solid #514a42; background:white; transform-origin:top left; }
    .image-wrap.dragging { cursor:grabbing; user-select:none; }
    .meta { display:grid; grid-template-columns:150px minmax(0,1fr); gap:7px 12px; font-size:13px; }
    .callout { border-left:4px solid var(--accent2); background:#fff5df; padding:11px 13px; line-height:1.35; margin-bottom:8px; }
    .callout.ok { border-color:var(--ok); background:#eef8ef; }
    .callout.bad { border-color:var(--bad); background:#fff0ef; }
    .table-wrap { overflow:auto; max-height:58vh; background:#fffdf8; border-top:1px solid var(--line); }
    table { border-collapse:collapse; width:100%; font-size:12px; }
    th,td { border:1px solid var(--line); padding:5px 7px; vertical-align:top; text-align:left; white-space:nowrap; }
    th { background:#efe4d1; position:sticky; top:0; z-index:1; }
    td:first-child, th:first-child { position:sticky; left:0; background:#fbf4e9; z-index:2; }
    th:first-child { background:#e6d9c3; z-index:3; }
    td.recon-col { background:#fff0b7; font-weight:700; }
    td.recon-total { background:#ffe0df; color:#8f2828; font-weight:800; outline:2px solid #8f2828; }
    th.recon-col { background:#f0ca62; color:#4e3300; }
    textarea { width:100%; min-height:98px; resize:vertical; border:1px solid var(--line); border-radius:12px; padding:10px; background:#fffdf8; font:14px/1.35 Menlo, Consolas, monospace; }
    .decision-bar { position:fixed; left:0; right:0; bottom:0; z-index:5; background:rgba(255,250,240,.98); backdrop-filter:blur(10px); border-top:1px solid var(--line); padding:10px 14px; box-shadow:0 -10px 30px rgba(49,34,16,.12); }
    .decision-form { display:grid; grid-template-columns:minmax(420px,1fr) minmax(620px,1.2fr); gap:12px; align-items:start; }
    .quick-actions { display:grid; grid-template-columns:repeat(4,minmax(130px,1fr)); gap:8px; }
    .quick-actions button { border-radius:14px; min-height:58px; font-size:15px; }
    .quick-actions small { display:block; opacity:.82; font-weight:400; margin-top:3px; }
    @media (max-width:1120px) {
      body { padding-bottom:240px; }
      header { display:block; }
      main, .decision-form { grid-template-columns:1fr; }
      .image-pane { position:relative; top:0; max-height:none; }
    }
  </style>
</head>
<body>
  <header>
    <div>
      <div class="muted">Reconciliation highlight review | {{ index + 1 }} / {{ total }} | filter={{ filter_mode }}</div>
      <h1>{{ row.province }} {{ row.year }} | {{ row.title }}</h1>
      <div class="mono muted">{{ row.table_id }}</div>
    </div>
    <div class="toolbar">
      <a class="btn secondary" href="{{ url_for('decimal_review_page', idx=decimal_first, filter='residue') }}">Decimal review</a>
      {% if prev_idx is not none %}<a class="btn secondary" data-prev href="{{ url_for('reconciliation_review_page', idx=prev_idx, filter=filter_mode) }}">Prev</a>{% endif %}
      {% if next_idx is not none %}<a class="btn secondary" data-next href="{{ url_for('reconciliation_review_page', idx=next_idx, filter=filter_mode) }}">Next</a>{% endif %}
      {% for mode,label in filters %}
        <a class="btn ghost" href="{{ url_for('reconciliation_review_page', idx=first_index(mode), filter=mode) }}">{{ label }}</a>
      {% endfor %}
    </div>
  </header>

  <main>
    <aside class="image-pane">
      <div class="card">
        <div class="pad">
          <h2>Source Image</h2>
          <div class="muted">{{ row.image or row.source_image_local or "not local" }}</div>
        </div>
        {% if has_image %}
          <div class="pad toolbar image-toolbar">
            <button type="button" class="secondary" onclick="zoomImage(-25)">Zoom -</button>
            <button type="button" class="secondary" onclick="zoomImage(25)">Zoom +</button>
            <button type="button" class="secondary" onclick="setImageZoom(100)">Fit</button>
            <button type="button" class="secondary" onclick="setImageZoom(175)">175%</button>
            <span class="mono muted" id="zoomReadout">100%</span>
            <a class="btn secondary" href="{{ url_for('reconciliation_image_file', idx=index) }}" target="_blank">Open image</a>
          </div>
          <div class="image-wrap" id="imageViewport"><img id="sourceImage" src="{{ url_for('reconciliation_image_file', idx=index) }}" draggable="false"></div>
        {% else %}
          <div class="pad muted">No local source image found.</div>
        {% endif %}
      </div>
    </aside>

    <section>
      <div class="card">
        <div class="pad">
          <h2>Failure Banners</h2>
          <div>
            <span class="pill bad">{{ row.n_failures }} failures</span>
            <span class="pill">{{ decision.human_decision or "unreviewed" }}</span>
            <span class="pill">{{ row.table_family }}</span>
            <span class="pill">{{ row.format_class }}</span>
          </div>
        </div>
        <div class="pad">
          {% for failure in failures %}
            <div class="callout bad">
              <strong>{{ failure.level }}:</strong>
              {{ failure.section }} · {{ failure.group }} · {{ failure.col }}:
              {{ failure.detail }}{% if failure.diff is not none %} (off by {{ failure.diff }}){% endif %}
            </div>
          {% endfor %}
          <div class="muted">These are reconciliation flags, not automatic errors. Large under-sums can be dropped rows/subtotals/reconstruction structure, while small diffs can be digit misreads.</div>
        </div>
        <div class="pad meta">
          <strong>Facts</strong><span>{{ original_count }} original, {{ cleaned_count }} cleaned</span>
          <strong>Excel</strong><span class="mono">{{ row.excel_path }}</span>
          <strong>Image</strong><span class="mono">{{ row.image }}</span>
        </div>
      </div>

      <div class="card">
        <div class="pad"><h2>Highlighted Whole Table</h2><div class="muted">Amber columns match failure section/group/sub-column. Red badges are Total/grand-total row cells in those columns. {{ matrix_status }}</div></div>
        <div class="table-wrap">{{ matrix_html|safe }}</div>
      </div>
    </section>
  </main>

  <form class="decision-bar" method="post" action="{{ url_for('save_reconciliation_review_page', idx=index, filter=filter_mode) }}">
    <input type="hidden" name="quick_decision" value="">
    <div class="decision-form">
      <div>
        <label><strong>Human notes / correction</strong></label>
        <textarea name="human_notes" placeholder="If real digit error, write exact correction. If structural/subtotal artifact, note that and accept.">{{ decision.human_notes }}</textarea>
      </div>
      <div>
        <label><strong>Decision</strong></label>
        <div class="quick-actions">
          <button type="button" onclick="quickSubmit('accept_artifact')">Accept artifact<small>1</small></button>
          <button type="button" onclick="quickSubmit('fix_cell')">Fix-cell<small>2</small></button>
          <button class="warn" type="button" onclick="quickSubmit('rerun_extract')">Rerun extract<small>3</small></button>
          <button class="bad" type="button" onclick="quickSubmit('exclude_for_now')">Exclude<small>4</small></button>
        </div>
      </div>
    </div>
  </form>
  <script>
    let imageZoom = 100;
    function imageEls() {
      return {
        viewport: document.getElementById('imageViewport'),
        image: document.getElementById('sourceImage'),
        readout: document.getElementById('zoomReadout')
      };
    }
    function setImageZoom(value) {
      const els = imageEls();
      if (!els.viewport || !els.image) return;
      imageZoom = Math.max(50, Math.min(500, value));
      els.image.style.width = imageZoom + '%';
      if (els.readout) els.readout.textContent = imageZoom + '%';
    }
    function zoomImage(delta) { setImageZoom(imageZoom + delta); }
    (function initImageViewer() {
      const els = imageEls();
      if (!els.viewport || !els.image) return;
      setImageZoom(100);
      els.viewport.addEventListener('wheel', function(evt) {
        if (!(evt.ctrlKey || evt.metaKey)) return;
        evt.preventDefault();
        zoomImage(evt.deltaY < 0 ? 25 : -25);
      }, { passive:false });
      let dragging = false, startX = 0, startY = 0, startLeft = 0, startTop = 0;
      els.viewport.addEventListener('pointerdown', function(evt) {
        if (evt.button !== 0) return;
        dragging = true; startX = evt.clientX; startY = evt.clientY;
        startLeft = els.viewport.scrollLeft; startTop = els.viewport.scrollTop;
        els.viewport.classList.add('dragging'); els.viewport.setPointerCapture(evt.pointerId);
      });
      els.viewport.addEventListener('pointermove', function(evt) {
        if (!dragging) return;
        els.viewport.scrollLeft = startLeft - (evt.clientX - startX);
        els.viewport.scrollTop = startTop - (evt.clientY - startY);
      });
      function stopDrag(evt) {
        dragging = false; els.viewport.classList.remove('dragging');
        try { els.viewport.releasePointerCapture(evt.pointerId); } catch (err) {}
      }
      els.viewport.addEventListener('pointerup', stopDrag);
      els.viewport.addEventListener('pointercancel', stopDrag);
    })();
    function quickSubmit(decision) {
      const form = document.querySelector('form');
      const quick = document.querySelector('input[name="quick_decision"]');
      if (quick) quick.value = decision;
      if (form) form.submit();
    }
    document.addEventListener('keydown', function(evt) {
      if (evt.target && ['TEXTAREA','INPUT','SELECT'].includes(evt.target.tagName)) return;
      if (evt.key === 'j') { const el=document.querySelector('a[data-next]'); if (el) window.location=el.href; }
      if (evt.key === 'k') { const el=document.querySelector('a[data-prev]'); if (el) window.location=el.href; }
      if (evt.key === '1') quickSubmit('accept_artifact');
      if (evt.key === '2') quickSubmit('fix_cell');
      if (evt.key === '3') quickSubmit('rerun_extract');
      if (evt.key === '4') quickSubmit('exclude_for_now');
      if (evt.key === '+' || evt.key === '=') zoomImage(25);
      if (evt.key === '-') zoomImage(-25);
      if (evt.key === '0') setImageZoom(100);
    });
  </script>
</body>
</html>
"""


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def write_csv(path: Path, rows: list[dict[str, str]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def load_packets(path: Path) -> dict[str, dict[str, Any]]:
    packets = {}
    if not path.exists():
        return packets
    with path.open(encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            packet = json.loads(line)
            packets[packet.get("table_id", "")] = packet
    return packets


def by_id(rows: list[dict[str, str]], key: str = "table_id") -> dict[str, dict[str, str]]:
    return {row.get(key, ""): row for row in rows if row.get(key)}


def format_sanity_instruction(row: dict[str, str], focus_area: str = "") -> str:
    title = (row.get("title") or "").lower()
    fmt = (row.get("format_class") or "").lower()
    measure = (row.get("measure_category") or "").lower()
    denom = (row.get("denominator_hint") or "").lower()
    focus = (focus_area or "").lower()
    text = " ".join([title, fmt, measure, denom, focus])

    if (
        fmt == "distribution_of_1000"
        or measure == "distribution_of_1000"
        or "per_1000" in measure
        or any(token in text for token in ["per 1,000", "per 1000", "per 1k", "per 1 k", "per mille", "of 1,000", "of 1000", "of 1k"])
    ):
        return (
            "Format sanity: if this is a bounded distribution/proportion per 1,000, "
            "recorded cell values must be 0-1,000. A 10,000-scale value usually means "
            "a decimal, digit grouping, row alignment, or column alignment error. "
            "If the source is instead a sex ratio or other rate per 1,000, preserve that unit "
            "but still re-check any order-of-magnitude outlier."
        )
    if fmt == "distribution_of_10000" or any(token in text for token in ["per 10,000", "per 10000", "of 10,000", "of 10000"]):
        return (
            "Format sanity: if this is a bounded distribution per 10,000, recorded cell values "
            "must be 0-10,000. Do not treat those cells as raw population counts."
        )
    if any(token in text for token in ["percent", "per cent", "percentage", "%"]):
        return (
            "Format sanity: if this is a percent/per-cent table, values normally must be 0-100. "
            "Re-check any count-sized value against the source image. Preserve printed decimal points "
            "and decimal values; do not round percentages to integers."
        )
    if any(token in text for token in ["mortality", "death", "deaths", "birth", "births"]):
        return (
            "Format sanity: for mortality/birth/death tables, identify whether cells are raw events, "
            "rates, percentages, or ratios before extracting. Preserve decimal points/values and do not force rate "
            "tables into additive population-count constraints."
        )
    return (
        "Format sanity: infer the unit from the printed title and headers before extracting. "
        "Counts, per-1,000 distributions, per-10,000 distributions, percentages, rates, ratios, "
        "and indexes have different valid ranges and constraints. If the source prints decimal places, "
        "preserve the decimal value instead of rounding."
    )


def multi_table_instruction(focus_area: str) -> str:
    text = (focus_area or "").lower()
    if not any(token in text for token in ["multiple table", "multiple printed", "subtable", "split", "statement no", "statement number"]):
        return ""
    return (
        "Multi-table page instruction: extract every printed table visible on the image, "
        "but keep each printed table/subtable as a separate logical table or section with "
        "its own exact title/statement number, unit/base, rows, and columns. Do not merge "
        "values across printed subtables, and do not drop secondary subtables."
    )


def build_fix_instruction(row: dict[str, str], focus_area: str) -> str:
    focus = focus_area.strip() or "Use the human decision and Gemini reason to focus the rerun."
    split_instruction = multi_table_instruction(focus_area)
    return (
        f"Rerun or repair this extraction with the following human focus area: {focus}\n"
        f"{format_sanity_instruction(row, focus_area)}\n"
        + (f"{split_instruction}\n" if split_instruction else "")
        + f"Compare against the source image and preserve the printed table axes exactly. "
        f"Do not force this table into the current family if the title/headers imply a different format.\n"
        f"Current Gemini concern: {row.get('reason', '')}\n"
        f"Title: {row.get('title', '')}"
    )


def write_rerun_focus_queue(rows: list[dict[str, str]], decisions_csv: Path, queue_csv: Path) -> None:
    decisions = by_id(read_csv(decisions_csv))
    queue_rows = []
    rerunish = {"rerun", "reject", "hold", "image_compared_reject"}
    for row in rows:
        decision = decisions.get(row.get("table_id", ""), {})
        human_decision = decision.get("human_decision", "")
        focus_area = decision.get("rerun_focus_area", "").strip()
        if human_decision not in rerunish:
            continue
        queue_rows.append(
            {
                "table_id": row.get("table_id", ""),
                "human_decision": human_decision,
                "rerun_focus_area": focus_area,
                "llm_fix_instruction": build_fix_instruction(row, focus_area),
                "llm_decision": row.get("llm_decision", ""),
                "llm_reason": row.get("reason", ""),
                "llm_flags": row.get("flags", ""),
                "source_image_local": row.get("source_image_local", ""),
                "excel_path": row.get("excel_path", ""),
                "province": row.get("province", ""),
                "year": row.get("year", ""),
                "table_family": row.get("table_family", ""),
                "measure_category": row.get("measure_category", ""),
                "format_class": row.get("format_class", ""),
                "title": row.get("title", ""),
            }
        )
    write_csv(queue_csv, queue_rows, RERUN_FOCUS_FIELDS)


def resolve_local_path(value: str, base: Path | None = None) -> Path | None:
    if not value:
        return None
    p = Path(value)
    candidates = [p] if p.is_absolute() else []
    if base is not None and not p.is_absolute():
        candidates.append(base / p)
    if not p.is_absolute():
        candidates.append(Path(value))
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def resolve_excel(row: dict[str, str]) -> Path | None:
    return resolve_local_path(row.get("excel_path", ""), SENDOUT)


def resolve_image(row: dict[str, str]) -> Path | None:
    value = row.get("source_image_local", "") or row.get("source_image", "")
    path = resolve_local_path(value)
    if path:
        return path
    if "/census_toc/digitized_tables/" in value:
        rel = value.split("/census_toc/digitized_tables/", 1)[1]
        candidate = DIGITIZED_TABLES_REFERENCED / rel
        if candidate.exists():
            return candidate
    if "/digitized_tables/" in value:
        rel = value.split("/digitized_tables/", 1)[1]
        candidate = DIGITIZED_TABLES_REFERENCED / rel
        if candidate.exists():
            return candidate
    return None


def html_table(rows: list[list[Any]], max_rows: int = 90, max_cols: int = 60) -> str:
    if not rows:
        return "<div class='pad muted'>No rows available.</div>"
    rows = rows[:max_rows]
    rows = [row[:max_cols] for row in rows]
    out = ["<table>"]
    header = rows[0]
    out.append("<tr>" + "".join(f"<th>{html.escape(str(cell or ''))}</th>" for cell in header) + "</tr>")
    for row in rows[1:]:
        out.append("<tr>" + "".join(f"<td>{html.escape(str(cell or ''))}</td>" for cell in row) + "</tr>")
    out.append("</table>")
    return "\n".join(out)


def int_value(value: Any) -> int:
    try:
        return int(float(str(value or "").strip()))
    except Exception:
        return 0


def bool_value(value: Any) -> bool:
    return str(value or "").strip().lower() in {"1", "true", "yes", "y"}


def stable_assembled_table_id(group_id: str) -> str:
    digest = hashlib.sha1(group_id.encode("utf-8")).hexdigest()[:16]
    return f"assembled::{digest}"


def load_assembled_groups(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    groups = []
    with path.open(encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if line:
                groups.append(json.loads(line))
    return groups


def assembled_sort_key(group: dict[str, Any]) -> tuple[str, str, str, str, str]:
    return (
        str(group.get("province") or ""),
        str(group.get("year") or ""),
        str(group.get("table_family") or ""),
        str(group.get("title") or ""),
        str(group.get("group_id") or ""),
    )


def assembled_report(group: dict[str, Any]) -> dict[str, Any]:
    report = dict(group.get("assembled_constraint_report") or {})
    report.setdefault("total_checks", 0)
    report.setdefault("failed", 0)
    report.setdefault("all_passed", False)
    return report


def source_json_keys(value: Any) -> set[str]:
    text = str(value or "").strip()
    if not text:
        return set()
    keys = {text}
    path = Path(text)
    name = path.name
    stem = path.stem
    if name:
        keys.add(name)
    if stem:
        keys.add(stem)
        if stem.endswith("_oneshot"):
            keys.add(stem[: -len("_oneshot")])
        if "__" in stem:
            tail = stem.rsplit("__", 1)[1]
            keys.add(tail)
            if tail.endswith("_oneshot"):
                keys.add(tail[: -len("_oneshot")])
    for marker in ["/ocr_results/digitized_tables/", "/digitized_tables/"]:
        if marker in text:
            keys.add(text.split(marker, 1)[1])
    return {key for key in keys if key}


def source_json_match(left: Any, right: Any) -> bool:
    return bool(source_json_keys(left) & source_json_keys(right))


def image_path_from_source_json(source_json: Any) -> str:
    text = str(source_json or "").strip()
    if not text:
        return ""
    if CLUSTER_OCR_PREFIX in text:
        rel = text.split(CLUSTER_OCR_PREFIX, 1)[1]
        path = Path(rel)
        image_name = path.name.replace("_oneshot.json", ".png").replace(".json", ".png")
        return CLUSTER_DIGITIZED_PREFIX + str(path.parent / image_name)
    if "/ocr_results/digitized_tables/" in text:
        rel = text.split("/ocr_results/digitized_tables/", 1)[1]
        path = Path(rel)
        image_name = path.name.replace("_oneshot.json", ".png").replace(".json", ".png")
        return CLUSTER_DIGITIZED_PREFIX + str(path.parent / image_name)
    if text.endswith(".json") and "__" in Path(text).stem:
        stem = Path(text).stem
        dir_name, image_stem = stem.rsplit("__", 1)
        image_name = image_stem.replace("_oneshot", "") + ".png"
        return str(Path(LOCAL_DIGITIZED_PREFIX) / dir_name / image_name)
    return ""


def table_dir_name_from_path(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    for marker in ["/ocr_results/digitized_tables/", "/digitized_tables/"]:
        if marker in text:
            rel = text.split(marker, 1)[1]
            return Path(rel).parts[0] if Path(rel).parts else ""
    parts = Path(text).parts
    for idx, part in enumerate(parts):
        if part in {"digitized_tables", "ocr_results_referenced"} and idx + 1 < len(parts):
            if part == "ocr_results_referenced" and idx + 2 < len(parts) and parts[idx + 1] == "digitized_tables":
                return parts[idx + 2]
            return parts[idx + 1]
    if len(parts) >= 2:
        return parts[-2]
    return ""


def table_dir_doc_key(dir_name: str) -> str:
    match = re.match(r"(.+?)_t\d{4}_", dir_name or "")
    return match.group(1) if match else ""


def table_dir_index(dir_name: str) -> int:
    match = re.search(r"_t(\d{4})_", dir_name or "")
    return int(match.group(1)) if match else 999999


def read_local_ocr_title(path: Path) -> str:
    try:
        obj = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return ""
    metadata = obj.get("metadata") if isinstance(obj, dict) else {}
    return str((metadata or {}).get("title") or (metadata or {}).get("table_title") or "").strip()


def previous_table_context(anchor: dict[str, Any], max_images: int = 8) -> dict[str, Any]:
    anchor_dir = table_dir_name_from_path(anchor.get("source_json") or anchor.get("source_image"))
    doc_key = table_dir_doc_key(anchor_dir)
    current_index = table_dir_index(anchor_dir)
    if not doc_key or current_index == 999999 or not OCR_RESULTS_REFERENCED.exists():
        return {"dir_name": "", "images": []}

    candidates = []
    for path in OCR_RESULTS_REFERENCED.iterdir():
        if not path.is_dir():
            continue
        dir_name = path.name
        if table_dir_doc_key(dir_name) != doc_key:
            continue
        idx = table_dir_index(dir_name)
        if idx < current_index:
            candidates.append((idx, dir_name, path))
    if not candidates:
        return {"dir_name": "", "images": []}

    _, dir_name, ocr_dir = max(candidates, key=lambda item: item[0])
    digitized_dir = DIGITIZED_TABLES_REFERENCED / dir_name
    images = []
    for json_path in sorted(ocr_dir.glob("*.json"))[:max_images]:
        image_name = json_path.name.replace("_oneshot.json", ".png").replace(".json", ".png")
        image_path = digitized_dir / image_name
        images.append(
            {
                "view_index": len(images),
                "title": read_local_ocr_title(json_path),
                "source_json": str(json_path),
                "source_image": str(image_path),
                "local_path": str(image_path) if image_path.exists() else "",
                "exists": image_path.exists(),
            }
        )
    return {"dir_name": dir_name, "images": images}


def assembled_decision_map(path: Path) -> dict[str, dict[str, str]]:
    return by_id(read_csv(path), "group_id")


def flatten_assembled_section_rows(rows: list[dict[str, Any]]) -> list[list[Any]]:
    scalar_keys: list[str] = []
    value_keys: list[tuple[str, str]] = []
    scalar_seen = set()
    value_seen = set()

    for row in rows:
        for key, value in row.items():
            if isinstance(value, dict):
                for subkey in value:
                    pair = (str(key), str(subkey))
                    if pair not in value_seen:
                        value_seen.add(pair)
                        value_keys.append(pair)
            elif key not in scalar_seen:
                scalar_seen.add(key)
                scalar_keys.append(str(key))

    preferred_candidates = [
        "row_label",
        "age",
        "AGE",
        "Age",
        "RELIGION",
        "Religion",
        "Caste",
        "CASTE",
        "PROVINCE OR STATE.",
        "District",
        "__page__",
    ]
    preferred = [key for key in preferred_candidates if key in scalar_keys]
    scalar_keys = preferred + [key for key in scalar_keys if key not in preferred]

    header = scalar_keys + [f"{key} | {subkey}" for key, subkey in value_keys]
    output = [header]
    for row in rows:
        output.append(
            [row.get(key, "") for key in scalar_keys]
            + [
                (row.get(key) or {}).get(subkey, "") if isinstance(row.get(key), dict) else ""
                for key, subkey in value_keys
            ]
        )
    return output


def assembled_section_views(group: dict[str, Any]) -> list[dict[str, Any]]:
    views = []
    for idx, section in enumerate(group.get("sections") or [], start=1):
        stats = section.get("stats") or {}
        matrix = section.get("matrix") or flatten_assembled_section_rows(section.get("rows") or [])
        row_count = max(len(matrix) - 1, 0)
        col_count = len(matrix[0]) if matrix else 0
        views.append(
            {
                "index": idx,
                "name": section.get("name") or f"section_{idx}",
                "merge_strategy": section.get("merge_strategy", ""),
                "input_pages": stats.get("input_pages", ""),
                "output_rows": stats.get("output_rows", row_count),
                "shape": f"{row_count} rows x {col_count} cols",
                "html": html_table(matrix, max_rows=2000, max_cols=320),
            }
        )
    return views


def assembled_anchor_for_group(group: dict[str, Any], continuation_by_group_id: dict[str, list[dict[str, str]]]) -> dict[str, Any]:
    group_id = group.get("group_id", "")
    source_jsons = group.get("page_source_jsons") or []
    source_images = group.get("page_source_images") or []
    continuation_rows = continuation_by_group_id.get(group_id, [])
    candidates: list[dict[str, Any]] = []
    for row in continuation_rows:
        source_json = row.get("candidate_source_json") or row.get("auto_attached_to_source_json") or ""
        if not source_json:
            continue
        candidates.append(
            {
                "source_json": source_json,
                "title": row.get("candidate_title", ""),
                "reason": row.get("auto_attach_reason", ""),
                "is_member": any(source_json_match(source_json, member_json) for member_json in source_jsons),
                "score": 0,
            }
        )

    if not candidates and source_jsons:
        candidates.append(
            {
                "source_json": source_jsons[0],
                "title": group.get("title", ""),
                "reason": "first_ordered_member_page",
                "is_member": True,
                "score": 0,
            }
        )

    continuation_tokens = ("continued", "contd", "cont.", "concluded", "concld", "concld.", "concl.")
    for candidate in candidates:
        title = str(candidate.get("title") or "").lower()
        score = 0
        if candidate.get("is_member"):
            score += 50
        if title and not any(token in title for token in continuation_tokens):
            score += 25
        if candidate.get("reason"):
            score += 5
        candidate["score"] = score

    anchor = max(candidates, key=lambda item: item.get("score", 0), default={})
    anchor_source = anchor.get("source_json", "")
    anchor_image = ""
    for idx, member_json in enumerate(source_jsons):
        if source_json_match(anchor_source, member_json):
            anchor_source = member_json
            anchor_image = source_images[idx] if idx < len(source_images) else ""
            break
    if not anchor_image:
        anchor_image = image_path_from_source_json(anchor_source)
    return {
        "source_json": anchor_source,
        "source_image": anchor_image,
        "title": anchor.get("title") or group.get("title", ""),
        "reason": anchor.get("reason") or "first_ordered_member_page",
        "is_member": bool(anchor.get("is_member", True)),
    }


def assembled_member_images(group: dict[str, Any], anchor: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    source_images = group.get("page_source_images") or []
    source_jsons = group.get("page_source_jsons") or []
    max_len = max(len(source_images), len(source_jsons), 0)
    rows = []
    anchor = anchor or {}
    anchor_source = anchor.get("source_json", "")
    for idx in range(max_len):
        source_image = source_images[idx] if idx < len(source_images) else ""
        source_json = source_jsons[idx] if idx < len(source_jsons) else ""
        path = resolve_image({"source_image": source_image, "source_image_local": source_image})
        is_anchor = source_json_match(anchor_source, source_json)
        rows.append(
            {
                "member_index0": idx,
                "member_index": idx + 1,
                "source_image": source_image,
                "source_json": source_json,
                "local_path": str(path) if path else "",
                "exists": bool(path),
                "is_anchor": is_anchor,
                "reference_only": False,
            }
        )
    if anchor_source and not any(row["is_anchor"] for row in rows):
        source_image = anchor.get("source_image", "") or image_path_from_source_json(anchor_source)
        path = resolve_image({"source_image": source_image, "source_image_local": source_image})
        rows.insert(
            0,
            {
                "member_index0": -1,
                "member_index": 0,
                "source_image": source_image,
                "source_json": anchor_source,
                "local_path": str(path) if path else "",
                "exists": bool(path),
                "is_anchor": True,
                "reference_only": True,
            },
        )
    rows.sort(key=lambda row: (not row["is_anchor"], row["member_index0"] if row["member_index0"] >= 0 else -1))
    for view_idx, row in enumerate(rows):
        row["view_index"] = view_idx
    return rows


def assembled_review_index_row(group: dict[str, Any], decision: dict[str, str]) -> dict[str, Any]:
    report = assembled_report(group)
    review_decision = decision.get("review_decision", "")
    return {
        "group_id": group.get("group_id", ""),
        "table_id": stable_assembled_table_id(group.get("group_id", "")),
        "review_decision": review_decision,
        "review_notes": decision.get("review_notes", ""),
        "province": group.get("province", ""),
        "year": group.get("year", ""),
        "table_family": group.get("table_family", ""),
        "title": group.get("title", ""),
        "n_pages": group.get("n_pages", ""),
        "overall_strategy": group.get("overall_strategy", ""),
        "assembled_constraint_checks": report.get("total_checks", ""),
        "assembled_constraint_failures": report.get("failed", ""),
        "assembled_all_passed": report.get("all_passed", ""),
        "use_as_compiled_table": "1" if review_decision == "keep" else "0",
    }


def assembled_review_payload(group: dict[str, Any], decision: dict[str, str]) -> dict[str, Any]:
    payload = dict(group)
    payload["review_decision"] = decision.get("review_decision", "")
    payload["review_notes"] = decision.get("review_notes", "")
    payload["reviewer"] = decision.get("reviewer", "")
    payload["reviewed_at"] = decision.get("reviewed_at", "")
    payload["table_id"] = stable_assembled_table_id(group.get("group_id", ""))
    payload["use_as_compiled_table"] = decision.get("review_decision") == "keep"
    return payload


def write_assembled_review_exports(groups: list[dict[str, Any]], decisions_csv: Path, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    decisions = assembled_decision_map(decisions_csv)
    index_rows = []
    accepted = []
    rejected = []
    rerun = []
    grouping_ok_rerun_facts = []
    hold = []
    replacement_rows = []
    accepted_member_source_jsons = []
    for group in groups:
        group_id = group.get("group_id", "")
        decision = decisions.get(group_id, {})
        index_rows.append(assembled_review_index_row(group, decision))
        review_decision = decision.get("review_decision", "")
        payload = assembled_review_payload(group, decision)
        if review_decision == "keep":
            accepted.append(payload)
            table_id = stable_assembled_table_id(group_id)
            source_jsons = group.get("page_source_jsons") or []
            source_images = group.get("page_source_images") or []
            for idx, source_json in enumerate(source_jsons):
                accepted_member_source_jsons.append(source_json)
                replacement_rows.append(
                    {
                        "member_source_json": source_json,
                        "member_source_image": source_images[idx] if idx < len(source_images) else "",
                        "assembled_group_id": group_id,
                        "assembled_table_id": table_id,
                        "province": group.get("province", ""),
                        "year": group.get("year", ""),
                        "table_family": group.get("table_family", ""),
                        "title": group.get("title", ""),
                    }
                )
        elif review_decision == "drop":
            rejected.append(payload)
        elif review_decision == "rerun":
            rerun.append(payload)
        elif review_decision == "grouping_ok_rerun_facts":
            grouping_ok_rerun_facts.append(payload)
            rerun.append(payload)
        elif review_decision == "hold":
            hold.append(payload)

    write_csv(output_dir / "assembled_group_review_index.csv", index_rows, ASSEMBLED_REVIEW_INDEX_FIELDS)
    for name, payloads in [
        ("accepted_assembled_groups.jsonl", accepted),
        ("rejected_assembled_groups.jsonl", rejected),
        ("rerun_assembled_groups.jsonl", rerun),
        ("grouping_ok_rerun_fact_groups.jsonl", grouping_ok_rerun_facts),
        ("hold_assembled_groups.jsonl", hold),
    ]:
        with (output_dir / name).open("w", encoding="utf-8") as fh:
            for payload in payloads:
                fh.write(json.dumps(payload, ensure_ascii=False) + "\n")
    write_csv(
        output_dir / "accepted_assembled_member_replacement_map.csv",
        replacement_rows,
        [
            "member_source_json",
            "member_source_image",
            "assembled_group_id",
            "assembled_table_id",
            "province",
            "year",
            "table_family",
            "title",
        ],
    )
    (output_dir / "accepted_assembled_member_source_jsons.txt").write_text(
        "\n".join(sorted(set(accepted_member_source_jsons))) + ("\n" if accepted_member_source_jsons else ""),
        encoding="utf-8",
    )
    summary = {
        "groups_total": len(groups),
        "accepted": len(accepted),
        "rejected": len(rejected),
        "rerun": len(rerun),
        "grouping_ok_rerun_facts": len(grouping_ok_rerun_facts),
        "hold": len(hold),
        "accepted_member_pages_replaced": len(set(accepted_member_source_jsons)),
        "unreviewed": sum(1 for group in groups if not decisions.get(group.get("group_id", ""), {}).get("review_decision")),
        "updated_at": datetime.now().isoformat(timespec="seconds"),
    }
    (output_dir / "assembled_group_review_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")


def pipe_values(value: Any) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    return [part.strip() for part in text.split("|") if part.strip()]


def group_review_decision_class(decision: str) -> str:
    if decision in {"ok", "ok_unnumbered"}:
        return "ok"
    if decision in {"has_duplicate"}:
        return "bad"
    return "warn"


def image_item_from_source_json(source_json: str, label: str, role: str) -> dict[str, Any]:
    source_image = image_path_from_source_json(source_json)
    path = resolve_image({"source_image": source_image, "source_image_local": source_image})
    return {
        "label": label,
        "role": role,
        "source_json": source_json,
        "source_image": source_image,
        "local_path": str(path) if path else "",
        "exists": bool(path),
    }


def image_item_from_group_member(source_json: str, source_image: str, label: str) -> dict[str, Any]:
    path = resolve_image({"source_image": source_image, "source_image_local": source_image})
    if not path and source_json:
        source_image = source_image or image_path_from_source_json(source_json)
        path = resolve_image({"source_image": source_image, "source_image_local": source_image})
    return {
        "label": label,
        "role": "current",
        "source_json": source_json,
        "source_image": source_image,
        "local_path": str(path) if path else "",
        "exists": bool(path),
    }


def group_review_images(row: dict[str, str], current_group: dict[str, Any] | None) -> dict[str, list[dict[str, Any]]]:
    current_images: list[dict[str, Any]] = []
    proposed_images: list[dict[str, Any]] = []
    evidence_images: list[dict[str, Any]] = []

    if current_group:
        source_jsons = current_group.get("page_source_jsons") or []
        source_images = current_group.get("page_source_images") or []
        max_len = max(len(source_jsons), len(source_images), 0)
        for idx in range(max_len):
            current_images.append(
                image_item_from_group_member(
                    source_jsons[idx] if idx < len(source_jsons) else "",
                    source_images[idx] if idx < len(source_images) else "",
                    f"Current member {idx + 1}",
                )
            )

    correct_first = row.get("correct_first_source_json", "")
    proposed_sources = pipe_values(row.get("correct_member_order_source_jsons"))
    if not proposed_sources and correct_first:
        proposed_sources = [correct_first]
    for idx, source_json in enumerate(proposed_sources, start=1):
        label = f"Proposed order {idx}"
        if correct_first and source_json_match(source_json, correct_first):
            label += " (correct first)"
        proposed_images.append(image_item_from_source_json(source_json, label, "proposed"))

    evidence_sources: list[tuple[str, str]] = []
    for source_json in pipe_values(row.get("missing_preceding_source_jsons")):
        evidence_sources.append(("Missing preceding candidate", source_json))
    if row.get("recovered_head_source_json"):
        evidence_sources.append(("Recovered head", row.get("recovered_head_source_json", "")))
    for source_json in pipe_values(row.get("extra_member_source_jsons")):
        evidence_sources.append(("Extra member to remove", source_json))
    seen = set()
    for label, source_json in evidence_sources:
        if not source_json or source_json in seen:
            continue
        seen.add(source_json)
        evidence_images.append(image_item_from_source_json(source_json, label, "evidence"))

    all_images = current_images + proposed_images + evidence_images
    for view_idx, image in enumerate(all_images):
        image["view_index"] = view_idx
    return {
        "current": current_images,
        "proposed": proposed_images,
        "evidence": evidence_images,
        "all": all_images,
    }


def group_review_action_row(row: dict[str, str], decision: dict[str, str]) -> dict[str, str]:
    return {
        "group_id": row.get("group_id", ""),
        "human_decision": decision.get("human_decision", ""),
        "claude_decision": row.get("decision", ""),
        "confidence": row.get("confidence", ""),
        "correct_first_source_json": decision.get("approved_correct_first_source_json") or row.get("correct_first_source_json", ""),
        "member_order_source_jsons": decision.get("approved_member_order_source_jsons") or row.get("correct_member_order_source_jsons", ""),
        "merge_with_group_ids": decision.get("approved_merge_with_group_ids") or row.get("merge_with_group_ids", ""),
        "missing_preceding_source_jsons": decision.get("approved_missing_preceding_source_jsons") or row.get("missing_preceding_source_jsons", ""),
        "extra_member_source_jsons": decision.get("approved_extra_member_source_jsons") or row.get("extra_member_source_jsons", ""),
        "proposed_split_subgroups": decision.get("approved_split_subgroups") or row.get("proposed_split_subgroups", ""),
        "human_notes": decision.get("human_notes", ""),
        "province": row.get("province", ""),
        "year": row.get("year", ""),
        "table_family": row.get("table_family", ""),
    }


def write_group_review_exports(rows: list[dict[str, str]], decisions_csv: Path, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    decisions = by_id(read_csv(decisions_csv), "group_id")
    index_rows = []
    approved_rows = []
    merge_rows = []
    split_rows = []
    incomplete_rows = []
    manual_rows = []
    for row in rows:
        group_id = row.get("group_id", "")
        decision = decisions.get(group_id, {})
        action = group_review_action_row(row, decision)
        index_rows.append(action)
        human_decision = decision.get("human_decision", "")
        claude_decision = row.get("decision", "")
        if human_decision in {"approve_proposal", "approve_ok"}:
            approved_rows.append(action)
            if claude_decision == "merge_candidate":
                merge_rows.append(action)
            elif claude_decision == "split_candidate":
                split_rows.append(action)
            elif claude_decision in {"incomplete_not_digitized", "missing_preceding", "has_duplicate"}:
                incomplete_rows.append(action)
        elif human_decision in {"needs_manual", "rerun_or_reextract", "reject_proposal"}:
            manual_rows.append(action)

    write_csv(output_dir / "group_review_action_index.csv", index_rows, GROUP_REVIEW_ACTION_FIELDS)
    write_csv(output_dir / "group_review_approved_actions.csv", approved_rows, GROUP_REVIEW_ACTION_FIELDS)
    write_csv(output_dir / "group_review_merge_actions.csv", merge_rows, GROUP_REVIEW_ACTION_FIELDS)
    write_csv(output_dir / "group_review_split_actions.csv", split_rows, GROUP_REVIEW_ACTION_FIELDS)
    write_csv(output_dir / "group_review_incomplete_or_missing_actions.csv", incomplete_rows, GROUP_REVIEW_ACTION_FIELDS)
    write_csv(output_dir / "group_review_manual_or_reject_actions.csv", manual_rows, GROUP_REVIEW_ACTION_FIELDS)
    summary = {
        "groups_total": len(rows),
        "reviewed": sum(1 for row in rows if decisions.get(row.get("group_id", ""), {}).get("human_decision")),
        "approved": len(approved_rows),
        "merge_actions": len(merge_rows),
        "split_actions": len(split_rows),
        "incomplete_or_missing_actions": len(incomplete_rows),
        "manual_or_reject": len(manual_rows),
        "unreviewed": sum(1 for row in rows if not decisions.get(row.get("group_id", ""), {}).get("human_decision")),
        "updated_at": datetime.now().isoformat(timespec="seconds"),
    }
    (output_dir / "group_review_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")


def table_key_index(table_rows: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    index: dict[str, dict[str, str]] = {}
    for row in table_rows:
        values = [
            row.get("table_id", ""),
            row.get("source_json_local", ""),
            row.get("source_json_original", ""),
            row.get("source_image_local", ""),
            row.get("source_image_original", ""),
        ]
        for value in values:
            for key in source_json_keys(value):
                index.setdefault(key, row)
    return index


def find_table_row(raw_id: str, table_rows: list[dict[str, str]], key_index: dict[str, dict[str, str]]) -> dict[str, str]:
    raw_id = str(raw_id or "").strip()
    if raw_id in key_index:
        return key_index[raw_id]
    for key in sorted(source_json_keys(raw_id), key=len, reverse=True):
        if key in key_index:
            return key_index[key]
    for row in table_rows:
        table_id = row.get("table_id", "")
        if table_id and (table_id.startswith(raw_id) or raw_id.startswith(table_id)):
            return row
    return {}


def load_decimal_opus_results(opus_dir: Path, table_rows: list[dict[str, str]], key_index: dict[str, dict[str, str]]) -> dict[str, dict[str, Any]]:
    output: dict[str, dict[str, Any]] = {}
    if not opus_dir.exists():
        return output
    for path in sorted(opus_dir.glob("result_od_*.json")):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        if isinstance(data, list):
            items = []
            for idx, item in enumerate(data):
                if isinstance(item, dict):
                    raw_id = str(item.get("table_id") or item.get("id") or item.get("source_json") or f"{path.name}:{idx}")
                    items.append((raw_id, item))
        elif isinstance(data, dict):
            items = [(str(key), value) for key, value in data.items() if isinstance(value, dict)]
        else:
            items = []
        for raw_id, result in items:
            table = find_table_row(raw_id, table_rows, key_index)
            table_id = table.get("table_id") or raw_id
            record = dict(result)
            record["opus_result_file"] = str(path)
            record["opus_raw_id"] = raw_id
            output[table_id] = record
    return output


def load_decimal_review_rows(
    rollout_dir: Path,
    table_rows: list[dict[str, str]],
    key_index: dict[str, dict[str, str]],
    revalidated_tables: Path,
    opus_by_table_id: dict[str, dict[str, Any]],
) -> list[dict[str, str]]:
    fixed_path = rollout_dir / "decimal_fixed_tables.csv"
    needs_path = rollout_dir / "decimal_needs_review.csv"
    rejected_path = rollout_dir / "decimal_rejected.csv"
    rejected_by_id = {row.get("table_id", ""): row.get("reason", "") for row in read_csv(rejected_path)}
    rows: list[dict[str, str]] = []
    seen_queue_ids = set()

    def add_row(raw: dict[str, str], bucket: str, source_label: str) -> None:
        raw_id = raw.get("table_id", "")
        table = find_table_row(raw_id, table_rows, key_index)
        table_id = table.get("table_id") or raw_id
        opus = opus_by_table_id.get(table_id, {})
        queue_id = f"{bucket}::{raw_id or table_id}"
        if queue_id in seen_queue_ids:
            return
        seen_queue_ids.add(queue_id)
        rows.append(
            {
                "queue_id": queue_id,
                "review_bucket": bucket,
                "source_label": source_label,
                "raw_table_id": raw_id,
                "table_id": table_id,
                "province": table.get("province", ""),
                "year": table.get("year", ""),
                "title": table.get("title", ""),
                "table_family": table.get("table_family", ""),
                "format_class": table.get("format_class", ""),
                "measure_category": table.get("measure_category", ""),
                "denominator_hint": table.get("denominator_hint", ""),
                "effective_quality_status": table.get("effective_quality_status", ""),
                "n_checks": table.get("n_checks", ""),
                "failed_checks": table.get("failed_checks", ""),
                "facts_rows": table.get("facts_rows", ""),
                "source_image_local": table.get("source_image_local", "") or image_path_from_source_json(raw_id),
                "source_json_local": table.get("source_json_local", ""),
                "source_image_original": table.get("source_image_original", ""),
                "source_json_original": table.get("source_json_original", ""),
                "excel_path": table.get("excel_path", ""),
                "decimal_places": raw.get("decimal_places", "") or str(opus.get("decimal_places", "") or ""),
                "reason": raw.get("reason", ""),
                "rejected_reason": rejected_by_id.get(raw_id, ""),
                "validation_failures": raw.get("n_failures", ""),
                "validation_fail_levels": raw.get("fail_levels", ""),
                "opus_decimal_places": str(opus.get("decimal_places", "") or ""),
                "opus_uniform": str(opus.get("uniform", "") or ""),
                "opus_column_dps": json.dumps(opus.get("column_dps", ""), ensure_ascii=False) if opus else "",
                "opus_sample": json.dumps(opus.get("sample", ""), ensure_ascii=False) if opus else "",
                "opus_looks_like": str(opus.get("looks_like", "") or ""),
                "opus_note": str(opus.get("note", "") or ""),
                "opus_result_file": str(opus.get("opus_result_file", "") or ""),
            }
        )

    for row in read_csv(fixed_path):
        add_row(row, "auto_fixed", "decimal_fixed_tables.csv")
    for row in read_csv(needs_path):
        add_row(row, "needs_review", "decimal_needs_review.csv")
    for row in read_csv(revalidated_tables):
        if row.get("verified") != "fail":
            continue
        add_row(
            {
                "table_id": row.get("table_id", ""),
                "reason": f"post-cleaning validation failed; kind={row.get('kind', '')}; levels={row.get('fail_levels', '')}",
                "n_failures": row.get("n_failures", ""),
                "fail_levels": row.get("fail_levels", ""),
            },
            "validation_fail",
            "revalidated_tables.csv",
        )

    bucket_order = {"needs_review": 0, "validation_fail": 1, "auto_fixed": 2}
    rows.sort(
        key=lambda row: (
            bucket_order.get(row.get("review_bucket", ""), 9),
            row.get("province", ""),
            row.get("year", ""),
            row.get("title", ""),
            row.get("table_id", ""),
        )
    )
    return rows


FACT_COORD_FIELDS = [
    "section",
    "row_label_raw",
    "row_label_std",
    "row_dimension",
    "row_group_value",
    "column_group_raw",
    "column_group_std",
    "column_dimension",
    "column_group_value",
    "sub_column_raw",
    "sub_column_std",
    "sex",
    "civil_condition",
    "religion",
    "resolved_age_label_std",
    "measure_unit",
]


def load_fact_subset(path: Path, table_ids: set[str]) -> dict[str, list[dict[str, str]]]:
    facts = {table_id: [] for table_id in table_ids if table_id}
    if not path.exists() or not table_ids:
        return facts
    with path.open(newline="", encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        for row in reader:
            table_id = row.get("table_id", "")
            if table_id in facts:
                facts[table_id].append(row)
    return facts


def fact_coord_key(row: dict[str, Any]) -> tuple[str, ...]:
    return tuple(str(row.get(field, "") or "") for field in FACT_COORD_FIELDS)


def decimal_changed_rows(original: list[dict[str, str]], cleaned: list[dict[str, str]], max_rows: int = 300) -> list[list[Any]]:
    original_by_key = {fact_coord_key(row): row for row in original}
    cleaned_by_key = {fact_coord_key(row): row for row in cleaned}
    keys = list(original_by_key)
    for key in cleaned_by_key:
        if key not in original_by_key:
            keys.append(key)
    rows: list[list[Any]] = [
        [
            "change",
            "section",
            "row_label",
            "column_group",
            "sub_column",
            "sex",
            "civil_condition",
            "religion",
            "age",
            "measure_unit",
            "original_value",
            "cleaned_value",
        ]
    ]
    for key in keys:
        old = original_by_key.get(key)
        new = cleaned_by_key.get(key)
        old_value = old.get("value", "") if old else ""
        new_value = new.get("value", "") if new else ""
        if old and new and old_value == new_value:
            continue
        row = new or old or {}
        change = "changed" if old and new else ("removed" if old else "added")
        rows.append(
            [
                change,
                row.get("section", ""),
                row.get("row_label_raw", "") or row.get("row_label_std", ""),
                row.get("column_group_raw", "") or row.get("column_group_std", ""),
                row.get("sub_column_raw", "") or row.get("sub_column_std", ""),
                row.get("sex", ""),
                row.get("civil_condition", ""),
                row.get("religion", ""),
                row.get("resolved_age_label_std", ""),
                row.get("measure_unit", ""),
                old_value,
                new_value,
            ]
        )
        if len(rows) > max_rows:
            break
    return rows


def decimal_bucket_class(bucket: str) -> str:
    if bucket == "auto_fixed":
        return "ok"
    if bucket == "validation_fail":
        return "bad"
    return "warn"


def decimal_decision_map(path: Path) -> dict[str, dict[str, str]]:
    return by_id(read_csv(path), "queue_id")


def decimal_review_index_row(row: dict[str, str], decision: dict[str, str], image_exists: bool, original_count: int, cleaned_count: int, changed_count: int) -> dict[str, Any]:
    return {
        "queue_id": row.get("queue_id", ""),
        "review_bucket": row.get("review_bucket", ""),
        "human_decision": decision.get("human_decision", ""),
        "table_id": row.get("table_id", ""),
        "raw_table_id": row.get("raw_table_id", ""),
        "province": row.get("province", ""),
        "year": row.get("year", ""),
        "table_family": row.get("table_family", ""),
        "format_class": row.get("format_class", ""),
        "measure_category": row.get("measure_category", ""),
        "title": row.get("title", ""),
        "decimal_places": row.get("decimal_places", ""),
        "reason": row.get("reason", ""),
        "rejected_reason": row.get("rejected_reason", ""),
        "validation_failures": row.get("validation_failures", ""),
        "facts_original": original_count,
        "facts_cleaned": cleaned_count,
        "changed_values": changed_count,
        "image_exists": "1" if image_exists else "0",
    }


def write_decimal_review_exports(
    rows: list[dict[str, str]],
    decisions_csv: Path,
    output_dir: Path,
    original_facts_by_table: dict[str, list[dict[str, str]]],
    cleaned_facts_by_table: dict[str, list[dict[str, str]]],
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    decisions = decimal_decision_map(decisions_csv)
    index_rows = []
    rerun_rows = []
    accepted_rows = []
    held_rows = []
    excluded_rows = []
    for row in rows:
        table_id = row.get("table_id", "")
        original = original_facts_by_table.get(table_id, [])
        cleaned = cleaned_facts_by_table.get(table_id, [])
        changed = max(len(decimal_changed_rows(original, cleaned)) - 1, 0)
        image_exists = bool(resolve_image(row))
        decision = decisions.get(row.get("queue_id", ""), {})
        index_row = decimal_review_index_row(row, decision, image_exists, len(original), len(cleaned), changed)
        index_rows.append(index_row)
        human_decision = decision.get("human_decision", "")
        if human_decision == "decimal_ok":
            accepted_rows.append(index_row)
        elif human_decision in {"revert_to_original", "manual_scale_fix", "rerun_decimal_prompt"}:
            rerun_rows.append(index_row)
        elif human_decision == "hold":
            held_rows.append(index_row)
        elif human_decision == "exclude_for_now":
            excluded_rows.append(index_row)
    write_csv(output_dir / "decimal_review_index.csv", index_rows, DECIMAL_REVIEW_INDEX_FIELDS)
    write_csv(output_dir / "decimal_review_accepted.csv", accepted_rows, DECIMAL_REVIEW_INDEX_FIELDS)
    write_csv(output_dir / "decimal_review_rerun_or_manual.csv", rerun_rows, DECIMAL_REVIEW_INDEX_FIELDS)
    write_csv(output_dir / "decimal_review_hold.csv", held_rows, DECIMAL_REVIEW_INDEX_FIELDS)
    write_csv(output_dir / "decimal_review_exclude.csv", excluded_rows, DECIMAL_REVIEW_INDEX_FIELDS)
    summary = {
        "queue_total": len(rows),
        "auto_fixed": sum(1 for row in rows if row.get("review_bucket") == "auto_fixed"),
        "needs_review": sum(1 for row in rows if row.get("review_bucket") == "needs_review"),
        "validation_fail": sum(1 for row in rows if row.get("review_bucket") == "validation_fail"),
        "residue": sum(1 for row in rows if row.get("review_bucket") in {"needs_review", "validation_fail"}),
        "reviewed": sum(1 for row in rows if decisions.get(row.get("queue_id", ""), {}).get("human_decision")),
        "unreviewed": sum(1 for row in rows if not decisions.get(row.get("queue_id", ""), {}).get("human_decision")),
        "updated_at": datetime.now().isoformat(timespec="seconds"),
    }
    (output_dir / "decimal_review_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")


def load_decimal_decision_rows(
    queue_csv: Path,
    table_rows: list[dict[str, str]],
    key_index: dict[str, dict[str, str]],
) -> list[dict[str, str]]:
    queue_rows = read_csv(queue_csv)
    output = []
    for queue_row in queue_rows:
        table = find_table_row(queue_row.get("table_id", ""), table_rows, key_index)
        row = dict(queue_row)
        row["queue_id"] = queue_row.get("table_id", "")
        row["table_id"] = table.get("table_id") or queue_row.get("table_id", "")
        row["raw_table_id"] = queue_row.get("table_id", "")
        row["province"] = table.get("province", "")
        row["year"] = table.get("year", "")
        row["title"] = table.get("title", "")
        row["table_family"] = table.get("table_family", "")
        row["format_class"] = table.get("format_class", "")
        row["measure_category"] = table.get("measure_category", "")
        row["excel_path"] = table.get("excel_path", "")
        row["source_image_local"] = queue_row.get("image", "")
        output.append(row)
    order = {
        "dp1_uniform": 0,
        "dp2_uniform": 1,
        "dp4_uniform": 2,
        "NON_uniform_mixed": 3,
        "dp0_suspect": 4,
    }
    output.sort(key=lambda row: (order.get(row.get("category", ""), 9), row.get("province", ""), row.get("year", ""), row.get("title", ""), row.get("table_id", "")))
    return output


def decimal_dp_decision_map(path: Path) -> dict[str, dict[str, str]]:
    return by_id(read_csv(path), "queue_id")


def decimal_decision_image(row: dict[str, str]) -> Path | None:
    image = row.get("image") or row.get("source_image_local") or ""
    return resolve_local_path(image) or resolve_image({"source_image": image, "source_image_local": image})


def decimal_decision_label(decision: dict[str, str]) -> str:
    flag = decision.get("human_flag", "")
    dp = decision.get("human_dp", "")
    if flag:
        return flag
    if dp:
        return f"{dp} dp"
    return "unreviewed"


def decimal_decision_action(row: dict[str, str], decision: dict[str, str]) -> str:
    flag = decision.get("human_flag", "")
    dp = decision.get("human_dp", "")
    if flag in {"mixed", "weird"}:
        return "rerun_or_hand_fix"
    if flag == "already_correct":
        return "no_divide"
    if dp in {"1", "2", "3+", "3", "4"}:
        return "divide"
    if dp == "0":
        return "no_divide"
    return ""


def decimal_decision_agreement(row: dict[str, str], decision: dict[str, str]) -> str:
    dp = decision.get("human_dp", "")
    flag = decision.get("human_flag", "")
    opus_dp = row.get("opus_dp", "")
    if not dp or flag:
        return ""
    if dp == "3+" and opus_dp and opus_dp not in {"0", "1", "2"}:
        return "agree"
    return "agree" if dp == opus_dp else "disagree"


def decimal_dp_index_row(row: dict[str, str], decision: dict[str, str]) -> dict[str, Any]:
    return {
        "queue_id": row.get("queue_id", ""),
        "table_id": row.get("table_id", ""),
        "category": row.get("category", ""),
        "human_dp": decision.get("human_dp", ""),
        "human_flag": decision.get("human_flag", ""),
        "province": row.get("province", ""),
        "year": row.get("year", ""),
        "table_family": row.get("table_family", ""),
        "title": row.get("title", ""),
        "opus_dp": row.get("opus_dp", ""),
        "opus_uniform": row.get("opus_uniform", ""),
        "opus_sample": row.get("opus_sample", ""),
        "rate_min": row.get("rate_min", ""),
        "rate_max": row.get("rate_max", ""),
        "agreement": decimal_decision_agreement(row, decision),
        "action": decimal_decision_action(row, decision),
        "image_exists": "1" if decimal_decision_image(row) else "0",
    }


def write_decimal_decision_exports(rows: list[dict[str, str]], decisions_csv: Path, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    decisions = decimal_dp_decision_map(decisions_csv)
    index_rows = []
    divide_rows = []
    no_divide_rows = []
    rerun_rows = []
    for row in rows:
        decision = decisions.get(row.get("queue_id", ""), {})
        index_row = decimal_dp_index_row(row, decision)
        index_rows.append(index_row)
        action = index_row.get("action", "")
        if action == "divide":
            divide_rows.append(index_row)
        elif action == "no_divide":
            no_divide_rows.append(index_row)
        elif action == "rerun_or_hand_fix":
            rerun_rows.append(index_row)
    reviewed_rows = [row for row in index_rows if row.get("human_dp") or row.get("human_flag")]
    comparable = [row for row in reviewed_rows if row.get("agreement")]
    summary = {
        "queue_total": len(rows),
        "reviewed": len(reviewed_rows),
        "unreviewed": len(rows) - len(reviewed_rows),
        "divide": len(divide_rows),
        "no_divide": len(no_divide_rows),
        "rerun_or_hand_fix": len(rerun_rows),
        "opus_comparable": len(comparable),
        "opus_agree": sum(1 for row in comparable if row.get("agreement") == "agree"),
        "opus_disagree": sum(1 for row in comparable if row.get("agreement") == "disagree"),
        "updated_at": datetime.now().isoformat(timespec="seconds"),
    }
    write_csv(output_dir / "decimal_decision_review_index.csv", index_rows, DECIMAL_DP_REVIEW_INDEX_FIELDS)
    write_csv(output_dir / "decimal_decision_apply_divides.csv", divide_rows, DECIMAL_DP_REVIEW_INDEX_FIELDS)
    write_csv(output_dir / "decimal_decision_no_divide.csv", no_divide_rows, DECIMAL_DP_REVIEW_INDEX_FIELDS)
    write_csv(output_dir / "decimal_decision_rerun_or_hand_fix.csv", rerun_rows, DECIMAL_DP_REVIEW_INDEX_FIELDS)
    (output_dir / "decimal_decision_review_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")


def load_reconciliation_rows(
    annotations_path: Path,
    table_rows: list[dict[str, str]],
    key_index: dict[str, dict[str, str]],
) -> list[dict[str, Any]]:
    if not annotations_path.exists():
        return []
    try:
        annotations = json.loads(annotations_path.read_text(encoding="utf-8"))
    except Exception:
        return []
    output = []
    for raw_table_id, annotation in annotations.items():
        table = find_table_row(raw_table_id, table_rows, key_index)
        table_id = table.get("table_id") or raw_table_id
        failures = annotation.get("failures") or []
        row = {
            "table_id": table_id,
            "raw_table_id": raw_table_id,
            "image": annotation.get("image", ""),
            "n_failures": str(annotation.get("n_failures", len(failures))),
            "failures": failures,
            "province": table.get("province", ""),
            "year": table.get("year", ""),
            "title": table.get("title", ""),
            "table_family": table.get("table_family", ""),
            "format_class": table.get("format_class", ""),
            "measure_category": table.get("measure_category", ""),
            "source_image_local": annotation.get("image", "") or table.get("source_image_local", ""),
            "source_json_local": table.get("source_json_local", ""),
            "excel_path": table.get("excel_path", ""),
        }
        output.append(row)
    output.sort(key=lambda row: (row.get("province", ""), row.get("year", ""), row.get("title", ""), row.get("table_id", "")))
    return output


def resolve_reconciliation_image(row: dict[str, Any]) -> Path | None:
    image = str(row.get("image") or row.get("source_image_local") or "")
    path = resolve_local_path(image)
    if path:
        return path
    return resolve_image({"source_image": image, "source_image_local": image})


def normalize_recon_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def recon_is_total_row(row_key: tuple[str, str]) -> bool:
    label = normalize_recon_text(row_key[1])
    return label in {"total", "all ages", "grand total", "total population"} or "total" in label


def recon_failure_matches(failure: dict[str, Any], section: str, col_label: str) -> bool:
    failure_section = normalize_recon_text(failure.get("section"))
    failure_group = normalize_recon_text(failure.get("group"))
    failure_col = normalize_recon_text(failure.get("col"))
    section_text = normalize_recon_text(section)
    col_text = normalize_recon_text(col_label)
    if failure_section and failure_section != section_text:
        return False
    group_ok = not failure_group or failure_group in col_text
    col_ok = not failure_col or failure_col in col_text
    return group_ok and col_ok


def highlighted_reconciliation_matrix(
    facts: list[dict[str, Any]],
    failures: list[dict[str, Any]],
    max_unique_cols: int = 180,
) -> tuple[str, str]:
    if not facts:
        return "<div class='pad muted'>No facts available for this table.</div>", "no_facts_available"

    row_keys: list[tuple[str, str]] = []
    col_keys: list[str] = []
    seen_rows = set()
    seen_cols = set()
    values: dict[tuple[tuple[str, str], str], str] = {}
    conflicts = 0
    for fact in facts:
        section = clean_label(fact.get("section"))
        row_label = fact_row_label(fact)
        col_label = fact_col_label(fact)
        row_key = (section, row_label)
        if row_key not in seen_rows:
            seen_rows.add(row_key)
            row_keys.append(row_key)
        if col_label not in seen_cols and len(col_keys) < max_unique_cols:
            seen_cols.add(col_label)
            col_keys.append(col_label)
        if col_label not in seen_cols:
            continue
        key = (row_key, col_label)
        value = clean_label(fact.get("value"))
        old = values.get(key)
        if old and old != value:
            conflicts += 1
            values[key] = f"{old} / {value}"
        else:
            values[key] = value

    highlighted_cols: set[tuple[str, str]] = set()
    for row_key in row_keys:
        for col in col_keys:
            if any(recon_failure_matches(failure, row_key[0], col) for failure in failures):
                highlighted_cols.add((row_key[0], col))

    out = ["<table>"]
    out.append(
        "<tr><th>section</th><th>row_label</th>"
        + "".join(
            f"<th class='{'recon-col' if any((section, col) in highlighted_cols for section, _ in row_keys) else ''}'>{html.escape(str(col))}</th>"
            for col in col_keys
        )
        + "</tr>"
    )
    for row_key in row_keys:
        cells = [f"<td>{html.escape(row_key[0])}</td>", f"<td>{html.escape(row_key[1])}</td>"]
        for col in col_keys:
            classes = []
            if (row_key[0], col) in highlighted_cols:
                classes.append("recon-col")
                if recon_is_total_row(row_key):
                    classes.append("recon-total")
            value = html.escape(values.get((row_key, col), ""))
            cells.append(f"<td class='{' '.join(classes)}'>{value}</td>")
        out.append("<tr>" + "".join(cells) + "</tr>")
    out.append("</table>")
    status = f"highlighted_from_recorded_facts | facts={len(facts)} | rows={len(row_keys)} | columns={len(col_keys)} | highlighted_cells={len(highlighted_cols)}"
    if conflicts:
        status += f" | conflicting_duplicate_cells={conflicts}"
    if len(seen_cols) >= max_unique_cols:
        status += f" | columns_capped_at={max_unique_cols}"
    return "\n".join(out), status


def reconciliation_decision_map(path: Path) -> dict[str, dict[str, str]]:
    return by_id(read_csv(path), "table_id")


def reconciliation_index_row(
    row: dict[str, Any],
    decision: dict[str, str],
    image_exists: bool,
    original_count: int,
    cleaned_count: int,
) -> dict[str, Any]:
    failures = row.get("failures") or []
    return {
        "table_id": row.get("table_id", ""),
        "human_decision": decision.get("human_decision", ""),
        "province": row.get("province", ""),
        "year": row.get("year", ""),
        "table_family": row.get("table_family", ""),
        "title": row.get("title", ""),
        "n_failures": row.get("n_failures", ""),
        "levels": "|".join(sorted({str(failure.get("level", "")) for failure in failures if failure.get("level")})),
        "image_exists": "1" if image_exists else "0",
        "facts_original": original_count,
        "facts_cleaned": cleaned_count,
    }


def write_reconciliation_review_exports(
    rows: list[dict[str, Any]],
    decisions_csv: Path,
    output_dir: Path,
    original_facts_by_table: dict[str, list[dict[str, str]]],
    cleaned_facts_by_table: dict[str, list[dict[str, str]]],
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    decisions = reconciliation_decision_map(decisions_csv)
    index_rows = []
    accept_rows = []
    fix_rows = []
    rerun_rows = []
    exclude_rows = []
    for row in rows:
        table_id = row.get("table_id", "")
        decision = decisions.get(table_id, {})
        index_row = reconciliation_index_row(
            row,
            decision,
            bool(resolve_reconciliation_image(row)),
            len(original_facts_by_table.get(table_id, [])),
            len(cleaned_facts_by_table.get(table_id, [])),
        )
        index_rows.append(index_row)
        human_decision = decision.get("human_decision", "")
        if human_decision == "accept_artifact":
            accept_rows.append(index_row)
        elif human_decision == "fix_cell":
            fix_rows.append(index_row)
        elif human_decision == "rerun_extract":
            rerun_rows.append(index_row)
        elif human_decision == "exclude_for_now":
            exclude_rows.append(index_row)
    write_csv(output_dir / "reconciliation_review_index.csv", index_rows, RECONCILIATION_REVIEW_INDEX_FIELDS)
    write_csv(output_dir / "reconciliation_review_accept_artifact.csv", accept_rows, RECONCILIATION_REVIEW_INDEX_FIELDS)
    write_csv(output_dir / "reconciliation_review_fix_cell.csv", fix_rows, RECONCILIATION_REVIEW_INDEX_FIELDS)
    write_csv(output_dir / "reconciliation_review_rerun_extract.csv", rerun_rows, RECONCILIATION_REVIEW_INDEX_FIELDS)
    write_csv(output_dir / "reconciliation_review_exclude.csv", exclude_rows, RECONCILIATION_REVIEW_INDEX_FIELDS)
    summary = {
        "queue_total": len(rows),
        "reviewed": sum(1 for row in rows if decisions.get(row.get("table_id", ""), {}).get("human_decision")),
        "unreviewed": sum(1 for row in rows if not decisions.get(row.get("table_id", ""), {}).get("human_decision")),
        "accept_artifact": len(accept_rows),
        "fix_cell": len(fix_rows),
        "rerun_extract": len(rerun_rows),
        "exclude_for_now": len(exclude_rows),
        "updated_at": datetime.now().isoformat(timespec="seconds"),
    }
    (output_dir / "reconciliation_review_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")


def excel_sheet_dicts(path: Path | None, sheet_name: str, max_rows: int = 10000) -> tuple[list[dict[str, Any]], str]:
    if not path:
        return [], "excel_not_found"
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        if sheet_name not in wb.sheetnames:
            wb.close()
            return [], f"sheet_not_found:{sheet_name}"
        ws = wb[sheet_name]
        values = ws.iter_rows(max_row=min(ws.max_row, max_rows + 1), values_only=True)
        headers = [str(cell or "") for cell in next(values, [])]
        rows = []
        for row in values:
            rows.append({headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))})
        status = f"{path} | sheet={sheet_name} | facts_loaded={len(rows)} | original={ws.max_row}x{ws.max_column}"
        if ws.max_row > max_rows + 1:
            status += f" | truncated_to={max_rows}"
        wb.close()
        return rows, status
    except Exception as exc:
        return [], f"excel_read_error: {type(exc).__name__}: {exc}"


def excel_sheet_rows(path: Path | None, sheet_name: str, max_rows: int = 90, max_cols: int = 60) -> tuple[list[list[Any]], str]:
    if not path:
        return [], "excel_not_found"
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        name = sheet_name if sheet_name in wb.sheetnames else wb.sheetnames[0]
        ws = wb[name]
        rows = []
        for row in ws.iter_rows(max_row=min(ws.max_row, max_rows), max_col=min(ws.max_column, max_cols), values_only=True):
            rows.append(["" if cell is None else cell for cell in row])
        status = f"{path} | sheet={name} | original={ws.max_row}x{ws.max_column}"
        wb.close()
        return rows, status
    except Exception as exc:
        return [], f"excel_read_error: {type(exc).__name__}: {exc}"


def clean_label(value: Any) -> str:
    return str(value or "").strip()


def contained_casefold(needle: str, haystack: str) -> bool:
    needle = needle.casefold().strip()
    haystack = haystack.casefold().strip()
    return bool(needle and needle in haystack)


def fact_row_label(fact: dict[str, Any]) -> str:
    return clean_label(fact.get("row_label_raw")) or clean_label(fact.get("row_label_std")) or clean_label(fact.get("row_group_value")) or "(row)"


def fact_col_label(fact: dict[str, Any]) -> str:
    parts = []
    column = clean_label(fact.get("column_group_raw")) or clean_label(fact.get("column_group_std"))
    sub = clean_label(fact.get("sub_column_raw")) or clean_label(fact.get("sub_column_std"))
    if column:
        parts.append(column)
    if sub:
        parts.append(sub)
    raw_text = " ".join(parts)
    # Add inferred semantic fields only when they disambiguate repeated raw labels.
    # This keeps the display source-faithful while preventing collisions.
    for field, prefix in [("sex", "sex"), ("civil_condition", "civil"), ("religion", "religion")]:
        value = clean_label(fact.get(field))
        if value and value != "All" and not contained_casefold(value, raw_text):
            parts.append(f"{prefix}={value}")
    return " | ".join(parts) or "(value)"


def build_recorded_fact_matrix(facts: list[dict[str, Any]], max_unique_cols: int = 240) -> tuple[list[list[Any]], str]:
    if not facts:
        return [], "no_facts_available"
    row_keys: list[tuple[str, str]] = []
    col_keys: list[str] = []
    seen_rows = set()
    seen_cols = set()
    values: dict[tuple[tuple[str, str], str], str] = {}
    conflicts = 0
    for fact in facts:
        section = clean_label(fact.get("section"))
        row_label = fact_row_label(fact)
        col_label = fact_col_label(fact)
        row_key = (section, row_label)
        if row_key not in seen_rows:
            seen_rows.add(row_key)
            row_keys.append(row_key)
        if col_label not in seen_cols and len(col_keys) < max_unique_cols:
            seen_cols.add(col_label)
            col_keys.append(col_label)
        if col_label not in seen_cols:
            continue
        key = (row_key, col_label)
        value = clean_label(fact.get("value"))
        old = values.get(key)
        if old and old != value:
            conflicts += 1
            values[key] = f"{old} / {value}"
        else:
            values[key] = value
    rows: list[list[Any]] = [["section", "row_label"] + col_keys]
    for row_key in row_keys:
        rows.append([row_key[0], row_key[1]] + [values.get((row_key, col), "") for col in col_keys])
    status = f"built_from_recorded_facts | facts={len(facts)} | rows={len(row_keys)} | columns={len(col_keys)}"
    if conflicts:
        status += f" | conflicting_duplicate_cells={conflicts}"
    if len(seen_cols) >= max_unique_cols:
        status += f" | columns_capped_at={max_unique_cols}"
    return rows, status


def build_before_after_fact_matrix(
    original: list[dict[str, Any]],
    cleaned: list[dict[str, Any]],
    max_unique_cols: int = 120,
) -> tuple[list[list[Any]], str]:
    if not original and not cleaned:
        return [], "no_facts_available"
    row_keys: list[tuple[str, str]] = []
    col_keys: list[str] = []
    seen_rows = set()
    seen_cols = set()
    original_values: dict[tuple[tuple[str, str], str], str] = {}
    cleaned_values: dict[tuple[tuple[str, str], str], str] = {}
    conflicts = 0

    def ingest(facts: list[dict[str, Any]], values: dict[tuple[tuple[str, str], str], str]) -> None:
        nonlocal conflicts
        for fact in facts:
            section = clean_label(fact.get("section"))
            row_label = fact_row_label(fact)
            col_label = fact_col_label(fact)
            row_key = (section, row_label)
            if row_key not in seen_rows:
                seen_rows.add(row_key)
                row_keys.append(row_key)
            if col_label not in seen_cols and len(col_keys) < max_unique_cols:
                seen_cols.add(col_label)
                col_keys.append(col_label)
            if col_label not in seen_cols:
                continue
            key = (row_key, col_label)
            value = clean_label(fact.get("value"))
            old = values.get(key)
            if old and old != value:
                conflicts += 1
                values[key] = f"{old} / {value}"
            else:
                values[key] = value

    ingest(original, original_values)
    ingest(cleaned, cleaned_values)

    header = ["section", "row_label"]
    for col in col_keys:
        header.extend([f"{col} | original", f"{col} | cleaned"])
    rows: list[list[Any]] = [header]
    for row_key in row_keys:
        output_row: list[Any] = [row_key[0], row_key[1]]
        for col in col_keys:
            output_row.extend([original_values.get((row_key, col), ""), cleaned_values.get((row_key, col), "")])
        rows.append(output_row)
    status = (
        f"before_after_from_recorded_facts | original={len(original)} | cleaned={len(cleaned)} "
        f"| rows={len(row_keys)} | columns={len(col_keys)}"
    )
    if conflicts:
        status += f" | conflicting_duplicate_cells={conflicts}"
    if len(seen_cols) >= max_unique_cols:
        status += f" | columns_capped_at={max_unique_cols}"
    return rows, status


def fact_rows_table(facts: list[dict[str, Any]], max_rows: int = 160) -> list[list[Any]]:
    fields = [
        "section",
        "row_label_raw",
        "column_group_raw",
        "sub_column_raw",
        "sex",
        "civil_condition",
        "religion",
        "resolved_age_label_std",
        "measure_unit",
        "value",
    ]
    rows: list[list[Any]] = [fields]
    for fact in facts[:max_rows]:
        rows.append([fact.get(field, "") for field in fields])
    return rows


def packet_facts_rows(packet: dict[str, Any]) -> list[list[Any]]:
    facts = packet.get("summary", {}).get("sample_facts", [])
    if not facts:
        return []
    fields = [
        "section",
        "row_label_raw",
        "row_label_std",
        "column_group_raw",
        "sub_column_raw",
        "sex",
        "civil_condition",
        "religion",
        "resolved_age_label_std",
        "measure_unit",
        "value",
    ]
    rows = [fields]
    for fact in facts[:80]:
        rows.append([fact.get(field, "") for field in fields])
    return rows


def decision_class(decision: str) -> str:
    if decision == "pass":
        return "ok"
    if decision in {"reject_or_rerun", "api_error"}:
        return "bad"
    return "warn"


def filter_indices(rows: list[dict[str, str]], mode: str) -> list[int]:
    mode = (mode or "flagged").strip()
    if mode == "all":
        return list(range(len(rows)))
    if mode == "unresolved":
        return [i for i, r in enumerate(rows) if not r.get("human_decision")]
    if mode == "reviewed":
        return [i for i, r in enumerate(rows) if r.get("human_decision")]
    if mode == "flagged":
        return [i for i, r in enumerate(rows) if r.get("llm_decision") != "pass" or r.get("status") != "ok"]
    if mode in {"pass", "needs_image_compare", "needs_manual_review", "reject_or_rerun"}:
        return [i for i, r in enumerate(rows) if r.get("llm_decision") == mode]
    if mode == "api_error":
        return [i for i, r in enumerate(rows) if r.get("status") == "error" or r.get("llm_decision") == "api_error"]
    if mode == "mortality":
        return [i for i, r in enumerate(rows) if any(token in (r.get("title", "").lower()) for token in ["mortality", "death", "birth"])]
    if mode == "continuation":
        return [i for i, r in enumerate(rows) if r.get("is_continuation_candidate") == "1"]
    if mode == "continuation_manual":
        return [i for i, r in enumerate(rows) if r.get("continuation_needs_manual_multipage_review") == "1"]
    if mode == "priority":
        return [i for i, r in enumerate(rows) if int(r.get("priority_score") or 0) >= 30]
    if mode in {"accept", "reject", "rerun", "hold", "image_compared_accept", "image_compared_reject"}:
        return [i for i, r in enumerate(rows) if r.get("human_decision") == mode]
    return list(range(len(rows)))


def index_continuation_candidates(rows: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    by_table_id = by_id(rows)
    by_source_json = by_id(rows, "source_json")
    merged = dict(by_source_json)
    merged.update(by_table_id)
    return merged


def continuation_review_row(row: dict[str, str]) -> dict[str, str]:
    manual = row.get("needs_manual_multipage_review") == "1"
    table_id = row.get("table_id") or row.get("source_json") or row.get("source_image")
    return {
        "table_id": table_id,
        "source_json": row.get("source_json", ""),
        "source_image_local": row.get("source_image", ""),
        "excel_path": "",
        "province": row.get("province", ""),
        "year": row.get("year", ""),
        "table_family": row.get("table_family", ""),
        "title": row.get("title", ""),
        "status": "continuation_only",
        "llm_decision": "needs_manual_review" if manual else "needs_image_compare",
        "reason": row.get("auto_attach_reason", ""),
        "flags": "continuation_candidate",
        "priority_score": "35" if manual else "20",
        "is_continuation_candidate": "1",
        "continuation_auto_attach_reason": row.get("auto_attach_reason", ""),
        "continuation_candidate_title": row.get("candidate_title", ""),
        "continuation_candidate_source_json": row.get("candidate_source_json", ""),
        "continuation_assigned_group_id": row.get("assigned_group_id", ""),
        "continuation_needs_manual_multipage_review": row.get("needs_manual_multipage_review", ""),
    }


def create_app(
    results_csv: Path,
    packets_jsonl: Path,
    decisions_csv: Path,
    rerun_focus_queue: Path = DEFAULT_RERUN_FOCUS_QUEUE,
    continuation_candidates_csv: Path = DEFAULT_CONTINUATION_CANDIDATES,
    assembled_jsonl: Path = DEFAULT_ASSEMBLED_JSONL,
    assembled_decisions_csv: Path = DEFAULT_ASSEMBLED_DECISIONS,
    assembled_review_dir: Path = DEFAULT_ASSEMBLED_REVIEW_DIR,
    group_review_csv: Path = DEFAULT_GROUP_REVIEW_CSV,
    group_review_decisions_csv: Path = DEFAULT_GROUP_REVIEW_DECISIONS,
    decimal_rollout_dir: Path = DEFAULT_DECIMAL_ROLLOUT_DIR,
    decimal_decisions_csv: Path = DEFAULT_DECIMAL_DECISIONS,
    decimal_decision_queue_csv: Path = DEFAULT_DECIMAL_DECISION_QUEUE,
    decimal_dp_decisions_csv: Path = DEFAULT_DECIMAL_DP_DECISIONS,
    tables_confirmed_csv: Path = DEFAULT_TABLES_CONFIRMED,
    facts_confirmed_csv: Path = DEFAULT_FACTS_CONFIRMED,
    facts_cleaned_csv: Path = DEFAULT_FACTS_CLEANED,
    revalidated_tables_csv: Path = DEFAULT_REVALIDATED_TABLES,
    reconciliation_annotations_json: Path = DEFAULT_RECONCILIATION_ANNOTATIONS,
    reconciliation_decisions_csv: Path = DEFAULT_RECONCILIATION_DECISIONS,
) -> Flask:
    app = Flask(__name__)
    packets = load_packets(packets_jsonl)
    results = read_csv(results_csv)
    decisions_by_id = by_id(read_csv(decisions_csv))
    continuation_rows = read_csv(continuation_candidates_csv)
    continuation_by_id = index_continuation_candidates(continuation_rows)
    continuation_by_group_id: dict[str, list[dict[str, str]]] = {}
    for continuation in continuation_rows:
        group_id = continuation.get("assigned_group_id", "")
        if group_id:
            continuation_by_group_id.setdefault(group_id, []).append(continuation)
    rows = []
    seen_row_ids = set()
    for row in results:
        table_id = row.get("table_id", "")
        seen_row_ids.add(table_id)
        if row.get("source_json"):
            seen_row_ids.add(row.get("source_json", ""))
        merged = dict(row)
        decision = decisions_by_id.get(table_id, {})
        merged["human_decision"] = decision.get("human_decision", "")
        merged["rerun_focus_area"] = decision.get("rerun_focus_area", "")
        merged["human_notes"] = decision.get("human_notes", "")
        continuation = continuation_by_id.get(table_id) or continuation_by_id.get(row.get("source_json", ""))
        if continuation:
            merged["is_continuation_candidate"] = "1"
            merged["continuation_auto_attach_reason"] = continuation.get("auto_attach_reason", "")
            merged["continuation_candidate_title"] = continuation.get("candidate_title", "")
            merged["continuation_candidate_source_json"] = continuation.get("candidate_source_json", "")
            merged["continuation_assigned_group_id"] = continuation.get("assigned_group_id", "")
            merged["continuation_needs_manual_multipage_review"] = continuation.get("needs_manual_multipage_review", "")
        else:
            merged["is_continuation_candidate"] = "0"
            merged["continuation_auto_attach_reason"] = ""
            merged["continuation_candidate_title"] = ""
            merged["continuation_candidate_source_json"] = ""
            merged["continuation_assigned_group_id"] = ""
            merged["continuation_needs_manual_multipage_review"] = ""
        rows.append(merged)
    for continuation in continuation_rows:
        table_id = continuation.get("table_id") or continuation.get("source_json", "")
        source_json = continuation.get("source_json", "")
        if table_id in seen_row_ids or source_json in seen_row_ids:
            continue
        merged = continuation_review_row(continuation)
        decision = decisions_by_id.get(table_id, {})
        merged["human_decision"] = decision.get("human_decision", "")
        merged["rerun_focus_area"] = decision.get("rerun_focus_area", "")
        merged["human_notes"] = decision.get("human_notes", "")
        rows.append(merged)
    rows.sort(key=lambda r: (r.get("human_decision") != "", r.get("llm_decision") == "pass", -(int(r.get("priority_score") or 0)), r.get("province", ""), r.get("year", ""), r.get("table_id", "")))

    assembled_groups = sorted(
        [group for group in load_assembled_groups(assembled_jsonl) if int_value(group.get("n_pages")) > 1],
        key=assembled_sort_key,
    )
    assembled_groups_by_id = {group.get("group_id", ""): group for group in assembled_groups if group.get("group_id")}
    assembled_decisions_by_id = assembled_decision_map(assembled_decisions_csv)
    write_assembled_review_exports(assembled_groups, assembled_decisions_csv, assembled_review_dir)
    group_review_rows = read_csv(group_review_csv)
    group_review_decisions_by_id = by_id(read_csv(group_review_decisions_csv), "group_id")
    for group_row in group_review_rows:
        decision = group_review_decisions_by_id.get(group_row.get("group_id", ""), {})
        group_row["human_decision"] = decision.get("human_decision", "")
        group_row["human_notes"] = decision.get("human_notes", "")
    group_review_index_by_group_id = {
        group_row.get("group_id", ""): idx
        for idx, group_row in enumerate(group_review_rows)
        if group_row.get("group_id")
    }
    write_group_review_exports(group_review_rows, group_review_decisions_csv, group_review_decisions_csv.parent)
    confirmed_table_rows = read_csv(tables_confirmed_csv)
    confirmed_table_key_index = table_key_index(confirmed_table_rows)
    decimal_opus_by_table_id = load_decimal_opus_results(
        decimal_rollout_dir / "opus",
        confirmed_table_rows,
        confirmed_table_key_index,
    )
    decimal_review_rows = load_decimal_review_rows(
        decimal_rollout_dir,
        confirmed_table_rows,
        confirmed_table_key_index,
        revalidated_tables_csv,
        decimal_opus_by_table_id,
    )
    decimal_decisions_by_id = decimal_decision_map(decimal_decisions_csv)
    for decimal_row in decimal_review_rows:
        decision = decimal_decisions_by_id.get(decimal_row.get("queue_id", ""), {})
        decimal_row["human_decision"] = decision.get("human_decision", "")
        decimal_row["human_notes"] = decision.get("human_notes", "")
    decimal_table_ids = {row.get("table_id", "") for row in decimal_review_rows if row.get("table_id")}
    decimal_decision_rows = load_decimal_decision_rows(
        decimal_decision_queue_csv,
        confirmed_table_rows,
        confirmed_table_key_index,
    )
    decimal_dp_decisions_by_id = decimal_dp_decision_map(decimal_dp_decisions_csv)
    for decimal_decision_row in decimal_decision_rows:
        decision = decimal_dp_decisions_by_id.get(decimal_decision_row.get("queue_id", ""), {})
        decimal_decision_row["human_dp"] = decision.get("human_dp", decimal_decision_row.get("human_dp", ""))
        decimal_decision_row["human_flag"] = decision.get("human_flag", decimal_decision_row.get("human_flag", ""))
        decimal_decision_row["human_notes"] = decision.get("human_notes", decimal_decision_row.get("human_notes", ""))
    decimal_decision_table_ids = {row.get("table_id", "") for row in decimal_decision_rows if row.get("table_id")}
    reconciliation_rows = load_reconciliation_rows(
        reconciliation_annotations_json,
        confirmed_table_rows,
        confirmed_table_key_index,
    )
    reconciliation_decisions_by_id = reconciliation_decision_map(reconciliation_decisions_csv)
    for reconciliation_row in reconciliation_rows:
        decision = reconciliation_decisions_by_id.get(reconciliation_row.get("table_id", ""), {})
        reconciliation_row["human_decision"] = decision.get("human_decision", "")
        reconciliation_row["human_notes"] = decision.get("human_notes", "")
    reconciliation_table_ids = {row.get("table_id", "") for row in reconciliation_rows if row.get("table_id")}
    review_fact_table_ids = decimal_table_ids | decimal_decision_table_ids | reconciliation_table_ids
    original_facts_by_table = load_fact_subset(facts_confirmed_csv, review_fact_table_ids)
    cleaned_facts_by_table = load_fact_subset(facts_cleaned_csv, review_fact_table_ids)
    write_decimal_review_exports(
        decimal_review_rows,
        decimal_decisions_csv,
        decimal_rollout_dir,
        original_facts_by_table,
        cleaned_facts_by_table,
    )
    write_decimal_decision_exports(
        decimal_decision_rows,
        decimal_dp_decisions_csv,
        decimal_rollout_dir,
    )
    write_reconciliation_review_exports(
        reconciliation_rows,
        reconciliation_decisions_csv,
        reconciliation_annotations_json.parent,
        original_facts_by_table,
        cleaned_facts_by_table,
    )

    filters = [
        ("flagged", "Flagged"),
        ("unresolved", "Unresolved"),
        ("reviewed", "Reviewed"),
        ("needs_image_compare", "Image Compare"),
        ("needs_manual_review", "Manual Review"),
        ("reject_or_rerun", "Reject/Rerun"),
        ("continuation", "Continuation"),
        ("continuation_manual", "Cont. Manual"),
        ("mortality", "Mortality"),
        ("priority", "Priority"),
        ("pass", "Pass"),
        ("all", "All"),
    ]

    assembled_filters = [
        ("unreviewed", "Unreviewed"),
        ("reviewed", "Reviewed"),
        ("accepted", "Accepted"),
        ("grouping_ok_rerun_facts", "Grouping OK / Rerun Facts"),
        ("rejected", "Rejected"),
        ("rerun", "Rerun"),
        ("hold", "Hold"),
        ("failed", "Constraint Fails"),
        ("passed", "Constraint Pass"),
        ("improved", "Improved"),
        ("clean", "Clean"),
        ("all", "All"),
    ]

    group_review_filters = [
        ("unreviewed", "Unreviewed"),
        ("reviewed", "Reviewed"),
        ("action_needed", "Action Needed"),
        ("ok", "OK"),
        ("merge_candidate", "Merge"),
        ("split_candidate", "Split"),
        ("incomplete_not_digitized", "Incomplete"),
        ("missing_preceding", "Missing Head"),
        ("has_duplicate", "Duplicate"),
        ("medium", "Medium"),
        ("high", "High"),
        ("all", "All"),
    ]

    decimal_filters = [
        ("residue", "Residue 115"),
        ("needs_review", "Decimal Risk"),
        ("validation_fail", "Validation Fail"),
        ("auto_fixed", "Auto Fixed"),
        ("unreviewed", "Unreviewed"),
        ("reviewed", "Reviewed"),
        ("revert_to_original", "Revert Original"),
        ("decimal_ok", "Decimal OK"),
        ("rerun_decimal_prompt", "Rerun"),
        ("manual_scale_fix", "Manual Fix"),
        ("exclude_for_now", "Exclude"),
        ("hold", "Hold"),
        ("all", "All"),
    ]

    decimal_decision_filters = [
        ("unreviewed", "Unreviewed"),
        ("reviewed", "Reviewed"),
        ("dp1_uniform", "dp1 Uniform"),
        ("dp2_uniform", "dp2 Uniform"),
        ("dp4_uniform", "dp4 Uniform"),
        ("NON_uniform_mixed", "Mixed Opus"),
        ("dp0_suspect", "dp0 Suspect"),
        ("divide", "Apply Divide"),
        ("no_divide", "No Divide"),
        ("rerun_or_hand_fix", "Rerun/Hand Fix"),
        ("all", "All"),
    ]

    reconciliation_filters = [
        ("unreviewed", "Unreviewed"),
        ("reviewed", "Reviewed"),
        ("fix_cell", "Fix-cell"),
        ("rerun_extract", "Rerun"),
        ("accept_artifact", "Accept Artifact"),
        ("exclude_for_now", "Exclude"),
        ("all", "All"),
    ]

    def first_index(mode: str) -> int:
        idxs = filter_indices(rows, mode)
        return idxs[0] if idxs else 0

    def nav(idx: int, mode: str) -> tuple[int | None, int | None, int | None]:
        idxs = filter_indices(rows, mode)
        if idx not in idxs and idxs:
            idx = idxs[0]
        pos = idxs.index(idx) if idx in idxs else -1
        prev_idx = idxs[pos - 1] if pos > 0 else None
        next_idx = idxs[pos + 1] if pos >= 0 and pos + 1 < len(idxs) else None
        unresolved = filter_indices(rows, "unresolved")
        next_unresolved = next((i for i in unresolved if i > idx), unresolved[0] if unresolved else None)
        return prev_idx, next_idx, next_unresolved

    def next_target_before_save(idx: int, mode: str) -> int:
        idxs = filter_indices(rows, mode)
        if not idxs:
            return idx
        if idx not in idxs:
            return idxs[0]
        pos = idxs.index(idx)
        if pos + 1 < len(idxs):
            return idxs[pos + 1]
        if len(idxs) > 1:
            return idxs[0]
        return idx

    def save_decision(idx: int, human_decision: str, rerun_focus_area: str, human_notes: str) -> None:
        row = rows[idx]
        table_id = row.get("table_id", "")
        existing = by_id(read_csv(decisions_csv))
        existing[table_id] = {
            "table_id": table_id,
            "human_decision": human_decision,
            "rerun_focus_area": rerun_focus_area,
            "human_notes": human_notes,
            "reviewer": "local",
            "reviewed_at": datetime.now().isoformat(timespec="seconds"),
            "llm_decision_at_review": row.get("llm_decision", ""),
            "llm_reason_at_review": row.get("reason", ""),
        }
        write_csv(decisions_csv, list(existing.values()), DECISION_FIELDS)
        row["human_decision"] = human_decision
        row["rerun_focus_area"] = rerun_focus_area
        row["human_notes"] = human_notes
        write_rerun_focus_queue(rows, decisions_csv, rerun_focus_queue)

    def assembled_filter_indices(mode: str) -> list[int]:
        mode = (mode or "unreviewed").strip()
        if mode == "all":
            return list(range(len(assembled_groups)))
        output = []
        for idx, group in enumerate(assembled_groups):
            decision = assembled_decisions_by_id.get(group.get("group_id", ""), {}).get("review_decision", "")
            report = assembled_report(group)
            if mode == "unreviewed" and not decision:
                output.append(idx)
            elif mode == "reviewed" and decision:
                output.append(idx)
            elif mode == "accepted" and decision == "keep":
                output.append(idx)
            elif mode == "grouping_ok_rerun_facts" and decision == "grouping_ok_rerun_facts":
                output.append(idx)
            elif mode == "rejected" and decision == "drop":
                output.append(idx)
            elif mode == "rerun" and decision == "rerun":
                output.append(idx)
            elif mode == "hold" and decision == "hold":
                output.append(idx)
            elif mode == "failed" and int_value(report.get("failed")) > 0:
                output.append(idx)
            elif mode == "passed" and int_value(report.get("total_checks")) > 0 and int_value(report.get("failed")) == 0:
                output.append(idx)
            elif mode == "improved" and int_value(group.get("input_constraint_failures")) > int_value(report.get("failed")):
                output.append(idx)
            elif mode == "clean" and int_value(report.get("failed")) == 0:
                output.append(idx)
        return output

    def assembled_first_index(mode: str) -> int:
        idxs = assembled_filter_indices(mode)
        return idxs[0] if idxs else 0

    def assembled_nav(idx: int, mode: str) -> tuple[int | None, int | None]:
        idxs = assembled_filter_indices(mode)
        if idx not in idxs and idxs:
            idx = idxs[0]
        pos = idxs.index(idx) if idx in idxs else -1
        prev_idx = idxs[pos - 1] if pos > 0 else None
        next_idx = idxs[pos + 1] if pos >= 0 and pos + 1 < len(idxs) else None
        return prev_idx, next_idx

    def next_assembled_target_before_save(idx: int, mode: str) -> int:
        idxs = assembled_filter_indices(mode)
        if not idxs:
            return idx
        if idx not in idxs:
            return idxs[0]
        pos = idxs.index(idx)
        if pos + 1 < len(idxs):
            return idxs[pos + 1]
        if len(idxs) > 1:
            return idxs[0]
        return idx

    def save_assembled_decision(idx: int, review_decision: str, review_notes: str) -> None:
        group = assembled_groups[idx]
        group_id = group.get("group_id", "")
        existing = assembled_decision_map(assembled_decisions_csv)
        report = assembled_report(group)
        if review_decision == "keep" and int_value(report.get("failed")) > 0:
            review_decision = "grouping_ok_rerun_facts"
        existing[group_id] = {
            "group_id": group_id,
            "table_id": stable_assembled_table_id(group_id),
            "review_decision": review_decision,
            "review_notes": review_notes,
            "reviewer": "local",
            "reviewed_at": datetime.now().isoformat(timespec="seconds"),
            "province": group.get("province", ""),
            "year": group.get("year", ""),
            "table_family": group.get("table_family", ""),
            "title": group.get("title", ""),
            "n_pages": group.get("n_pages", ""),
            "overall_strategy": group.get("overall_strategy", ""),
            "assembled_constraint_checks": report.get("total_checks", ""),
            "assembled_constraint_failures": report.get("failed", ""),
            "assembled_all_passed": report.get("all_passed", ""),
        }
        write_csv(assembled_decisions_csv, list(existing.values()), ASSEMBLED_DECISION_FIELDS)
        assembled_decisions_by_id.clear()
        assembled_decisions_by_id.update(existing)
        write_assembled_review_exports(assembled_groups, assembled_decisions_csv, assembled_review_dir)

    def group_review_filter_indices(mode: str) -> list[int]:
        mode = (mode or "unreviewed").strip()
        if mode == "all":
            return list(range(len(group_review_rows)))
        output = []
        for idx, row in enumerate(group_review_rows):
            human_decision = group_review_decisions_by_id.get(row.get("group_id", ""), {}).get("human_decision", "")
            claude_decision = row.get("decision", "")
            confidence = row.get("confidence", "")
            if mode == "unreviewed" and not human_decision:
                output.append(idx)
            elif mode == "reviewed" and human_decision:
                output.append(idx)
            elif mode == "action_needed" and (claude_decision not in {"ok", "ok_unnumbered"} or confidence != "high"):
                output.append(idx)
            elif mode == "ok" and claude_decision in {"ok", "ok_unnumbered"}:
                output.append(idx)
            elif mode in {"merge_candidate", "split_candidate", "incomplete_not_digitized", "missing_preceding", "has_duplicate"} and claude_decision == mode:
                output.append(idx)
            elif mode in {"high", "medium"} and confidence == mode:
                output.append(idx)
            elif mode in {"approve_proposal", "approve_ok", "needs_manual", "rerun_or_reextract", "reject_proposal"} and human_decision == mode:
                output.append(idx)
        return output

    def group_review_first_index(mode: str) -> int:
        idxs = group_review_filter_indices(mode)
        return idxs[0] if idxs else 0

    def group_review_nav(idx: int, mode: str) -> tuple[int | None, int | None]:
        idxs = group_review_filter_indices(mode)
        if idx not in idxs and idxs:
            idx = idxs[0]
        pos = idxs.index(idx) if idx in idxs else -1
        prev_idx = idxs[pos - 1] if pos > 0 else None
        next_idx = idxs[pos + 1] if pos >= 0 and pos + 1 < len(idxs) else None
        return prev_idx, next_idx

    def next_group_review_target_before_save(idx: int, mode: str) -> int:
        idxs = group_review_filter_indices(mode)
        if not idxs:
            return idx
        if idx not in idxs:
            return idxs[0]
        pos = idxs.index(idx)
        if pos + 1 < len(idxs):
            return idxs[pos + 1]
        if len(idxs) > 1:
            return idxs[0]
        return idx

    def group_review_merge_targets(row: dict[str, str]) -> list[dict[str, Any]]:
        targets = []
        for group_id in pipe_values(row.get("merge_with_group_ids")):
            target_row_idx = group_review_index_by_group_id.get(group_id)
            target_row = group_review_rows[target_row_idx] if target_row_idx is not None else {}
            target_group = assembled_groups_by_id.get(group_id, {})
            targets.append(
                {
                    "group_id": group_id,
                    "review_idx": target_row_idx,
                    "status": "in 1901 review CSV" if target_row else ("in assembled JSONL" if target_group else "not found locally"),
                    "title": target_group.get("title", "") or target_row.get("notes", ""),
                    "decision": target_row.get("decision", ""),
                    "confidence": target_row.get("confidence", ""),
                }
            )
        return targets

    def save_group_review_decision(idx: int, human_decision: str, human_notes: str) -> None:
        row = group_review_rows[idx]
        group_id = row.get("group_id", "")
        existing = by_id(read_csv(group_review_decisions_csv), "group_id")
        existing[group_id] = {
            "group_id": group_id,
            "human_decision": human_decision,
            "human_notes": human_notes,
            "reviewer": "local",
            "reviewed_at": datetime.now().isoformat(timespec="seconds"),
            "claude_decision": row.get("decision", ""),
            "confidence": row.get("confidence", ""),
            "approved_correct_first_source_json": row.get("correct_first_source_json", ""),
            "approved_member_order_source_jsons": row.get("correct_member_order_source_jsons", ""),
            "approved_merge_with_group_ids": row.get("merge_with_group_ids", ""),
            "approved_missing_preceding_source_jsons": row.get("missing_preceding_source_jsons", ""),
            "approved_extra_member_source_jsons": row.get("extra_member_source_jsons", ""),
            "approved_split_subgroups": row.get("proposed_split_subgroups", ""),
            "province": row.get("province", ""),
            "year": row.get("year", ""),
            "table_family": row.get("table_family", ""),
        }
        write_csv(group_review_decisions_csv, list(existing.values()), GROUP_REVIEW_DECISION_FIELDS)
        group_review_decisions_by_id.clear()
        group_review_decisions_by_id.update(existing)
        row["human_decision"] = human_decision
        row["human_notes"] = human_notes
        write_group_review_exports(group_review_rows, group_review_decisions_csv, group_review_decisions_csv.parent)

    def decimal_filter_indices(mode: str) -> list[int]:
        mode = (mode or "residue").strip()
        if mode == "all":
            return list(range(len(decimal_review_rows)))
        output = []
        for idx, row in enumerate(decimal_review_rows):
            human_decision = decimal_decisions_by_id.get(row.get("queue_id", ""), {}).get("human_decision", "")
            bucket = row.get("review_bucket", "")
            if mode == "residue" and bucket in {"needs_review", "validation_fail"}:
                output.append(idx)
            elif mode in {"needs_review", "validation_fail", "auto_fixed"} and bucket == mode:
                output.append(idx)
            elif mode == "unreviewed" and not human_decision:
                output.append(idx)
            elif mode == "reviewed" and human_decision:
                output.append(idx)
            elif mode in {"revert_to_original", "decimal_ok", "manual_scale_fix", "rerun_decimal_prompt", "exclude_for_now", "hold"} and human_decision == mode:
                output.append(idx)
        return output

    def decimal_first_index(mode: str) -> int:
        idxs = decimal_filter_indices(mode)
        return idxs[0] if idxs else 0

    def decimal_nav(idx: int, mode: str) -> tuple[int | None, int | None]:
        idxs = decimal_filter_indices(mode)
        if idx not in idxs and idxs:
            idx = idxs[0]
        pos = idxs.index(idx) if idx in idxs else -1
        prev_idx = idxs[pos - 1] if pos > 0 else None
        next_idx = idxs[pos + 1] if pos >= 0 and pos + 1 < len(idxs) else None
        return prev_idx, next_idx

    def next_decimal_target_before_save(idx: int, mode: str) -> int:
        idxs = decimal_filter_indices(mode)
        if not idxs:
            return idx
        if idx not in idxs:
            return idxs[0]
        pos = idxs.index(idx)
        if pos + 1 < len(idxs):
            return idxs[pos + 1]
        if len(idxs) > 1:
            return idxs[0]
        return idx

    def save_decimal_decision(idx: int, human_decision: str, human_notes: str) -> None:
        row = decimal_review_rows[idx]
        queue_id = row.get("queue_id", "")
        existing = decimal_decision_map(decimal_decisions_csv)
        existing[queue_id] = {
            "queue_id": queue_id,
            "human_decision": human_decision,
            "human_notes": human_notes,
            "reviewer": "local",
            "reviewed_at": datetime.now().isoformat(timespec="seconds"),
            "review_bucket": row.get("review_bucket", ""),
            "table_id": row.get("table_id", ""),
            "raw_table_id": row.get("raw_table_id", ""),
            "province": row.get("province", ""),
            "year": row.get("year", ""),
            "table_family": row.get("table_family", ""),
            "title": row.get("title", ""),
            "decimal_places": row.get("decimal_places", ""),
            "reason": row.get("reason", "") or row.get("rejected_reason", ""),
            "validation_failures": row.get("validation_failures", ""),
        }
        write_csv(decimal_decisions_csv, list(existing.values()), DECIMAL_DECISION_FIELDS)
        decimal_decisions_by_id.clear()
        decimal_decisions_by_id.update(existing)
        row["human_decision"] = human_decision
        row["human_notes"] = human_notes
        write_decimal_review_exports(
            decimal_review_rows,
            decimal_decisions_csv,
            decimal_rollout_dir,
            original_facts_by_table,
            cleaned_facts_by_table,
        )

    def decimal_decision_filter_indices(mode: str) -> list[int]:
        mode = (mode or "unreviewed").strip()
        if mode == "all":
            return list(range(len(decimal_decision_rows)))
        output = []
        for idx, row in enumerate(decimal_decision_rows):
            decision = decimal_dp_decisions_by_id.get(row.get("queue_id", ""), {})
            human_dp = decision.get("human_dp", row.get("human_dp", ""))
            human_flag = decision.get("human_flag", row.get("human_flag", ""))
            reviewed = bool(human_dp or human_flag)
            action = decimal_decision_action(row, {"human_dp": human_dp, "human_flag": human_flag})
            if mode == "unreviewed" and not reviewed:
                output.append(idx)
            elif mode == "reviewed" and reviewed:
                output.append(idx)
            elif mode in {"dp1_uniform", "dp2_uniform", "dp4_uniform", "NON_uniform_mixed", "dp0_suspect"} and row.get("category") == mode:
                output.append(idx)
            elif mode in {"divide", "no_divide", "rerun_or_hand_fix"} and action == mode:
                output.append(idx)
        return output

    def decimal_decision_first_index(mode: str) -> int:
        idxs = decimal_decision_filter_indices(mode)
        return idxs[0] if idxs else 0

    def decimal_decision_nav(idx: int, mode: str) -> tuple[int | None, int | None]:
        idxs = decimal_decision_filter_indices(mode)
        if idx not in idxs and idxs:
            idx = idxs[0]
        pos = idxs.index(idx) if idx in idxs else -1
        prev_idx = idxs[pos - 1] if pos > 0 else None
        next_idx = idxs[pos + 1] if pos >= 0 and pos + 1 < len(idxs) else None
        return prev_idx, next_idx

    def next_decimal_decision_target_before_save(idx: int, mode: str) -> int:
        idxs = decimal_decision_filter_indices(mode)
        if not idxs:
            return idx
        if idx not in idxs:
            return idxs[0]
        pos = idxs.index(idx)
        if pos + 1 < len(idxs):
            return idxs[pos + 1]
        if len(idxs) > 1:
            return idxs[0]
        return idx

    def save_decimal_dp_decision(idx: int, human_dp: str, human_flag: str, human_notes: str) -> None:
        row = decimal_decision_rows[idx]
        queue_id = row.get("queue_id", "")
        existing = decimal_dp_decision_map(decimal_dp_decisions_csv)
        existing[queue_id] = {
            "queue_id": queue_id,
            "table_id": row.get("table_id", ""),
            "human_dp": human_dp,
            "human_flag": human_flag,
            "human_notes": human_notes,
            "reviewer": "local",
            "reviewed_at": datetime.now().isoformat(timespec="seconds"),
            "category": row.get("category", ""),
            "opus_dp": row.get("opus_dp", ""),
            "opus_uniform": row.get("opus_uniform", ""),
            "opus_sample": row.get("opus_sample", ""),
            "rate_min": row.get("rate_min", ""),
            "rate_max": row.get("rate_max", ""),
            "province": row.get("province", ""),
            "year": row.get("year", ""),
            "table_family": row.get("table_family", ""),
            "title": row.get("title", ""),
        }
        write_csv(decimal_dp_decisions_csv, list(existing.values()), DECIMAL_DP_DECISION_FIELDS)
        decimal_dp_decisions_by_id.clear()
        decimal_dp_decisions_by_id.update(existing)
        row["human_dp"] = human_dp
        row["human_flag"] = human_flag
        row["human_notes"] = human_notes
        write_decimal_decision_exports(decimal_decision_rows, decimal_dp_decisions_csv, decimal_rollout_dir)

    def reconciliation_filter_indices(mode: str) -> list[int]:
        mode = (mode or "unreviewed").strip()
        if mode == "all":
            return list(range(len(reconciliation_rows)))
        output = []
        for idx, row in enumerate(reconciliation_rows):
            human_decision = reconciliation_decisions_by_id.get(row.get("table_id", ""), {}).get("human_decision", "")
            if mode == "unreviewed" and not human_decision:
                output.append(idx)
            elif mode == "reviewed" and human_decision:
                output.append(idx)
            elif mode in {"accept_artifact", "fix_cell", "rerun_extract", "exclude_for_now"} and human_decision == mode:
                output.append(idx)
        return output

    def reconciliation_first_index(mode: str) -> int:
        idxs = reconciliation_filter_indices(mode)
        return idxs[0] if idxs else 0

    def reconciliation_nav(idx: int, mode: str) -> tuple[int | None, int | None]:
        idxs = reconciliation_filter_indices(mode)
        if idx not in idxs and idxs:
            idx = idxs[0]
        pos = idxs.index(idx) if idx in idxs else -1
        prev_idx = idxs[pos - 1] if pos > 0 else None
        next_idx = idxs[pos + 1] if pos >= 0 and pos + 1 < len(idxs) else None
        return prev_idx, next_idx

    def next_reconciliation_target_before_save(idx: int, mode: str) -> int:
        idxs = reconciliation_filter_indices(mode)
        if not idxs:
            return idx
        if idx not in idxs:
            return idxs[0]
        pos = idxs.index(idx)
        if pos + 1 < len(idxs):
            return idxs[pos + 1]
        if len(idxs) > 1:
            return idxs[0]
        return idx

    def save_reconciliation_decision(idx: int, human_decision: str, human_notes: str) -> None:
        row = reconciliation_rows[idx]
        table_id = row.get("table_id", "")
        existing = reconciliation_decision_map(reconciliation_decisions_csv)
        existing[table_id] = {
            "table_id": table_id,
            "human_decision": human_decision,
            "human_notes": human_notes,
            "reviewer": "local",
            "reviewed_at": datetime.now().isoformat(timespec="seconds"),
            "province": row.get("province", ""),
            "year": row.get("year", ""),
            "table_family": row.get("table_family", ""),
            "title": row.get("title", ""),
            "n_failures": row.get("n_failures", ""),
        }
        write_csv(reconciliation_decisions_csv, list(existing.values()), RECONCILIATION_DECISION_FIELDS)
        reconciliation_decisions_by_id.clear()
        reconciliation_decisions_by_id.update(existing)
        row["human_decision"] = human_decision
        row["human_notes"] = human_notes
        write_reconciliation_review_exports(
            reconciliation_rows,
            reconciliation_decisions_csv,
            reconciliation_annotations_json.parent,
            original_facts_by_table,
            cleaned_facts_by_table,
        )

    @app.route("/")
    def index():
        return redirect(url_for("review_page", idx=first_index("flagged"), filter="flagged"))

    @app.route("/review/<int:idx>")
    def review_page(idx: int):
        if idx < 0 or idx >= len(rows):
            abort(404)
        filter_mode = request.args.get("filter", "flagged")
        row = rows[idx]
        packet = packets.get(row.get("table_id", ""), {})
        excel = resolve_excel(row)
        image = resolve_image(row)
        fact_dicts, facts_status = excel_sheet_dicts(excel, "Facts")
        matrix_rows, matrix_status = build_recorded_fact_matrix(fact_dicts)
        if matrix_rows:
            excel_status = f"{facts_status} | {matrix_status}"
        else:
            matrix_rows, legacy_status = excel_sheet_rows(excel, "Matrix")
            excel_status = f"{facts_status} | fallback_legacy_matrix | {legacy_status}"
        facts_rows = fact_rows_table(fact_dicts)
        if len(facts_rows) == 1:
            facts_rows = packet_facts_rows(packet)
        prev_idx, next_idx, next_unresolved_idx = nav(idx, filter_mode)
        packet_summary = json.dumps(
            {
                "deterministic_audit": packet.get("deterministic_audit", {}),
                "numeric_summary": {
                    k: packet.get("summary", {}).get(k)
                    for k in [
                        "fact_rows",
                        "nonzero_rows",
                        "zero_rows",
                        "min_value",
                        "max_value",
                        "sex_values",
                        "civil_condition_values",
                        "religion_values",
                        "age_labels",
                        "heuristic_flags",
                        "suspicious_examples",
                    ]
                },
                "recommended_checks": packet.get("recommended_checks", []),
            },
            indent=2,
            ensure_ascii=False,
        )
        return render_template_string(
            APP_HTML,
            row=row,
            index=idx,
            total=len(rows),
            filter_mode=filter_mode,
            filters=filters,
            first_index=first_index,
            assembled_total=len(assembled_groups),
            assembled_first_index=assembled_first_index,
            group_review_total=len(group_review_rows),
            group_review_first_index=group_review_first_index,
            decimal_total=len(decimal_review_rows),
            decimal_first_index=decimal_first_index,
            reconciliation_total=len(reconciliation_rows),
            reconciliation_first_index=reconciliation_first_index,
            prev_idx=prev_idx,
            next_idx=next_idx,
            next_unresolved_idx=next_unresolved_idx,
            has_image=bool(image),
            has_excel=bool(excel),
            decision_class=decision_class(row.get("llm_decision", "")),
            matrix_html=html_table(matrix_rows, max_rows=260, max_cols=180),
            facts_html=html_table(facts_rows, max_rows=170, max_cols=40),
            excel_status=excel_status,
            packet_summary=packet_summary,
        )

    @app.post("/review/<int:idx>/save")
    def save_page(idx: int):
        if idx < 0 or idx >= len(rows):
            abort(404)
        filter_mode = request.args.get("filter", "flagged")
        target = next_target_before_save(idx, filter_mode)
        human_decision = request.form.get("quick_decision") or request.form.get("human_decision", "")
        rerun_focus_area = request.form.get("rerun_focus_area", "")
        human_notes = request.form.get("human_notes", "")
        save_decision(idx, human_decision, rerun_focus_area, human_notes)
        return redirect(url_for("review_page", idx=target, filter=filter_mode))

    @app.route("/source-image/<int:idx>")
    def source_image_file(idx: int):
        if idx < 0 or idx >= len(rows):
            abort(404)
        path = resolve_image(rows[idx])
        if not path:
            abort(404)
        return send_file(path)

    @app.route("/excel/<int:idx>")
    def excel_file(idx: int):
        if idx < 0 or idx >= len(rows):
            abort(404)
        path = resolve_excel(rows[idx])
        if not path:
            abort(404)
        return send_file(path, as_attachment=False)

    @app.route("/assembled-review")
    def assembled_review_home():
        if not assembled_groups:
            return "No assembled groups loaded.", 404
        filter_mode = request.args.get("filter", "unreviewed")
        return redirect(url_for("assembled_review_page", idx=assembled_first_index(filter_mode), filter=filter_mode))

    @app.route("/assembled-review/<int:idx>")
    def assembled_review_page(idx: int):
        if not assembled_groups:
            return "No assembled groups loaded.", 404
        if idx < 0 or idx >= len(assembled_groups):
            abort(404)
        filter_mode = request.args.get("filter", "unreviewed")
        valid = assembled_filter_indices(filter_mode)
        if valid and idx not in valid:
            idx = valid[0]
        group = assembled_groups[idx]
        group_id = group.get("group_id", "")
        decision = assembled_decisions_by_id.get(group_id, {})
        report = assembled_report(group)
        anchor = assembled_anchor_for_group(group, continuation_by_group_id)
        context = previous_table_context(anchor)
        display_title = anchor.get("title") or group.get("title", "")
        prev_idx, next_idx = assembled_nav(idx, filter_mode)
        return render_template_string(
            ASSEMBLED_HTML,
            group=group,
            table_id=stable_assembled_table_id(group_id),
            display_title=display_title,
            anchor=anchor,
            decision=decision,
            report=report,
            index=idx,
            total=len(assembled_groups),
            filter_mode=filter_mode,
            filters=assembled_filters,
            first_index=assembled_first_index,
            normal_first_index=first_index("continuation"),
            group_review_total=len(group_review_rows),
            group_review_first_index=group_review_first_index,
            decimal_total=len(decimal_review_rows),
            decimal_first_index=decimal_first_index,
            reconciliation_total=len(reconciliation_rows),
            reconciliation_first_index=reconciliation_first_index,
            prev_idx=prev_idx,
            next_idx=next_idx,
            member_images=assembled_member_images(group, anchor),
            previous_context=context,
            sections=assembled_section_views(group),
            section_count=len(group.get("sections") or []),
        )

    @app.post("/assembled-review/<int:idx>/save")
    def save_assembled_page(idx: int):
        if idx < 0 or idx >= len(assembled_groups):
            abort(404)
        filter_mode = request.args.get("filter", "unreviewed")
        target = next_assembled_target_before_save(idx, filter_mode)
        review_decision = request.form.get("quick_decision") or request.form.get("review_decision", "")
        review_notes = request.form.get("review_notes", "")
        save_assembled_decision(idx, review_decision, review_notes)
        return redirect(url_for("assembled_review_page", idx=target, filter=filter_mode))

    @app.route("/assembled-member-image/<int:idx>/<int:member_idx>")
    def assembled_member_image_file(idx: int, member_idx: int):
        if idx < 0 or idx >= len(assembled_groups):
            abort(404)
        group = assembled_groups[idx]
        anchor = assembled_anchor_for_group(group, continuation_by_group_id)
        images = assembled_member_images(group, anchor)
        if member_idx < 0 or member_idx >= len(images):
            abort(404)
        path_text = images[member_idx].get("local_path", "")
        path = Path(path_text) if path_text else None
        if not path or not path.exists():
            abort(404)
        return send_file(path)

    @app.route("/assembled-context-image/<int:idx>/<int:context_idx>")
    def assembled_context_image_file(idx: int, context_idx: int):
        if idx < 0 or idx >= len(assembled_groups):
            abort(404)
        group = assembled_groups[idx]
        anchor = assembled_anchor_for_group(group, continuation_by_group_id)
        context = previous_table_context(anchor)
        images = context.get("images") or []
        if context_idx < 0 or context_idx >= len(images):
            abort(404)
        path_text = images[context_idx].get("local_path", "")
        path = Path(path_text) if path_text else None
        if not path or not path.exists():
            abort(404)
        return send_file(path)

    @app.route("/group-review")
    def group_review_home():
        if not group_review_rows:
            return "No group review CSV loaded.", 404
        filter_mode = request.args.get("filter", "unreviewed")
        return redirect(url_for("group_review_page", idx=group_review_first_index(filter_mode), filter=filter_mode))

    @app.route("/group-review/<int:idx>")
    def group_review_page(idx: int):
        if not group_review_rows:
            return "No group review CSV loaded.", 404
        if idx < 0 or idx >= len(group_review_rows):
            abort(404)
        filter_mode = request.args.get("filter", "unreviewed")
        valid = group_review_filter_indices(filter_mode)
        if valid and idx not in valid:
            idx = valid[0]
        row = group_review_rows[idx]
        group_id = row.get("group_id", "")
        decision = group_review_decisions_by_id.get(group_id, {})
        current_group = assembled_groups_by_id.get(group_id)
        current_report = assembled_report(current_group or {})
        images = group_review_images(row, current_group)
        prev_idx, next_idx = group_review_nav(idx, filter_mode)
        return render_template_string(
            GROUP_REVIEW_HTML,
            row=row,
            decision=decision,
            decision_class=group_review_decision_class(row.get("decision", "")),
            index=idx,
            total=len(group_review_rows),
            filter_mode=filter_mode,
            filters=group_review_filters,
            first_index=group_review_first_index,
            assembled_first=assembled_first_index("all"),
            prev_idx=prev_idx,
            next_idx=next_idx,
            current_group=current_group,
            current_report=current_report,
            current_images=images["current"],
            proposed_images=images["proposed"],
            evidence_images=images["evidence"],
            merge_targets=group_review_merge_targets(row),
        )

    @app.post("/group-review/<int:idx>/save")
    def save_group_review_page(idx: int):
        if idx < 0 or idx >= len(group_review_rows):
            abort(404)
        filter_mode = request.args.get("filter", "unreviewed")
        target = next_group_review_target_before_save(idx, filter_mode)
        human_decision = request.form.get("quick_decision") or request.form.get("human_decision", "")
        human_notes = request.form.get("human_notes", "")
        save_group_review_decision(idx, human_decision, human_notes)
        return redirect(url_for("group_review_page", idx=target, filter=filter_mode))

    @app.route("/group-review-image/<int:idx>/<int:image_idx>")
    def group_review_image_file(idx: int, image_idx: int):
        if idx < 0 or idx >= len(group_review_rows):
            abort(404)
        row = group_review_rows[idx]
        current_group = assembled_groups_by_id.get(row.get("group_id", ""))
        images = group_review_images(row, current_group)["all"]
        if image_idx < 0 or image_idx >= len(images):
            abort(404)
        path_text = images[image_idx].get("local_path", "")
        path = Path(path_text) if path_text else None
        if not path or not path.exists():
            abort(404)
        return send_file(path)

    @app.route("/decimal-review")
    def decimal_review_home():
        if not decimal_review_rows:
            return "No decimal review rows loaded.", 404
        filter_mode = request.args.get("filter", "residue")
        return redirect(url_for("decimal_review_page", idx=decimal_first_index(filter_mode), filter=filter_mode))

    @app.route("/decimal-review/<int:idx>")
    def decimal_review_page(idx: int):
        if not decimal_review_rows:
            return "No decimal review rows loaded.", 404
        if idx < 0 or idx >= len(decimal_review_rows):
            abort(404)
        filter_mode = request.args.get("filter", "residue")
        valid = decimal_filter_indices(filter_mode)
        if valid and idx not in valid:
            idx = valid[0]
        row = decimal_review_rows[idx]
        decision = decimal_decisions_by_id.get(row.get("queue_id", ""), {})
        table_id = row.get("table_id", "")
        original_facts = original_facts_by_table.get(table_id, [])
        cleaned_facts = cleaned_facts_by_table.get(table_id, [])
        changed_rows = decimal_changed_rows(original_facts, cleaned_facts, max_rows=500)
        comparison_matrix, comparison_status = build_before_after_fact_matrix(original_facts, cleaned_facts)
        original_matrix, original_status = build_recorded_fact_matrix(original_facts)
        cleaned_matrix, cleaned_status = build_recorded_fact_matrix(cleaned_facts)
        prev_idx, next_idx = decimal_nav(idx, filter_mode)
        image = resolve_image(row)
        opus_summary = "; ".join(
            part
            for part in [
                f"dp={row.get('opus_decimal_places', '')}" if row.get("opus_decimal_places") else "",
                f"uniform={row.get('opus_uniform', '')}" if row.get("opus_uniform") else "",
                f"looks_like={row.get('opus_looks_like', '')}" if row.get("opus_looks_like") else "",
                row.get("opus_note", ""),
            ]
            if part
        )
        return render_template_string(
            DECIMAL_REVIEW_HTML,
            row=row,
            decision=decision,
            index=idx,
            total=len(decimal_review_rows),
            filter_mode=filter_mode,
            filters=decimal_filters,
            first_index=decimal_first_index,
            normal_first_index=first_index("flagged"),
            prev_idx=prev_idx,
            next_idx=next_idx,
            has_image=bool(image),
            bucket_class=decimal_bucket_class(row.get("review_bucket", "")),
            original_count=len(original_facts),
            cleaned_count=len(cleaned_facts),
            changed_count=max(len(changed_rows) - 1, 0),
            comparison_matrix_html=html_table(comparison_matrix, max_rows=240, max_cols=240),
            comparison_status=comparison_status,
            changed_html=html_table(changed_rows, max_rows=500, max_cols=24),
            original_matrix_html=html_table(original_matrix, max_rows=240, max_cols=180),
            cleaned_matrix_html=html_table(cleaned_matrix, max_rows=240, max_cols=180),
            original_status=original_status,
            cleaned_status=cleaned_status,
            opus_summary=opus_summary or "no Opus result matched",
            opus_sample=row.get("opus_sample", ""),
        )

    @app.post("/decimal-review/<int:idx>/save")
    def save_decimal_review_page(idx: int):
        if idx < 0 or idx >= len(decimal_review_rows):
            abort(404)
        filter_mode = request.args.get("filter", "residue")
        target = next_decimal_target_before_save(idx, filter_mode)
        human_decision = request.form.get("quick_decision") or request.form.get("human_decision", "")
        human_notes = request.form.get("human_notes", "")
        save_decimal_decision(idx, human_decision, human_notes)
        return redirect(url_for("decimal_review_page", idx=target, filter=filter_mode))

    @app.route("/decimal-image/<int:idx>")
    def decimal_image_file(idx: int):
        if idx < 0 or idx >= len(decimal_review_rows):
            abort(404)
        path = resolve_image(decimal_review_rows[idx])
        if not path:
            abort(404)
        return send_file(path)

    @app.route("/decimal-decision-review")
    def decimal_decision_home():
        if not decimal_decision_rows:
            return "No decimal decision queue loaded.", 404
        filter_mode = request.args.get("filter", "unreviewed")
        return redirect(url_for("decimal_decision_page", idx=decimal_decision_first_index(filter_mode), filter=filter_mode))

    @app.route("/decimal-decision-review/<int:idx>")
    def decimal_decision_page(idx: int):
        if not decimal_decision_rows:
            return "No decimal decision queue loaded.", 404
        if idx < 0 or idx >= len(decimal_decision_rows):
            abort(404)
        filter_mode = request.args.get("filter", "unreviewed")
        valid = decimal_decision_filter_indices(filter_mode)
        if valid and idx not in valid:
            idx = valid[0]
        row = decimal_decision_rows[idx]
        decision = decimal_dp_decisions_by_id.get(row.get("queue_id", ""), {})
        display_decision = {
            "human_dp": decision.get("human_dp", row.get("human_dp", "")),
            "human_flag": decision.get("human_flag", row.get("human_flag", "")),
            "human_notes": decision.get("human_notes", row.get("human_notes", "")),
        }
        table_id = row.get("table_id", "")
        facts = cleaned_facts_by_table.get(table_id, []) or original_facts_by_table.get(table_id, [])
        matrix_rows, matrix_status = build_recorded_fact_matrix(facts)
        prev_idx, next_idx = decimal_decision_nav(idx, filter_mode)
        category_class = "warn" if row.get("category") in {"NON_uniform_mixed", "dp0_suspect"} else "ok"
        return render_template_string(
            DECIMAL_DECISION_HTML,
            row=row,
            decision=display_decision,
            decision_label=decimal_decision_label(display_decision),
            index=idx,
            total=len(decimal_decision_rows),
            filter_mode=filter_mode,
            filters=decimal_decision_filters,
            first_index=decimal_decision_first_index,
            decimal_first=decimal_first_index("residue"),
            prev_idx=prev_idx,
            next_idx=next_idx,
            has_image=bool(decimal_decision_image(row)),
            category_class=category_class,
            facts_count=len(facts),
            matrix_html=html_table(matrix_rows, max_rows=260, max_cols=180),
            matrix_status=matrix_status,
        )

    @app.post("/decimal-decision-review/<int:idx>/save")
    def save_decimal_decision_page(idx: int):
        if idx < 0 or idx >= len(decimal_decision_rows):
            abort(404)
        filter_mode = request.args.get("filter", "unreviewed")
        target = next_decimal_decision_target_before_save(idx, filter_mode)
        human_dp = request.form.get("quick_dp", "")
        human_flag = request.form.get("quick_flag", "")
        human_notes = request.form.get("human_notes", "")
        save_decimal_dp_decision(idx, human_dp, human_flag, human_notes)
        return redirect(url_for("decimal_decision_page", idx=target, filter=filter_mode))

    @app.route("/decimal-decision-image/<int:idx>")
    def decimal_decision_image_file(idx: int):
        if idx < 0 or idx >= len(decimal_decision_rows):
            abort(404)
        path = decimal_decision_image(decimal_decision_rows[idx])
        if not path:
            abort(404)
        return send_file(path)

    @app.route("/reconciliation-review")
    def reconciliation_review_home():
        if not reconciliation_rows:
            return "No reconciliation annotations loaded.", 404
        filter_mode = request.args.get("filter", "unreviewed")
        return redirect(url_for("reconciliation_review_page", idx=reconciliation_first_index(filter_mode), filter=filter_mode))

    @app.route("/reconciliation-review/<int:idx>")
    def reconciliation_review_page(idx: int):
        if not reconciliation_rows:
            return "No reconciliation annotations loaded.", 404
        if idx < 0 or idx >= len(reconciliation_rows):
            abort(404)
        filter_mode = request.args.get("filter", "unreviewed")
        valid = reconciliation_filter_indices(filter_mode)
        if valid and idx not in valid:
            idx = valid[0]
        row = reconciliation_rows[idx]
        table_id = row.get("table_id", "")
        decision = reconciliation_decisions_by_id.get(table_id, {})
        cleaned_facts = cleaned_facts_by_table.get(table_id, [])
        original_facts = original_facts_by_table.get(table_id, [])
        facts_for_display = cleaned_facts or original_facts
        matrix_html, matrix_status = highlighted_reconciliation_matrix(facts_for_display, row.get("failures") or [])
        prev_idx, next_idx = reconciliation_nav(idx, filter_mode)
        return render_template_string(
            RECONCILIATION_REVIEW_HTML,
            row=row,
            failures=row.get("failures") or [],
            decision=decision,
            index=idx,
            total=len(reconciliation_rows),
            filter_mode=filter_mode,
            filters=reconciliation_filters,
            first_index=reconciliation_first_index,
            decimal_first=decimal_first_index("residue"),
            prev_idx=prev_idx,
            next_idx=next_idx,
            has_image=bool(resolve_reconciliation_image(row)),
            original_count=len(original_facts),
            cleaned_count=len(cleaned_facts),
            matrix_html=matrix_html,
            matrix_status=matrix_status,
        )

    @app.post("/reconciliation-review/<int:idx>/save")
    def save_reconciliation_review_page(idx: int):
        if idx < 0 or idx >= len(reconciliation_rows):
            abort(404)
        filter_mode = request.args.get("filter", "unreviewed")
        target = next_reconciliation_target_before_save(idx, filter_mode)
        human_decision = request.form.get("quick_decision") or request.form.get("human_decision", "")
        human_notes = request.form.get("human_notes", "")
        save_reconciliation_decision(idx, human_decision, human_notes)
        return redirect(url_for("reconciliation_review_page", idx=target, filter=filter_mode))

    @app.route("/reconciliation-image/<int:idx>")
    def reconciliation_image_file(idx: int):
        if idx < 0 or idx >= len(reconciliation_rows):
            abort(404)
        path = resolve_reconciliation_image(reconciliation_rows[idx])
        if not path:
            abort(404)
        return send_file(path)

    return app


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--results-csv", type=Path, default=DEFAULT_RESULTS)
    parser.add_argument("--packets-jsonl", type=Path, default=DEFAULT_PACKETS)
    parser.add_argument("--decisions-csv", type=Path, default=DEFAULT_DECISIONS)
    parser.add_argument("--rerun-focus-queue", type=Path, default=DEFAULT_RERUN_FOCUS_QUEUE)
    parser.add_argument("--continuation-candidates-csv", type=Path, default=DEFAULT_CONTINUATION_CANDIDATES)
    parser.add_argument("--assembled-jsonl", type=Path, default=DEFAULT_ASSEMBLED_JSONL)
    parser.add_argument("--assembled-decisions-csv", type=Path, default=DEFAULT_ASSEMBLED_DECISIONS)
    parser.add_argument("--assembled-review-dir", type=Path, default=DEFAULT_ASSEMBLED_REVIEW_DIR)
    parser.add_argument("--group-review-csv", type=Path, default=DEFAULT_GROUP_REVIEW_CSV)
    parser.add_argument("--group-review-decisions-csv", type=Path, default=DEFAULT_GROUP_REVIEW_DECISIONS)
    parser.add_argument("--decimal-rollout-dir", type=Path, default=DEFAULT_DECIMAL_ROLLOUT_DIR)
    parser.add_argument("--decimal-decisions-csv", type=Path, default=DEFAULT_DECIMAL_DECISIONS)
    parser.add_argument("--decimal-decision-queue-csv", type=Path, default=DEFAULT_DECIMAL_DECISION_QUEUE)
    parser.add_argument("--decimal-dp-decisions-csv", type=Path, default=DEFAULT_DECIMAL_DP_DECISIONS)
    parser.add_argument("--tables-confirmed-csv", type=Path, default=DEFAULT_TABLES_CONFIRMED)
    parser.add_argument("--facts-confirmed-csv", type=Path, default=DEFAULT_FACTS_CONFIRMED)
    parser.add_argument("--facts-cleaned-csv", type=Path, default=DEFAULT_FACTS_CLEANED)
    parser.add_argument("--revalidated-tables-csv", type=Path, default=DEFAULT_REVALIDATED_TABLES)
    parser.add_argument("--reconciliation-annotations-json", type=Path, default=DEFAULT_RECONCILIATION_ANNOTATIONS)
    parser.add_argument("--reconciliation-decisions-csv", type=Path, default=DEFAULT_RECONCILIATION_DECISIONS)
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=APP_PORT)
    args = parser.parse_args()
    if not args.results_csv.exists():
        raise SystemExit(f"Missing results CSV: {args.results_csv}. Run run_gemini_plausibility_review.py first.")
    app = create_app(
        args.results_csv,
        args.packets_jsonl,
        args.decisions_csv,
        args.rerun_focus_queue,
        args.continuation_candidates_csv,
        args.assembled_jsonl,
        args.assembled_decisions_csv,
        args.assembled_review_dir,
        args.group_review_csv,
        args.group_review_decisions_csv,
        args.decimal_rollout_dir,
        args.decimal_decisions_csv,
        args.decimal_decision_queue_csv,
        args.decimal_dp_decisions_csv,
        args.tables_confirmed_csv,
        args.facts_confirmed_csv,
        args.facts_cleaned_csv,
        args.revalidated_tables_csv,
        args.reconciliation_annotations_json,
        args.reconciliation_decisions_csv,
    )
    print(f"Review app: http://{args.host}:{args.port}")
    print(f"Decisions CSV: {args.decisions_csv}")
    print(f"Rerun focus queue: {args.rerun_focus_queue}")
    print(f"Continuation candidates: {args.continuation_candidates_csv}")
    print(f"Assembled JSONL: {args.assembled_jsonl}")
    print(f"Assembled decisions CSV: {args.assembled_decisions_csv}")
    print(f"Assembled review exports: {args.assembled_review_dir}")
    print(f"Group review CSV: {args.group_review_csv}")
    print(f"Group review decisions CSV: {args.group_review_decisions_csv}")
    print(f"Decimal rollout dir: {args.decimal_rollout_dir}")
    print(f"Decimal decisions CSV: {args.decimal_decisions_csv}")
    print(f"Decimal decision queue CSV: {args.decimal_decision_queue_csv}")
    print(f"Decimal DP decisions CSV: {args.decimal_dp_decisions_csv}")
    print(f"Facts cleaned CSV: {args.facts_cleaned_csv}")
    print(f"Reconciliation annotations: {args.reconciliation_annotations_json}")
    print(f"Reconciliation decisions CSV: {args.reconciliation_decisions_csv}")
    app.run(host=args.host, port=args.port, debug=False)


if __name__ == "__main__":
    main()
